"""Importação do BD_Operações (.xlsx enviado pela tela) para o ensaio do Nexus.

É a ÚNICA porta de Excel do Nexus, e é temporária: existe para o ensaio e para a virada (Levi, 29/09/2026:
"ensaio e depois virada"). O Nexus não lê o OneDrive; quem manda o arquivo é o admin, pela tela.

O que a importação garante:
- nada se perde: valor fora do tipo vira Legado (fica o texto, com aviso);
- fórmula e digitação na mesma coluna se distinguem: onde era fórmula, o campo fica AUTOMÁTICO; onde alguém
  digitou por cima, vira sobreposição (76 técnicos da aba Operações estavam digitados);
- id estável entre importações: usina pelo IDUsina, pessoa pelo e-mail (ou pelo nome), equipe pelo nome como o
  PROCV do Excel casa (sem diferenciar maiúscula). Importar de novo no ensaio não cria registro nem versão;
- o que o Excel mostrava é comparado com o que o Nexus calcula, valor a valor, e a diferença vem com o motivo
  quando ele é conhecido.
"""
import hashlib
import io
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any

from . import calculos, tipos
from .calculos import ERRO, UFS, Contexto
from .esquema import ENTIDADES, PESSOAS, USINAS
from .servico import QUEM_IMPORTACAO, Carga, chave_texto
from .tipos import Legado

ABAS = ("Operações", "Relação Geral Colaboradores", "Auxiliar", "Parametros")
VINCULOS = ["Colaborador de campo", "Supervisor", "Gestor de contrato"]
# A fórmula de receita de 70 usinas: preço fixo por MWp × potência contratual.
_RE_PRECO = re.compile(r"^=\s*([\d.,]+)\s*\*\s*Operacoes\[\[#This Row\],\[POT[ÊE]NCIA CONTRATUAL \(MWp\)\]\]\s*$",
                       re.I)
# Conta com número fixo ("=9.633/3", "=2160+1080"): é digitação feita com calculadora, não regra.
_RE_CONTA = re.compile(r"^=[\d\s.,+\-*/()]+$")


def chave_cab(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def chave_equipe(s) -> str:
    """Como o PROCV do Excel casa texto: sem diferenciar maiúscula. As pontas saem (a importação tira)."""
    return re.sub(r" +", " ", str(s or "").strip()).lower()


@dataclass
class Celula:
    valor: Any = None               # o que o Excel mostrava (valor em cache)
    formula: str | None = None      # o texto da fórmula, quando a célula era fórmula


class Aba:
    def __init__(self, cabecalho, linhas, validacoes=None):
        self.cabecalho = [str(c).strip() if c is not None else "" for c in cabecalho]
        self.linhas = linhas
        self.validacoes = validacoes or {}
        self._idx = {}
        for i, c in enumerate(self.cabecalho):
            k = chave_cab(c)
            if k and k not in self._idx:
                self._idx[k] = i

    def indice(self, nome) -> int | None:
        return self._idx.get(chave_cab(nome))

    def celula(self, linha, nome) -> Celula:
        i = self.indice(nome)
        if i is None or i >= len(linha):
            return Celula()
        return linha[i]

    def coluna(self, nome) -> list:
        return [self.celula(ln, nome).valor for ln in self.linhas]


@dataclass
class Proposta:
    carga: Carga
    resumo: dict
    avisos: list
    comparacao: list
    origem: dict
    pessoas_por_id: dict = field(default_factory=dict)


# ── leitura do arquivo ─────────────────────────────────────────────────────

def _texto_formula(f):
    if isinstance(f, str) and f.startswith("="):
        return f
    texto = getattr(f, "text", None)                     # ArrayFormula do openpyxl
    return texto if isinstance(texto, str) else None


def ler_xlsx(dados: bytes) -> dict:
    """As quatro abas que o cadastro usa, cada célula com o valor mostrado e a fórmula (quando havia)."""
    import openpyxl

    wf = openpyxl.load_workbook(io.BytesIO(dados), data_only=False)
    wv = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
    por_chave = {chave_cab(n): n for n in wv.sheetnames}
    out = {}
    for nome in ABAS:
        real = por_chave.get(chave_cab(nome))
        if real is None:
            continue
        wsf, wsv = wf[real], wv[real]
        rf, rv = list(wsf.iter_rows(values_only=True)), list(wsv.iter_rows(values_only=True))
        if not rv:
            continue
        cab = list(rv[0])
        linhas = []
        for lf, lv in zip(rf[1:], rv[1:]):
            cels = [Celula(lv[j] if j < len(lv) else None, _texto_formula(lf[j] if j < len(lf) else None))
                    for j in range(len(cab))]
            if any(c.valor not in (None, "") for c in cels):
                linhas.append(cels)
        validacoes = {}
        for dv in getattr(wsf.data_validations, "dataValidation", []) or []:
            f1 = str(dv.formula1 or "")
            if dv.type != "list" or not f1.startswith('"'):
                continue
            primeira = str(dv.sqref).split()[0].split(":")[0]
            col = openpyxl.utils.cell.coordinate_from_string(primeira)[0]
            j = openpyxl.utils.column_index_from_string(col) - 1
            if j < len(cab) and cab[j]:
                itens = validacoes.setdefault(str(cab[j]).strip(), [])
                for x in f1.strip('"').split(","):
                    if x.strip() and x.strip() not in itens:
                        itens.append(x.strip())
        out[nome] = Aba(cab, linhas, validacoes)
    return out


# ── montagem ───────────────────────────────────────────────────────────────

def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _juntar(*fontes) -> list:
    out = []
    for fonte in fontes:
        for v in fonte or []:
            t = _texto(v)
            if t and t not in out:
                out.append(t)
    return out


def _proximo(ids):
    """Ids novos em número simples, 1, 2, 3 (Levi, 29/09). Id antigo com prefixo (E-001) não entra na conta."""
    n = max((int(i) for i in ids if str(i).isdigit()), default=0)
    while True:
        n += 1
        yield str(n)


def montar(abas: dict, srv, arquivo: str = "", dados: bytes | None = None) -> Proposta:
    avisos, formula_valor, faltando = [], Counter(), set()
    ops = abas.get("Operações") or Aba([], [])
    pes = abas.get("Relação Geral Colaboradores") or Aba([], [])
    aux = abas.get("Auxiliar") or Aba([], [])
    par = abas.get("Parametros") or Aba([], [])
    atuais = {eid: (srv.registros(eid, incluir_excluidos=True) if srv else []) for eid in ENTIDADES}

    def cel(aba, linha, coluna):
        if aba.indice(coluna) is None and aba.linhas:
            faltando.add(f"{coluna} ({'Operações' if aba is ops else 'Relação Geral Colaboradores'})")
        return aba.celula(linha, coluna)

    # ── equipes: a "Equipe Cluster" das usinas e o "Cluster" das pessoas ──
    nomes_eq, vistos = [], set()
    for t in [*_juntar(ops.coluna("Equipe Cluster")), *_juntar(pes.coluna("Cluster"))]:
        k = chave_equipe(t)
        if k not in vistos:
            vistos.add(k)
            nomes_eq.append(t)
    eq_atual = {chave_equipe(r.valores.get("nome")): r for r in atuais["equipes"]}
    gera_eq = _proximo([r.id for r in atuais["equipes"]])
    equipes, eq_id = [], {}
    for i, nome in enumerate(nomes_eq, start=1):
        r = eq_atual.get(chave_equipe(nome))
        id_ = r.id if r else next(gera_eq)
        eq_id[chave_equipe(nome)] = id_
        equipes.append({"id": id_, "ordem": i, "valores": {
            "nome": nome, "observacao": r.valores.get("observacao") if r else None}})

    def equipe_de(texto):
        t = _texto(texto)
        if not t or t == "0":
            return None
        return eq_id.get(chave_equipe(t)) or Legado(t)

    # ── pessoas: a Relação Geral, e depois quem só existe nas listas do Auxiliar ──
    p_atual_email = {str(r.valores.get("email")).lower(): r for r in atuais["pessoas"]
                     if isinstance(r.valores.get("email"), str) and r.valores.get("email")}
    p_atual_nome = defaultdict(list)
    for r in atuais["pessoas"]:
        p_atual_nome[(chave_texto(r.valores.get("nome")), r.valores.get("vinculo"))].append(r)
    gera_p = _proximo([r.id for r in atuais["pessoas"]])
    pessoas, usados, sup_digitado = [], set(), {}

    def id_pessoa(email, nome, vinculo):
        r = p_atual_email.get(str(email).lower()) if email else None
        if r is None:
            cands = [x for x in p_atual_nome.get((chave_texto(nome), vinculo), []) if x.id not in usados]
            r = cands[0] if len(cands) == 1 else None
        id_ = r.id if r and r.id not in usados else next(gera_p)
        usados.add(id_)
        return id_

    for i, ln in enumerate(pes.linhas, start=1):
        valores = {}
        for c in PESSOAS.campos:
            if not c.coluna_bd or c.modo == "calculado":
                continue
            x = cel(pes, ln, c.coluna_bd)
            if c.modo == "automatico" and x.formula:
                valores[c.id] = None
                continue
            if c.id == "equipe":
                valores[c.id] = equipe_de(x.valor)
            elif c.id == "supervisor":
                t = _texto(x.valor)
                valores[c.id] = None
                if t and t != "0":
                    sup_digitado[i] = t
            else:
                valores[c.id] = tipos.do_excel(c.tipo, x.valor)
        valores["vinculo"] = "Colaborador de campo"
        id_ = id_pessoa(valores.get("email") if isinstance(valores.get("email"), str) else None,
                        valores.get("nome"), "Colaborador de campo")
        pessoas.append({"id": id_, "ordem": i, "valores": valores})

    def _np(p):
        v = p["valores"]
        n = v.get("nome_padrao") if not tipos._vazio(v.get("nome_padrao")) else calculos.nome_padrao(v.get("nome"))
        return "" if n is ERRO else n

    def achar_pessoa(texto):
        k = chave_texto(texto)
        if not k:
            return None
        cands = [p for p in pessoas if k in (chave_texto(p["valores"].get("nome")), chave_texto(_np(p)))]
        if len(cands) == 1:
            return cands[0]["id"]
        if len(cands) > 1:
            return None
        toks = k.split()
        if len(toks) >= 2:
            por_ponta = [p for p in pessoas
                         if (lambda t: len(t) >= 2 and t[0] == toks[0] and t[-1] == toks[-1])(
                             chave_texto(p["valores"].get("nome")).split())]
            if len(por_ponta) == 1:
                return por_ponta[0]["id"]
        return None

    ordem = len(pessoas)
    criados = Counter()

    def nova_pessoa(nome, vinculo, origem):
        nonlocal ordem
        ordem += 1
        id_ = id_pessoa(None, nome, vinculo)
        pessoas.append({"id": id_, "ordem": ordem, "valores": {"nome": nome, "vinculo": vinculo}})
        criados[origem] += 1
        return id_

    for coluna, vinculo in (("Supervisor", "Supervisor"), ("Gestor de Contrato", "Gestor de contrato")):
        for nome in _juntar(aux.coluna(coluna)):
            if achar_pessoa(nome) is None:
                nova_pessoa(nome, vinculo, f"lista {coluna} do Auxiliar")
    for i, t in sup_digitado.items():
        pessoas[i - 1]["valores"]["supervisor"] = tipos.marcador(t) or achar_pessoa(t) or Legado(t)

    # ── clientes: a coluna CLIENTE vira cadastro com ID (separa usinas de mesmo nome e clientes diferentes) ──
    cl_atual = {chave_texto(r.valores.get("nome")): r for r in atuais["clientes"]}
    gera_cl = _proximo([r.id for r in atuais["clientes"]])
    clientes, cl_id = [], {}
    for nome in _juntar(ops.coluna("CLIENTE")):
        k = chave_texto(nome)
        if k in cl_id:
            continue
        r = cl_atual.get(k)
        cl_id[k] = r.id if r else next(gera_cl)
        clientes.append({"id": cl_id[k], "ordem": len(clientes) + 1, "valores": {
            "nome": nome, "observacao": r.valores.get("observacao") if r else None}})

    # ── usinas: id do Nexus é número; o IDUsina da planilha fica no "ID do BD" e é por ele que se casa ──
    u_atual = {str(r.valores.get("id_bd")): r for r in atuais["usinas"] if r.valores.get("id_bd")}
    gera_u = _proximo([r.id for r in atuais["usinas"]])
    usinas = []
    for i, ln in enumerate(ops.linhas, start=1):
        id_bd = _texto(cel(ops, ln, "IDUsina").valor)
        if not id_bd:
            avisos.append(f"Linha {i + 1} da aba Operações sem IDUsina: não importada.")
            continue
        id_ = u_atual[id_bd].id if id_bd in u_atual else next(gera_u)
        valores = {}
        for c in USINAS.campos:
            if not c.coluna_bd or c.modo == "calculado":
                continue
            x = cel(ops, ln, c.coluna_bd)
            if c.id == "receita_mensal" and x.formula:
                m = _RE_PRECO.match(x.formula)
                if m:
                    valores["preco_mwp"] = float(m.group(1).replace(",", "."))
                    valores["receita_mensal"] = None
                    continue
            if c.modo == "automatico" and x.formula and c.id != "receita_mensal":
                valores[c.id] = None
                continue
            if c.tipo == "ref":
                t = _texto(x.valor)
                if not t or t == "0":
                    valores[c.id] = None
                elif tipos.marcador(t):
                    valores[c.id] = tipos.marcador(t)       # "N/A" digitado: a usina não tem, não é pessoa
                elif c.id == "equipe":
                    valores[c.id] = equipe_de(t)
                elif c.id == "cliente":
                    valores[c.id] = cl_id.get(chave_texto(t)) or Legado(t)
                else:
                    achada = achar_pessoa(t)
                    if achada is None and c.id in ("responsavel_om", "gestor_contrato"):
                        vinc = "Supervisor" if c.id == "responsavel_om" else "Gestor de contrato"
                        achada = nova_pessoa(t, vinc, f"coluna {c.coluna_bd}")
                    valores[c.id] = achada or Legado(t)
                continue
            if x.formula:
                formula_valor[c.coluna_bd] += 1
            valores[c.id] = tipos.do_excel(c.tipo, x.valor)
        valores.setdefault("preco_mwp", None)
        usinas.append({"id": id_, "ordem": i, "valores": valores})

    # ── listas ──
    val = ops.validacoes
    listas = {
        "status_usina": _juntar(val.get("STATUS"), aux.coluna("CONSTRUÇÃO"), ops.coluna("STATUS")),
        "tipologia": _juntar(val.get("TIPOLOGIA"), ops.coluna("TIPOLOGIA")),
        "faixa_potencia": _juntar(aux.coluna("Faixa Potência"), ops.coluna("FAIXA POTÊNCIA")),
        "cluster": _juntar(ops.coluna("CLUSTER")),
        "regiao": _juntar(ops.coluna("REGIÃO")),
        "uf": _juntar(UFS, ops.coluna("UF")),
        "cargos": _juntar(pes.coluna("Cargo"), [calculos.CARGO_TECNICO, calculos.CARGO_ELETRICISTA,
                                                calculos.CARGO_MANTENEDOR]),
        "status_contratacao": _juntar(["Ativo", "Desligado"], pes.coluna("Status de Contratação")),
        "vinculo": list(VINCULOS),
        "estrutura_zeladoria": _juntar(aux.coluna("Estrutura Equipe Zeladoria")),
        "estrutura_mpa_mps": _juntar(aux.coluna("Estrutura Equipe MPA MPS")),
        "filtro_locais": _juntar(aux.coluna("Filtro Locais")),
        "status_operacional": _juntar(par.coluna("Status Operacional")),
        "modalidade": _juntar(par.coluna("Modalidade")),
        "status_contrato": _juntar(par.coluna("Status Contrato")),
    }
    if srv:
        # Valor incluído à mão no Nexus fica (a lista é domínio, não registro).
        for lista, valores in srv.listas().items():
            if lista in listas:
                listas[lista] = _juntar(listas[lista], valores)

    for coluna, n in sorted(formula_valor.items()):
        avisos.append(f"{coluna}: {n} célula(s) eram fórmula de conta ou de outra coluna e entraram como o valor "
                      "que o Excel mostrava.")
    for origem, n in sorted(criados.items()):
        avisos.append(f"{n} pessoa(s) criada(s) a partir da {origem} (só o nome; complete a ficha).")
    for f in sorted(faltando):
        avisos.append(f"Coluna não encontrada: {f}.")

    carga = Carga(entidades={"clientes": clientes, "equipes": equipes, "pessoas": pessoas, "usinas": usinas},
                  listas=listas,
                  origem={"arquivo": arquivo, "sha256": hashlib.sha256(dados).hexdigest()[:16] if dados else ""})
    comparacao = _comparar_calculados(ops, pes, carga)
    resumo = _resumo(carga, atuais)
    origem = {"arquivo": arquivo, "abas": {n: len(a.linhas) for n, a in abas.items()},
              "sha256": carga.origem["sha256"],
              "quando": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}
    return Proposta(carga=carga, resumo=resumo, avisos=avisos, comparacao=comparacao, origem=origem,
                    pessoas_por_id={p["id"]: p for p in pessoas})


def _repr(ent, valores: dict) -> dict:
    return {c.id: tipos.para_api(c.tipo, valores.get(c.id)) for c in ent.campos if c.modo != "calculado"}


def _resumo(carga: Carga, atuais: dict) -> dict:
    out = {}
    for eid, regs in carga.entidades.items():
        ent = ENTIDADES[eid]
        vivos = {r.id: r for r in atuais[eid] if not r.excluido}
        novos, alterados, iguais = [], [], 0
        for r in regs:
            a = vivos.get(r["id"])
            if a is None:
                novos.append(r["id"])
            elif _repr(ent, a.valores) != _repr(ent, r["valores"]) or a.ordem != r["ordem"]:
                alterados.append(r["id"])
            else:
                iguais += 1
        vindos = {r["id"] for r in regs}
        out[eid] = {
            "novos": len(novos), "alterados": len(alterados), "iguais": iguais,
            "sairiam": len([i for i in vivos if i not in vindos]),
            "lista_novos": novos, "lista_alterados": alterados,
            "lista_sairiam": sorted(i for i in vivos if i not in vindos),
            "editados_no_nexus": sorted(i for i, r in vivos.items()
                                        if r.alterado_por and r.alterado_por != QUEM_IMPORTACAO),
        }
    return out


# ── comparação Excel × Nexus ───────────────────────────────────────────────

def _vazio_cmp(v) -> bool:
    if isinstance(v, (list, tuple)):
        return all(_vazio_cmp(x) for x in v)
    return (v is None or v is ERRO or (isinstance(v, str) and (v.strip() == "" or v.strip().startswith("#")))
            or (isinstance(v, (int, float)) and not isinstance(v, bool) and v == 0))


def _sem_espaco(v) -> str:
    return re.sub(r"\s+", "", str(v))


def comparar(tipo: str, excel, nexus) -> bool:
    """O valor que o Excel mostrava é o mesmo que o Nexus calcula? 0, vazio e erro do Excel são "sem valor"
    (pessoa sem cluster: o PROCX casa com uma linha vazia e mostra 0). Espaço sobrando não é diferença: o Nexus
    apara o texto na importação, e a planilha tinha cidade com espaço no fim."""
    if _vazio_cmp(excel) or _vazio_cmp(nexus):
        return _vazio_cmp(excel) and _vazio_cmp(nexus)
    if tipo == "telefone":
        dig = lambda v: re.sub(r"\D", "", str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
        return dig(excel) == dig(nexus)
    if tipo in ("moeda", "numero"):
        try:
            return abs(float(excel) - float(nexus)) < 0.01
        except (TypeError, ValueError):
            return False
    if tipo == "ref":
        return chave_texto(excel) in {chave_texto(x) for x in (nexus if isinstance(nexus, (list, tuple)) else [nexus])}
    return _sem_espaco(excel) == _sem_espaco(nexus)


def _comparar_calculados(ops: Aba, pes: Aba, carga: Carga) -> list:
    usinas = carga.entidades["usinas"]
    pessoas = carga.entidades["pessoas"]
    ctx = Contexto(usinas=usinas, pessoas=pessoas, equipes=carga.entidades["equipes"])
    por_id = {p["id"]: p for p in pessoas}
    calc_p = {p["id"]: ctx.calcular_pessoa(p) for p in pessoas}

    def nomes(pid):
        p = por_id.get(pid) if isinstance(pid, str) else None
        if not p:
            return [pid.texto] if isinstance(pid, Legado) else []
        n = calc_p[p["id"]]["nome_padrao"]
        return [x for x in (n if n is not ERRO else None, p["valores"].get("nome")) if x]

    def desligada(nome_excel):
        k = chave_texto(nome_excel)
        return any(chave_texto(p["valores"].get("status")) == "desligado"
                   and k in (chave_texto(p["valores"].get("nome")), chave_texto(nomes(p["id"])[0] if nomes(p["id"]) else ""))
                   for p in pessoas)

    saida = []

    def bloco(entidade, campo, tipo, aba, linhas_regs, valor_nexus, motivo=None):
        c = ENTIDADES[entidade].campo(campo)
        total, iguais, dif = 0, 0, []
        for ln, reg in linhas_regs:
            x = aba.celula(ln, c.coluna_bd)
            if not x.formula:
                continue
            total += 1
            nv = valor_nexus(reg)
            if comparar(tipo, x.valor, nv):
                iguais += 1
            else:
                mostra = nv[0] if isinstance(nv, list) and nv else ("" if nv is ERRO or nv is None or nv == [] else nv)
                dif.append({"id": reg["id"], "id_bd": reg["valores"].get("id_bd"),
                            "titulo": reg["valores"].get("nome") or reg["id"],
                            "excel": x.valor, "nexus": mostra, "motivo": motivo(x.valor, reg) if motivo else ""})
        if total:
            saida.append({"entidade": entidade, "campo": campo, "rotulo": c.rotulo, "total": total,
                          "iguais": iguais, "diferentes": dif})

    por_bd = {u["valores"].get("id_bd"): u for u in usinas}
    linhas_u = [(ln, por_bd[_texto(ops.celula(ln, "IDUsina").valor)]) for ln in ops.linhas
                if _texto(ops.celula(ln, "IDUsina").valor) in por_bd]
    calc_u = {u["id"]: ctx.calcular_usina(u) for u in usinas}

    def motivo_pessoa(excel, reg):
        return "O Excel pegava uma pessoa desligada (a primeira da planilha)." if desligada(excel) else ""

    for campo in ("tecnico_om", "eletricista_om", "mantenedor_om"):
        bloco("usinas", campo, "ref", ops, linhas_u, lambda r, c=campo: nomes(calc_u[r["id"]][c]), motivo_pessoa)
    for campo in ("contato_tecnico", "contato_eletricista", "contato_mantenedor"):
        bloco("usinas", campo, "telefone", ops, linhas_u, lambda r, c=campo: calc_u[r["id"]][c])
    bloco("usinas", "base_equipe", "texto", ops, linhas_u, lambda r: calc_u[r["id"]]["base_equipe"],
          lambda e, r: "O Excel pegava a cidade da primeira pessoa da equipe, mesmo desligada.")
    bloco("usinas", "receita_mensal", "moeda", ops, linhas_u, lambda r: calc_u[r["id"]]["receita_mensal"])
    bloco("usinas", "receita_contratual", "moeda", ops, linhas_u, lambda r: calc_u[r["id"]]["receita_contratual"])
    bloco("usinas", "localizacao", "texto", ops, linhas_u, lambda r: calc_u[r["id"]]["localizacao"])

    linhas_p = list(zip(pes.linhas, pessoas[:len(pes.linhas)]))
    bloco("pessoas", "nome_padrao", "texto", pes, linhas_p, lambda r: calc_p[r["id"]]["nome_padrao"])
    bloco("pessoas", "supervisor", "ref", pes, linhas_p, lambda r: nomes(calc_p[r["id"]]["supervisor"]))
    for campo in ("cidade_uf_endereco", "cidade_uf_fallback", "cidade_estado_base"):
        bloco("pessoas", campo, "texto", pes, linhas_p, lambda r, c=campo: calc_p[r["id"]][c])
    return saida
