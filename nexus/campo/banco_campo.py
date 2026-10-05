"""Os fechamentos que o Nexus calcula, na API do PG (db_performace), workbook `campo_nexus` (Levi, 04/10/2026: "não
aceito esse JSON como banco de dados pois temos a API do PG"; "pode gravar na API do PG e ligar as telas").

Uma linha por TAREFA que entrou na fila de verificação do Fracttal: a nota da régua do App calculada com o que o
Fracttal guarda (nota_fracttal.py) e os campos do Fracttal que as telas usam, os mesmos que o `_enriquecer_qlog` do App
copia toda madrugada para o registro dele (tipo, criticidade, status, área, ativo, datas, durações, avaliação).

O nome do técnico vai CIFRADO (NEXUS_CHAVE_CADASTRO, o cofre do cadastro): a leitura da API não pede credencial. Sem a
chave, vai só o id do Fracttal. E-mail não vai: quem casa o técnico com o cadastro é a tela, na hora.

Gravação: o caminho provado do cadastro (cadastro/banco.py): xlsx com o cabeçalho na linha 1, célula vazia = None,
POST /api/workbooks/campo_nexus/sync-xlsx?replace=true. Só o Nexus grava neste workbook, então trocar tudo de uma vez
não apaga nada de ninguém. Depois de gravar, lê de volta e confere linha a linha (regra: conferir antes e depois).
"""
import io
from datetime import datetime, timedelta, timezone

WORKBOOK = "campo_nexus"
NOME = "Campo · App (Nexus): fechamentos calculados pelo Fracttal"
ABA = "fechamentos"
ABA_ATUALIZACAO = "atualizacao"
# Levi, 05/10/2026: "esses dados de ronda têm sim que ir para a API, guardados como informação de ronda, salvando o
# número da OS e todas essas informações". Uma linha por OS de ronda, com o que o App escreveu nela no Fracttal.
ABA_RONDAS = "rondas"
BASE_API = "https://app.gridco.com.br/db_performace"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BRT = timezone(timedelta(hours=-3))

INTEIRAS = ("id_tarefa", "id_os", "tecnico_id", "status_os", "status_tarefa", "dur_prev_min", "dur_real_min", "rating",
            "nota", "nota_placar", "pontos_placar", "n_fotos", "n_desc", "ronda_duracao_min", "usina_id", "pessoa_id")
SIM_NAO = ("pelo_app", "ronda", "gps_fotos", "sub_ok", "todas_sub", "obs_ok", "devolvida", "completa", "ronda_lida")
COLUNAS = ("id_tarefa", "os", "id_os", "tarefa", "tecnico_id", "tecnico_cifrado", "usina_id", "pessoa_id", "equipe",
           "regiao", "ativo", "codigo",
           "tipo", "crit", "status_os", "status_tarefa", "inicio", "fim", "verificacao", "aprovacao", "dur_prev_min",
           "dur_real_min", "rating", "pelo_app", "ronda", "ronda_dia", "ronda_usina", "ronda_duracao_min", "ronda_hora_ini",
           "ronda_hora_fim", "ronda_respostas", "ronda_pendencias", "ronda_lida", "nota", "nota_itens", "nota_placar",
           "pontos_placar", "n_fotos", "n_desc", "gps_fotos", "sub_ok", "todas_sub", "obs_ok", "devolvida", "lido_em",
           "regua")
# a aba de rondas: (coluna da aba, coluna da linha da tarefa)
COLUNAS_RONDAS = (("os", "os"), ("id_os", "id_os"), ("id_tarefa", "id_tarefa"), ("dia", "ronda_dia"),
                  ("usina", "ronda_usina"), ("usina_id", "usina_id"), ("equipe", "equipe"), ("regiao", "regiao"),
                  ("pessoa_id", "pessoa_id"), ("tecnico_id", "tecnico_id"),
                  ("tecnico_cifrado", "tecnico_cifrado"), ("qualidade", "nota"), ("duracao_real_min", "ronda_duracao_min"),
                  ("hora_inicio", "ronda_hora_ini"), ("hora_fim", "ronda_hora_fim"), ("respostas", "ronda_respostas"),
                  ("pendencias_e_achados", "ronda_pendencias"), ("n_fotos", "n_fotos"), ("gps_fotos", "gps_fotos"),
                  ("status_os", "status_os"), ("fechada_em", "fim"), ("aprovada_em", "aprovacao"), ("lido_em", "lido_em"))
CAB_ATUALIZACAO = ("publicado_em", "tarefas", "pelo_app", "com_nota", "completa", "pendentes", "fila", "aprovadas_janela",
                   "janela_dias", "sem_usina_id", "sem_pessoa_id", "fonte", "regua", "como_ler")


class BancoErro(RuntimeError):
    pass


def _int(v):
    if v is None or v == "" or isinstance(v, bool):
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _sim_nao(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in ("sim", "true", "1")


def normalizar(linha: dict) -> dict:
    """A linha no formato do Nexus: inteiros como int, sim/não como bool, vazio como None."""
    out = {}
    for c in COLUNAS:
        v = linha.get(c)
        if c in INTEIRAS:
            v = _int(v)
        elif c in SIM_NAO:
            v = _sim_nao(v)
        elif v is not None:
            v = str(v)
            v = v if v.strip() else None
        out[c] = v
    return out


def _celula(c, v):
    if v is None or v == "":
        return None                     # o sync recusa texto vazio (cadastro/tipos.py, 03/09/2026)
    if c in SIM_NAO:
        return "sim" if v else "não"
    return v


def xlsx_bytes(linhas: list[dict], resumo: dict) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(ABA)
    ws.append(list(COLUNAS))
    for l in sorted(linhas, key=lambda x: x["id_tarefa"]):
        ws.append([_celula(c, l.get(c)) for c in COLUNAS])
    ws.freeze_panes = "A2"
    ws = wb.create_sheet(ABA_RONDAS)
    ws.append([a for a, _ in COLUNAS_RONDAS])
    for l in sorted((l for l in linhas if l.get("ronda")), key=lambda x: x["id_tarefa"]):
        ws.append([_celula(c, l.get(c)) for _, c in COLUNAS_RONDAS])
    ws.freeze_panes = "A2"
    ws = wb.create_sheet(ABA_ATUALIZACAO)
    ws.append(list(CAB_ATUALIZACAO))
    ws.append([_celula(c, resumo.get(c)) for c in CAB_ATUALIZACAO])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _abas(base, s) -> dict:
    r = s.get(f"{base}/api/sheets", timeout=60)
    r.raise_for_status()
    return {x["sheet_name"]: x["id"] for x in r.json() if x.get("workbook_key") == WORKBOOK}


def _linhas_da_aba(base, s, sid) -> list[dict]:
    out, offset = [], 0
    while True:
        r = s.get(f"{base}/api/sheets/{sid}/rows", params={"limit": 1000, "offset": offset}, timeout=60)
        r.raise_for_status()
        rows = r.json()
        rows = rows.get("rows", rows) if isinstance(rows, dict) else rows
        out += [dict(zip(x.get("headers") or [], x.get("values") or [])) for x in rows]
        if len(rows) < 1000:
            return out
        offset += 1000


def ler_tudo(base: str = BASE_API, sessao=None) -> tuple[dict, dict]:
    """({id_tarefa: linha}, a linha da aba de atualização). Workbook ainda inexistente = vazio."""
    import requests
    s = sessao or requests
    base = base.rstrip("/")
    abas = _abas(base, s)
    out, atualizacao = {}, {}
    if ABA in abas:
        for bruta in _linhas_da_aba(base, s, abas[ABA]):
            l = normalizar(bruta)
            if l["id_tarefa"] is not None:
                out[l["id_tarefa"]] = l
    if ABA_ATUALIZACAO in abas:
        atualizacao = next(iter(_linhas_da_aba(base, s, abas[ABA_ATUALIZACAO])), {})
    return out, atualizacao


def ler_rondas(base: str = BASE_API, sessao=None) -> list[dict]:
    """As linhas da aba de rondas, como a API devolve."""
    import requests
    s = sessao or requests
    base = base.rstrip("/")
    abas = _abas(base, s)
    return _linhas_da_aba(base, s, abas[ABA_RONDAS]) if ABA_RONDAS in abas else []


def ler(base: str = BASE_API, sessao=None) -> dict:
    """{id_tarefa: linha} do que está no banco."""
    return ler_tudo(base, sessao)[0]


def _canon(l: dict) -> tuple:
    return tuple((c, l.get(c)) for c in COLUNAS)


def conferir(gravadas: list[dict], lidas: dict) -> list[str]:
    """O que voltou do banco é o que foi gravado? Devolve as diferenças (vazio = confere)."""
    dif = []
    esperadas = {l["id_tarefa"]: normalizar(l) for l in gravadas}
    if set(esperadas) != set(lidas):
        faltam, sobram = set(esperadas) - set(lidas), set(lidas) - set(esperadas)
        dif.append(f"tarefas: {len(faltam)} não voltaram, {len(sobram)} a mais no banco")
    for k, l in esperadas.items():
        if k in lidas and _canon(l) != _canon(lidas[k]):
            cols = [c for c in COLUNAS if l.get(c) != lidas[k].get(c)]
            dif.append(f"tarefa {k}: {', '.join(cols)}")
    return dif


def gravar(linhas: list[dict], resumo: dict, *, base: str, token: str, sessao=None) -> dict:
    """Cria o workbook só se faltar (a API não apaga workbook), sincroniza e confere lendo de volta."""
    import requests
    s = sessao or requests
    base = base.rstrip("/")
    h = {"Authorization": f"Bearer {token}"}
    r = s.get(f"{base}/api/workbooks", headers=h, timeout=60)
    r.raise_for_status()
    if WORKBOOK not in {w.get("key") for w in r.json()}:
        s.post(f"{base}/api/workbooks", headers=h, json={"key": WORKBOOK, "display_name": NOME},
               timeout=60).raise_for_status()
    r = s.post(f"{base}/api/workbooks/{WORKBOOK}/sync-xlsx", headers=h, params={"replace": "true"},
               files={"file": (f"{WORKBOOK}.xlsx", xlsx_bytes(linhas, resumo), MIME_XLSX)}, timeout=300)
    r.raise_for_status()
    dif = conferir(linhas, ler(base, s))
    rondas = {str(l["os"]) for l in linhas if l.get("ronda")}
    voltaram = [str(r.get("os")) for r in ler_rondas(base, s)]
    if sorted(voltaram) != sorted(str(l["os"]) for l in linhas if l.get("ronda")) or set(voltaram) != rondas:
        dif.append(f"rondas: {len(rondas)} OS gravadas, {len(voltaram)} linhas no banco")
    if dif:
        raise BancoErro("o banco não devolveu o que foi gravado: " + "; ".join(dif[:5]))
    return {"gravadas": len(linhas), "conferidas": len(linhas)}


def resumo(linhas: list[dict], regua: str, *, completa=False, pendentes=0, fila=0, aprovadas=0, janela_dias=90) -> dict:
    """A aba de atualização. `completa` = a fila de verificação e as aprovadas da janela estão TODAS calculadas: é o
    que as telas pedem para sair do painel do App."""
    return {"publicado_em": datetime.now(_BRT).isoformat(timespec="seconds"), "tarefas": len(linhas),
            "pelo_app": sum(1 for l in linhas if l.get("pelo_app")),
            "com_nota": sum(1 for l in linhas if l.get("nota") is not None), "completa": bool(completa),
            "pendentes": int(pendentes), "fila": int(fila), "aprovadas_janela": int(aprovadas),
            "janela_dias": int(janela_dias),
            # governança: o que não ligou ao cadastro fica à vista (usina sem de-para, técnico sem ficha ou homônimo)
            "sem_usina_id": sum(1 for l in linhas if l.get("usina_id") is None),
            "sem_pessoa_id": sum(1 for l in linhas if l.get("pessoa_id") is None),
            "fonte": "Fracttal (REST, só leitura): work_orders, work_orders_subtasks, work_orders_attachments",
            "regua": regua,
            "como_ler": "uma linha por tarefa que entrou na fila de verificação; nota = a do painel do App "
                        "(_qualidade_os), nota_placar = régua de 01/09/2026 (_qualidade_v2); pelo_app = fechada pelo "
                        "App (marca das fotos); tecnico_cifrado só abre no Nexus"}
