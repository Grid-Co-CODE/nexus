"""Etapa 2: gerar a semana no Nexus, em SOMBRA.

O motor é o do PCM (programacao_v7.py e os dois leitores do Fracttal), copiado sem mudança em motor/. O Nexus:
1. confere o que o motor precisa: credencial, insumos e o próprio motor;
2. monta uma pasta por rodada, FORA do OneDrive, e escreve nela os insumos que moram no Nexus (insumos.py:
   prioridades, confiabilidade, feriados e observações), no formato que o motor lê, e a AUXILIAR, que sai
   do cadastro do Nexus (auxiliar.py). Da pasta do PCM só vêm as durações aprendidas. Na sombra, o Nexus importa os arquivos da pasta do PCM antes, para as duas
   gerações partirem do mesmo ponto;
3. escreve o histórico das programações a partir do BANCO (`historico_banco.py`, desde 08/10/2026), só com as semanas
   de antes da gerada: na 2ª geração da mesma semana o motor marcava tudo como reprogramado (415 tarefas na semana 40),
   e assim tanto faz rodar antes ou depois do PCM. Banco fora do ar: a reserva guardada no Nexus, voltada para antes
   da semana (`preparar_historico`);
4. roda o motor num SUBPROCESSO, com ambiente próprio. O FRACTTAL_BASE_URL do motor tem /api/ e o do OS Creator, que
   roda dentro do Nexus, não; no mesmo processo um quebraria o outro. As variáveis NEXUS_* não vão para o motor;
5. compara a planilha que saiu com a oficial da mesma semana, se ela já existir.
Nada é publicado: a semana do campo continua sendo a oficial do PCM.

`gerar_com_foto` (ferramentas/gerar_semana_foto.py) faz o mesmo SEM ler o Fracttal: o motor recebe a foto que ele mesmo
gravou na pasta do PCM e roda sem credencial. É a prova de mudança no motor ou nos insumos: ler o Fracttal horas depois
muda a semana (S41: 157 tarefas a mais e 280 horários) e não prova nada.

Credencial (decisão do Levi, 30/09: "pode usar do OS Creator, mesma lógica porém sem a parte de login"): o
client_credentials do Fracttal que o OS Creator lê do .env da pasta dele. A leitura pela API não precisa do login de
pessoa (o RPC); o motor do PCM usa o mesmo método.
"""
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from dotenv import dotenv_values

from . import auxiliar as A
from . import comparar
from . import fonte
from . import historico_banco as HB
from . import insumos as I

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
MOTOR = AQUI / "motor" / "programacao_v7.py"
ENV_OS_CREATOR = RAIZ / "nexus" / "torres" / "oscreator" / "os_creator" / ".env"
# Fora do OneDrive e fora do AppData, como o ensaio do cadastro (o Claude desktop virtualiza o AppData).
TRABALHO_PADRAO = Path(r"C:\GridcoAuto\nexus\pcm") if os.name == "nt" else RAIZ / "dados" / "pcm"
# A pasta do PCM sincronizada nesta máquina: é de lá que o PCM gera. Na sombra, os insumos vêm dela.
ORIGEM_PADRAO = (Path.home() / "OneDrive - GRID CO" / "Área de Trabalho" / "Grid Co_ - 4. O&M" / "11.Pré-Operação"
                 / "6. PCM" / "09. Programação Semanal")
# O que ainda vem da pasta do PCM. Prioridades, Confiabilidade, Histórico, Feriados e Observações moram no Nexus
# desde 30/09/2026 (insumos.py) e são escritos na rodada a partir de lá. A AUXILIAR sai do cadastro desde 02/10.
INSUMOS = (  # (arquivo, obrigatório)
    ("duracoes_aprendidas.json", False),      # modo sombra no motor: só preenche uma coluna
)
URL_API = "https://app.fracttal.com/api/"
# 1 pedido por segundo. Fora do GitHub o motor anda a 0,15 s, rápido demais para os 200/min que robôs e App de Campo
# dividem na EMPRESA inteira desde 24/09.
INTERVALO_S = "1.0"
TEMPO_MAX_S = 60 * 60
SEMANA_RE = re.compile(r"^\d{4}-W\d{2}$")
RODADA_RE = re.compile(r"^\d{8}-\d{6}-\d{4}-W\d{2}$")
_BRT = timezone(timedelta(hours=-3))
# A foto do Fracttal: os caches que o motor (fonte_bd_api) grava na pasta de trabalho ao ler o Fracttal. Com o
# .cache_semanal_api.pkl "fresco" o motor nem abre conexão: devolve o DataFrame da foto, e as coordenadas saem do
# _usinas_coordenadas_cache.json. Os outros vão junto para a rodada ser a mesma pasta que o PCM tinha.
FOTO = (".cache_semanal_api.pkl", ".cache_hist_api.pkl", ".cache_bd_api.pkl", "_ativos_classificacao_cache.json",
        "_usinas_coordenadas_cache.json")
FOTO_OBRIGATORIA = ".cache_semanal_api.pkl"
# "Fresco" para sempre (~100 anos): o TTL do motor é pela idade do arquivo, e a foto tem horas ou dias
TTL_FOTO = {"PROG_API_TTL_MIN": "52560000", "PROG_HIST_TTL_H": "876000", "GESTAO_ATIVOS_TTL_H": "876000"}
# Sem credencial e com o Fracttal num endereço que recusa conexão: se o motor tentar ler, ele FALHA, nunca lê. As vazias
# também barram um .env perdido: o carregar_env do motor usa setdefault, e variável que já existe (mesmo vazia) vence.
SEM_FRACTTAL = {"FRACTTAL_CLIENT_ID": "", "FRACTTAL_CLIENT_SECRET": "", "FRACTTAL_JWT_TOKEN": "",
                "FRACTTAL_BASE_URL": "http://127.0.0.1:9/sem-fracttal/"}

_trava = threading.Lock()
_rodando: dict[str, tuple[str, threading.Thread]] = {}


class Ocupado(RuntimeError):
    pass


class NaoPronto(RuntimeError):
    pass


@dataclass
class Credencial:
    ok: bool
    onde: str
    client_id: str = field(default="", repr=False)
    client_secret: str = field(default="", repr=False)


def _cfg(config, chave, padrao=None):
    return config.get(chave) or os.environ.get(chave) or padrao


def pasta_trabalho(config) -> Path:
    return Path(_cfg(config, "NEXUS_PCM_TRABALHO", TRABALHO_PADRAO))


def pasta_origem(config) -> Path | None:
    if config.get("TESTING") and not config.get("NEXUS_PCM_ORIGEM"):
        return None
    return Path(_cfg(config, "NEXUS_PCM_ORIGEM", ORIGEM_PADRAO))


def caminho_motor(config) -> Path:
    return Path(_cfg(config, "NEXUS_PCM_MOTOR", MOTOR))


def credencial(config) -> Credencial:
    """Onde está o client_credentials do Fracttal, sem nunca expor o valor. Ordem: ambiente do processo (servidor),
    .env do Nexus, .env do OS Creator dentro do Nexus, e um .env apontado por NEXUS_FRACTTAL_ENV."""
    if config.get("TESTING"):
        fontes = [("teste", config.get("NEXUS_PCM_FRACTTAL_TESTE") or {})]
    else:
        fontes = [("ambiente", os.environ), ("o .env do Nexus", dotenv_values(RAIZ / ".env"))]
        if ENV_OS_CREATOR.exists():
            fontes.append(("o .env do OS Creator", dotenv_values(ENV_OS_CREATOR)))
        extra = _cfg(config, "NEXUS_FRACTTAL_ENV")
        if extra and Path(extra).exists():
            fontes.append((f"o .env de {Path(extra).parent.name}", dotenv_values(extra)))
    for onde, vals in fontes:
        cid = (vals.get("FRACTTAL_CLIENT_ID") or "").strip()
        seg = (vals.get("FRACTTAL_CLIENT_SECRET") or "").strip()
        if cid and seg:
            return Credencial(True, onde, cid, seg)
    return Credencial(False, "")


def horario_de_campo(agora: datetime | None = None) -> bool:
    """Seg a sex, das 06h às 18h: o motor lê o Fracttal inteiro e divide a cota com o App de Campo."""
    a = (agora or datetime.now(_BRT)).astimezone(_BRT)
    return a.weekday() < 5 and 6 <= a.hour < 18


def semana_padrao(hoje: date | None = None) -> str:
    """A da próxima segunda-feira, como o motor faz sem --semana."""
    h = hoje or datetime.now(_BRT).date()
    seg = h + timedelta(days=(7 - h.weekday()) or 7)
    iso = seg.isocalendar()
    return f"{iso[0]}-W{iso[1]:02d}"


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()[:12]


def _quando(ts: float) -> str:
    return datetime.fromtimestamp(ts, _BRT).strftime("%d/%m %H:%M")


def conferir(config, semana: str | None = None) -> dict:
    itens = []
    m = caminho_motor(config)
    itens.append({"nome": "Motor do PCM", "ok": m.exists(), "obrigatorio": True,
                  "detalhe": "cópia idêntica do programacao_v7.py em nexus/pcm/motor" if m.exists() else "motor não encontrado"})
    c = credencial(config)
    itens.append({"nome": "Credencial do Fracttal", "ok": c.ok, "obrigatorio": True,
                  "detalhe": f"encontrada em {c.onde}" if c.ok else
                  "falta: copie o .env do OS Creator para nexus/torres/oscreator/os_creator/.env"})
    origem = pasta_origem(config)
    tem_origem = bool(origem and origem.is_dir())
    # Opcional desde 02/10: a AUXILIAR sai do cadastro e os insumos moram no Nexus. A pasta só traz a oficial (para
    # comparar) e as durações (sombra). Obrigatória, ela travava a geração no servidor, onde o OneDrive não existe.
    itens.append({"nome": "Pasta do PCM", "ok": True, "obrigatorio": False,
                  "detalhe": "a mesma de onde o PCM gera: serve para comparar com a oficial" if tem_origem else
                  "não está nesta máquina: gera mesmo assim, sem comparar com a oficial"})
    for nome, obrig in INSUMOS:
        p = origem / nome if origem else None
        existe = bool(p and p.exists())
        itens.append({"nome": nome, "ok": existe or not obrig, "obrigatorio": obrig,
                      "detalhe": ("na pasta do PCM, atualizado em " + _quando(p.stat().st_mtime)) if existe else
                      ("falta" if obrig else "opcional, ausente")})
    itens.append(A.estado(config))
    itens += I.estado(pasta_trabalho(config), origem if tem_origem else None, semana or semana_padrao())
    return {"pronto": all(i["ok"] for i in itens), "itens": itens}


def preparar_historico(src, dst, semana: str) -> None:
    """O histórico como era ANTES da semana: tira a semana de cada tarefa e refaz a contagem."""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    cab = list(next(it))
    ix = {c: i for i, c in enumerate(cab)}
    novas = []
    for r in it:
        if not r or not r[ix["task_key"]]:
            continue
        r = list(r)
        semanas = sorted(s for s in str(r[ix["weeks"]] or "").split(",") if s.strip() and s.strip() != semana)
        if not semanas:
            continue
        r[ix["weeks"]] = ",".join(semanas)
        r[ix["count"]] = len(semanas)
        r[ix["first_week"]] = semanas[0]
        r[ix["last_week"]] = semanas[-1]
        novas.append(r)
    wb.close()
    out = openpyxl.Workbook()
    wo = out.active
    wo.append(cab)
    for r in novas:
        wo.append(r)
    out.save(dst)


def _ambiente(cred: Credencial, pasta: Path, saida: Path) -> dict:
    fora = ("PCM_PROG_DIR", "PCM_OUTPUT", "PCM_WEEK_FORCE")
    env = {k: v for k, v in os.environ.items()
           if not k.startswith(("NEXUS_", "FRACTTAL_")) and k not in fora}
    env.update({"FRACTTAL_CLIENT_ID": cred.client_id, "FRACTTAL_CLIENT_SECRET": cred.client_secret,
                "FRACTTAL_BASE_URL": URL_API, "FRACTTAL_INTERVALO_S": INTERVALO_S,
                "PCM_PROG_DIR": str(pasta), "PCM_OUTPUT": str(saida),
                "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1"})
    # No Linux o TZ põe o motor no horário de Brasília. No Windows o C runtime não entende "America/Sao_Paulo" e o
    # relógio do motor saía 4 h adiantado (S41 de teste, 02/10/2026: rodou às 17:28 e o log dizia 21:28): depois das
    # 20h o "hoje" do motor viraria o dia seguinte e a idade das tarefas andaria um dia. Lá o relógio do PC já é BRT.
    if os.name != "nt":
        env["TZ"] = "America/Sao_Paulo"
    else:
        env.pop("TZ", None)
    return env


def _gravar(pasta: Path, st: dict) -> None:
    tmp = pasta / "status.json.tmp"
    tmp.write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(pasta / "status.json")


def _oficial(origem: Path | None, semana: str) -> Path | None:
    if not origem:
        return None
    p = origem / f"Programação Semana {int(semana.split('W')[1]):02d}.xlsx"
    return p if p.exists() else None


def leitor_do_repositorio(config):
    """Lê um arquivo da raiz do repositório do PCM (público: sem token). Nos testes, só o falso da config."""
    if config.get("TESTING"):
        return config.get("NEXUS_PCM_REPO_TESTE")
    repo, ramo = fonte.repositorio(config)

    def ler(nome: str) -> bytes | None:
        import requests
        from urllib.parse import quote
        r = requests.get(f"https://raw.githubusercontent.com/{repo}/{ramo}/{quote(nome)}", timeout=60)
        if r.status_code == 404:
            return None
        r.raise_for_status()
        return r.content
    return ler


def _historico(config, pasta: Path, semana: str, trabalho: Path, origem: Path | None, modo: str = "banco") -> dict:
    """Escreve o Historico_Programacoes.xlsx da rodada e devolve o carimbo. `modo` "banco" (o padrão desde 08/10/2026):
    o fato_programacao do banco, com o plano publicado da semana que o banco ainda não fechou; se o banco não responder,
    a reserva guardada no Nexus. `modo` "nexus": a reserva direto (o histórico de antes de 08/10, para comparar)."""
    destino = pasta / HB.NOME
    reserva = I.carregar(trabalho).get("historico")
    erro = None
    if modo == "banco":
        try:
            rel = HB.materializar(config, destino, semana, reserva=reserva, origem=origem,
                                  ler_repo=leitor_do_repositorio(config))
        except HB.HistoricoErro as ex:
            erro = str(ex)
        else:
            # o relatório por semana (de onde veio cada uma, quantas chaves) fica na rodada para conferir depois
            (pasta / "historico_banco.json").write_text(json.dumps(rel, ensure_ascii=False, indent=1), encoding="utf-8")
            return {"nome": I.NOMES["historico"], "sha": _sha(destino), "atualizado": (rel.get("banco_ate") or "")[:16],
                    "detalhe": rel["detalhe"], "aviso": bool(rel["avisos"])}
    if not reserva:
        raise NaoPronto("histórico: " + (f"o banco não respondeu ({erro}) e " if erro else "") + "não há reserva no Nexus")
    preparar_historico(pasta / "historico_do_nexus.xlsx", destino, semana)
    m = reserva.get("meta") or {}
    return {"nome": I.NOMES["historico"], "sha": _sha(destino), "atualizado": (m.get("importado_em") or "")[:16],
            "detalhe": "reserva do Nexus, voltada para antes da semana" + (f": o banco não respondeu ({erro})" if erro
                                                                           else " (pedida na rodada)"),
            "aviso": bool(erro)}


def _preparar(config, semana: str, pasta: Path, trabalho: Path, origem: Path | None, historico: str = "banco") -> list:
    """Escreve na pasta da rodada tudo o que o motor lê, menos o Fracttal. Devolve o carimbo de cada insumo (de onde
    veio e de quando é o dado)."""
    insumos = []
    for nome, _obrig in INSUMOS:
        p = origem / nome if origem else None
        if not (p and p.exists()):
            continue
        shutil.copy2(p, pasta / nome)
        insumos.append({"nome": nome, "sha": _sha(p), "atualizado": _quando(p.stat().st_mtime),
                        "detalhe": "da pasta do PCM"})
    # os outros saem do Nexus, no formato que o motor lê
    try:
        aux = A.escrever(A.servico(config), pasta / A.NOME)
        insumos.append({"nome": "Cadastro de usinas (AUXILIAR)", "sha": _sha(pasta / A.NOME),
                        "atualizado": datetime.now(_BRT).isoformat(timespec="minutes"),
                        "detalhe": f"do cadastro do Nexus: {aux['usinas']} usinas, "
                                   f"{aux['sem_responsavel']} sem responsável O&M"})
        insumos += I.materializar(trabalho, pasta, semana)
    except A.SemCadastro as ex:
        raise NaoPronto(f"cadastro de usinas: {ex}") from ex
    except I.InsumoErro as ex:
        raise NaoPronto(str(ex)) from ex
    insumos.append(_historico(config, pasta, semana, trabalho, origem, historico))
    return insumos


def iniciar(config, semana: str) -> dict:
    if not SEMANA_RE.match(semana or ""):
        raise ValueError("semana inválida")
    conf = conferir(config, semana)
    if not conf["pronto"]:
        raise NaoPronto("falta: " + ", ".join(i["nome"] for i in conf["itens"] if not i["ok"]))
    trabalho = pasta_trabalho(config)
    chave = str(trabalho.resolve())
    with _trava:
        ativo = _rodando.get(chave)
        if ativo and ativo[1].is_alive():
            raise Ocupado(ativo[0])
        rid = datetime.now(_BRT).strftime("%Y%m%d-%H%M%S") + "-" + semana
        pasta = trabalho / "geracoes" / rid
        (pasta / "saida").mkdir(parents=True, exist_ok=False)
        origem = pasta_origem(config)
        try:
            insumos = _preparar(config, semana, pasta, trabalho, origem)
        except NaoPronto:
            shutil.rmtree(pasta, ignore_errors=True)
            raise
        cred = credencial(config)
        st = {"id": rid, "semana": semana, "estado": "rodando", "inicio": datetime.now(_BRT).isoformat(timespec="seconds"),
              "credencial": cred.onde, "insumos": insumos, "motor": str(caminho_motor(config))}
        _gravar(pasta, st)
        saida = pasta / "saida" / "sombra.xlsx"
        t = threading.Thread(target=_executar, name=f"pcm-{rid}", daemon=True,
                             args=(pasta, st, _ambiente(cred, pasta, saida), caminho_motor(config), origem, saida))
        _rodando[chave] = (rid, t)
        t.start()
    return st


def _executar(pasta: Path, st: dict, env: dict, motor: Path, origem: Path, saida: Path) -> None:
    t0 = time.monotonic()
    log = pasta / "log_motor.txt"
    try:
        with open(log, "wb") as f:
            r = subprocess.run([sys.executable, str(motor), "--semana", st["semana"]], cwd=str(pasta), env=env,
                               stdout=f, stderr=subprocess.STDOUT, timeout=TEMPO_MAX_S)
        st["codigo"] = r.returncode
        if r.returncode == 0 and saida.exists():
            st["estado"] = "ok"
            st["resumo"] = comparar.resumo(saida)
            # o que olhar antes de publicar (blocos fora da semana, horas estouradas por equipe e dia e por OS...)
            try:
                st["conferencia"] = comparar.conferir(saida, st["semana"])
            except Exception as ex:      # noqa: BLE001 — a conferência não derruba a rodada; a publicação a exige
                st["conferencia_erro"] = f"{type(ex).__name__}: {str(ex)[:200]}"
            oficial = _oficial(origem, st["semana"])
            if oficial:
                shutil.copy2(oficial, pasta / "oficial.xlsx")
                st["oficial"] = {"arquivo": oficial.name, "atualizado": _quando(oficial.stat().st_mtime)}
                st["comparacao"] = comparar.comparar(pasta / "oficial.xlsx", saida)
        else:
            st["estado"] = "falhou"
            st["erro"] = _cauda(log, 12) or f"o motor terminou com código {r.returncode} sem planilha"
    except subprocess.TimeoutExpired:
        st["estado"] = "falhou"
        st["erro"] = f"o motor passou de {TEMPO_MAX_S // 60} min e foi parado"
    except Exception as ex:      # noqa: BLE001 — qualquer falha fica no status, nunca derruba o Nexus
        st["estado"] = "falhou"
        st["erro"] = f"{type(ex).__name__}: {str(ex)[:300]}"
    st["fim"] = datetime.now(_BRT).isoformat(timespec="seconds")
    st["duracao_s"] = int(time.monotonic() - t0)
    _gravar(pasta, st)


def _ambiente_foto(pasta: Path, saida: Path) -> dict:
    """O ambiente da rodada normal, sem credencial do Fracttal, com o endereço dele numa porta fechada e a foto
    "fresca" para sempre: o motor lê a foto e, se tentar o Fracttal, falha (nunca lê)."""
    env = _ambiente(Credencial(False, ""), pasta, saida)
    env.update(SEM_FRACTTAL)
    env.update(TTL_FOTO)
    return env


def gerar_com_foto(config, semana: str, foto: Path | None = None, historico: str = "banco", rotulo: str = "") -> dict:
    """Gera a semana com a FOTO do Fracttal e espera terminar. Os insumos são os da geração da tela (`_preparar`: os do
    Nexus, a AUXILIAR do cadastro, o histórico de antes da semana); o Fracttal é a foto que o motor do PCM gravou na
    pasta dele (`FOTO`), copiada para a rodada (a pasta do PCM só é lida).

    É o método da prova da S41 (02/10/2026): com a mesma foto da geração oficial, agenda e pendentes saíram idênticos.
    Não gasta cota do Fracttal e pode rodar no horário de campo. A rodada leva o sufixo "-foto" e não entra na lista da
    tela (`RODADA_RE`): é prova, não geração. `historico`: "banco" (o da geração) ou "nexus" (a reserva, o histórico de
    antes de 08/10), para comparar os dois com a mesma foto."""
    if not SEMANA_RE.match(semana or ""):
        raise ValueError("semana inválida")
    if historico not in ("banco", "nexus"):
        raise ValueError("histórico: 'banco' ou 'nexus'")
    origem = pasta_origem(config)
    foto = Path(foto) if foto else origem
    if not (foto and (foto / FOTO_OBRIGATORIA).exists()):
        raise NaoPronto(f"não há foto do Fracttal ({FOTO_OBRIGATORIA}) em {foto}")
    trabalho = pasta_trabalho(config)
    sufixo = re.sub(r"[^a-z0-9]+", "-", (rotulo or "").lower()).strip("-")
    rid = datetime.now(_BRT).strftime("%Y%m%d-%H%M%S") + f"-{semana}-foto" + (f"-{sufixo}" if sufixo else "")
    pasta = trabalho / "geracoes" / rid
    (pasta / "saida").mkdir(parents=True, exist_ok=False)
    try:
        insumos = _preparar(config, semana, pasta, trabalho, origem, historico)
    except NaoPronto:
        shutil.rmtree(pasta, ignore_errors=True)
        raise
    copiados = {}
    for nome in FOTO:
        p = foto / nome
        if p.exists():
            shutil.copy2(p, pasta / nome)       # copy2: a idade do arquivo vai junto (o TTL do motor é por ela)
            copiados[nome] = {"sha": _sha(pasta / nome), "de": _quando(p.stat().st_mtime)}
    st = {"id": rid, "semana": semana, "estado": "rodando", "inicio": datetime.now(_BRT).isoformat(timespec="seconds"),
          "credencial": f"nenhuma: foto do Fracttal de {copiados[FOTO_OBRIGATORIA]['de']}", "insumos": insumos,
          "motor": str(caminho_motor(config)), "historico": historico,
          "foto": {"pasta": str(foto), "arquivos": copiados}}
    _gravar(pasta, st)
    saida = pasta / "saida" / "sombra.xlsx"
    _executar(pasta, st, _ambiente_foto(pasta, saida), caminho_motor(config), origem, saida)
    # a prova de que o Fracttal não foi lido: o motor regrava o cache quando lê, e aqui ele saiu como entrou
    st["foto"]["intacta"] = all(_sha(pasta / n) == c["sha"] for n, c in copiados.items())
    _gravar(pasta, st)
    return st


def _cauda(p: Path, n: int) -> str:
    try:
        linhas = p.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return ""
    return "\n".join(linhas[-n:]).strip()


def _pasta_rodada(config, rid: str) -> Path | None:
    if not RODADA_RE.match(rid or ""):
        return None
    p = pasta_trabalho(config) / "geracoes" / rid
    return p if (p / "status.json").exists() else None


def ler_status(config, rid: str) -> dict | None:
    p = _pasta_rodada(config, rid)
    if not p:
        return None
    try:
        return json.loads((p / "status.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def log(config, rid: str, ultimas: int | None = None) -> str:
    p = _pasta_rodada(config, rid)
    if not p:
        return ""
    if ultimas:
        return _cauda(p / "log_motor.txt", ultimas)
    try:
        return (p / "log_motor.txt").read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""


def ativa(config, rid: str) -> bool:
    """A rodada está de fato rodando neste processo? Um status "rodando" de um Nexus que reiniciou no meio não está."""
    ativo = _rodando.get(str(pasta_trabalho(config).resolve()))
    return bool(ativo and ativo[0] == rid and ativo[1].is_alive())


def aguardar(config, rid: str, timeout: float) -> dict | None:
    ativo = _rodando.get(str(pasta_trabalho(config).resolve()))
    if ativo and ativo[0] == rid:
        ativo[1].join(timeout)
    return ler_status(config, rid)


def ultimas(config, n: int = 8) -> list[dict]:
    base = pasta_trabalho(config) / "geracoes"
    if not base.is_dir():
        return []
    out = []
    for p in sorted((x for x in base.iterdir() if RODADA_RE.match(x.name)), reverse=True)[:n]:
        st = ler_status(config, p.name)
        if st:
            out.append(st)
    return out
