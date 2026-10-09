"""O histórico das programações que o motor lê (`Historico_Programacoes.xlsx`), tirado do BANCO.

Levi, 08/10/2026: "PUXE O HISTÓRICO". Até aqui o histórico era o arquivo do PC do programador do PCM, importado para o
Nexus. Esse arquivo guarda a UNIÃO de todas as gerações de cada semana (a W41 tem as duas de 02/10, 08:26 e 18:42:
1.286 tarefas contra as 1.119 que foram ao campo; a W40, 1.306 contra 1.109) e perdeu a W35 inteira. A programação que
foi ao campo está no banco desde 08/10 (`nexus_programacao · fato_programacao`, W21 em diante,
`nexus/dados/programacao.py`); daqui sai o histórico.

**O que o motor lê** (`programacao_v7.py`, linhas 742 e 926-938): a 1ª aba, cabeçalho na linha 1, colunas `task_key`,
`first_week`, `last_week`, `count`, `weeks`. Ele usa só `count` (o "Nº vezes programada" é `count + 1`, a
"Reprogramada" é `count > 0`, e o RPN dinâmico, em sombra, soma `0,75 × count`) e `weeks` (para regravar). A chave é
`task_key(os, codigo) = f"{int(os)}|{codigo.strip()}"`: OS + código do equipamento, SEM a tarefa. Uma semana conta
para a chave quando a chave teve bloco na agenda dela (o motor grava só as linhas agendadas, nunca as pendentes).

**A chave do banco é a do motor:** o fato guarda a OS e o `codigo_ativo` (maiúsculo e sem espaço, a espinha da dimensão
de equipamento). As chaves do arquivo de histórico do PC não têm código com espaço nem minúscula (só 5 de OS negativa,
"nan", que nunca foram ao App). Na foto do Fracttal de 08/10, 3 de 6.226 códigos têm ("Transformador 2"), nenhum na
W42; se um deles for programado, a chave do banco não o acharia (contaria como 1ª vez).

**Semana em andamento.** No arquivo do App, a semana ATIVA muda ao longo dela: à noite o robô do PCM tira da agenda o que
não foi feito ("Rolagem: sem capacidade", "Deslocada por OS #...") e põe Nova OS; quando a semana fecha, ele a refaz pelo
plano publicado (W40: 858 chaves em 02/10 19:30 e 1.356 às 22:10; a W41 foi de 1.556 em 05/10 a 971 em 08/10). O motor
gera a semana N durante a N-1, então a N-1 do banco é sempre um retrato do meio da semana, sem as tarefas que rolaram
(na W41 de 08/10, 247 tarefas de OS vivas). Por isso a semana que ainda não tinha acabado quando o banco foi gravado
(`atualizacao.gerado_em`) usa o PLANO PUBLICADO dela: o guardado no Nexus (a reserva, `insumos.json`; a W41 de lá é a
S41 que foi ao campo) ou, na falta, a planilha oficial da pasta do PCM. Sem nenhum dos dois, vale a do banco, com aviso.
Semana fechada no banco é o plano publicado menos as OS canceladas depois (o robô tira; a OS cancelada não volta ao
Fracttal, então não muda geração nenhuma). Até agosto o robô também deixava na semana fechada a Nova OS que entrou nela
(1.189 chaves de W22 a W35): essas contam, porque foram ao campo naquela semana.

Só lê: GET na API do banco (leitura aberta) e, para a semana em andamento, o plano publicado.

**De onde vem o plano publicado (09/10/2026, "full Nexus"):** a reserva do Nexus (o plano que o próprio Nexus publicou
ou o corrigido à mão); senão a `Programação Semana NN.xlsx` do REPOSITÓRIO do PCM (o que foi ao App, publicado pelo PC do
PCM ou pelo Nexus; funciona no servidor, onde a pasta do PCM não existe); senão a da pasta do PCM. A planilha do
repositório só vale se os dias dela forem os da semana (o nome não tem ano: a "Semana 43" de 2026 não serve para 2027).
"""
import io
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl

from ..dados import programacao as P

NOME = "Historico_Programacoes.xlsx"
CABECALHO = ["task_key", "first_week", "last_week", "count", "weeks"]
_BRT = timezone(timedelta(hours=-3))
_SEMANA_RE = re.compile(r"^(\d{4})-W(\d{2})$")

ORIGEM_BANCO = "banco"
ORIGEM_RESERVA = "plano publicado guardado no Nexus"
ORIGEM_OFICIAL = "planilha oficial do PCM"
ORIGEM_REPO = "planilha publicada no repositório do PCM"
ORIGEM_ANDAMENTO = "banco, semana em andamento (incompleta)"


class HistoricoErro(RuntimeError):
    pass


# ── chaves e semanas ────────────────────────────────────────────────────────────────────────────────────────────
def chave_da_tarefa(os_, codigo) -> str | None:
    """A chave do motor (`programacao_v7.task_key`): f"{int(os)}|{codigo.strip()}". OS que não é número ou código
    vazio = None (o motor nunca gravaria uma chave assim a partir do Fracttal)."""
    try:
        o = int(float(str(os_).strip()))
    except (TypeError, ValueError):
        return None
    c = "" if codigo is None else str(codigo).strip()
    return f"{o}|{c}" if c else None


def segunda(semana: str) -> date:
    m = _SEMANA_RE.match(semana or "")
    if not m:
        raise ValueError(f"semana inválida: {semana!r}")
    return date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1)


def fim_da_semana(semana: str) -> datetime:
    """A segunda-feira seguinte, 00:00 de Brasília: a partir daí a semana acabou."""
    return datetime.combine(segunda(semana) + timedelta(days=7), datetime.min.time(), _BRT)


def _semana_de(d: date) -> str:
    a, s, _ = d.isocalendar()
    return f"{a}-W{s:02d}"


def _quando(v) -> datetime | None:
    try:
        d = datetime.fromisoformat(str(v).strip().replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None
    return d if d.tzinfo else d.replace(tzinfo=_BRT)


def semanas_do_banco(linhas: list) -> tuple[dict, dict]:
    """({semana: {chaves}}, contagens) das linhas do `fato_programacao` como o banco devolve (dicionários; a API pode
    mandar número como texto)."""
    out, sem_chave = defaultdict(set), 0
    for l in linhas or []:
        w = P.semana_iso(l.get("data_id_semana"))
        k = chave_da_tarefa(l.get("os"), l.get("codigo_ativo"))
        if not w or not k:
            sem_chave += 1
            continue
        out[w].add(k)
    return dict(out), {"blocos": len(linhas or []), "sem_chave": sem_chave}


def semanas_da_reserva(tab: dict | None) -> dict:
    """{semana: {chaves}} do histórico guardado no Nexus (`insumos.json`, o mesmo formato do arquivo do motor)."""
    if not tab:
        return {}
    ix = {c: i for i, c in enumerate(tab.get("colunas") or [])}
    if "task_key" not in ix or "weeks" not in ix:
        return {}
    out = defaultdict(set)
    for r in tab.get("linhas") or []:
        if len(r) <= max(ix["task_key"], ix["weeks"]) or not r[ix["task_key"]]:
            continue
        for w in str(r[ix["weeks"]] or "").split(","):
            if w.strip():
                out[w.strip()].add(str(r[ix["task_key"]]))
    return dict(out)


_DIA_RE = re.compile(r"^\s*(\S+)[^(]*\((\d{2})/(\d{2})\)")
_NOMES_DIA = ("segunda", "terca", "quarta", "quinta", "sexta")


def planilha_da_semana(arquivo, semana: str) -> bool:
    """Os dias da planilha (a coluna "Dia" das abas de equipe: "Segunda-feira (12/10)") são os da semana? O nome do
    arquivo não tem ano: o dia da semana tem de bater com o dd/mm (a "Semana 43" de 2025 tem a segunda em 20/10, que em
    2026 é uma terça)."""
    import unicodedata
    seg = segunda(semana)
    dias = {(seg + timedelta(days=i)).strftime("%d/%m"): _NOMES_DIA[i] for i in range(5)}
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if ws.title.startswith("_"):
                continue
            it = ws.iter_rows(values_only=True)
            cab = next(it, None)
            if not cab or "Dia" not in cab:
                continue
            i = cab.index("Dia")
            for r in it:
                m = _DIA_RE.search(str(r[i] or "")) if r and len(r) > i else None
                if m:
                    nome = unicodedata.normalize("NFKD", m.group(1)).encode("ascii", "ignore").decode().lower()
                    return nome.startswith(dias.get(f"{m.group(2)}/{m.group(3)}", "-"))
    finally:
        wb.close()
    return False


def chaves_da_planilha(caminho: Path) -> set:
    """As chaves das tarefas AGENDADAS de uma `Programação Semana NN.xlsx` (as abas de equipe; `_Pendentes` e as outras
    abas `_` ficam fora, como no motor, que só grava no histórico as linhas agendadas)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ks = set()
    try:
        for ws in wb.worksheets:
            if ws.title.startswith("_"):
                continue
            it = ws.iter_rows(values_only=True)
            cab = next(it, None)
            if not cab or "OSs ID" not in cab or "Código Equipamento" not in cab:
                continue
            i_os, i_cod = cab.index("OSs ID"), cab.index("Código Equipamento")
            for r in it:
                k = chave_da_tarefa(r[i_os], r[i_cod]) if r and len(r) > max(i_os, i_cod) else None
                if k:
                    ks.add(k)
    finally:
        wb.close()
    return ks


# ── a montagem ──────────────────────────────────────────────────────────────────────────────────────────────────
def em_andamento(semana: str, banco_ate: datetime | None) -> bool:
    """A semana ainda não tinha acabado quando o banco foi gravado: o que está lá é o retrato do meio da semana."""
    return banco_ate is None or fim_da_semana(semana) > banco_ate


def semanas_sem_fechar(do_banco: dict, banco_ate: datetime | None, semana: str) -> list[str]:
    """As semanas anteriores à gerada que precisam do plano publicado: as do banco ainda em andamento e as que vieram
    depois da última semana do banco (o banco não foi gravado desde que elas saíram)."""
    antes = [w for w in do_banco if w < semana]
    out = [w for w in antes if em_andamento(w, banco_ate)]
    if antes:
        d = segunda(max(antes)) + timedelta(days=7)
        while _semana_de(d) < semana:
            out.append(_semana_de(d))
            d += timedelta(days=7)
    return sorted(set(out))


def montar(do_banco: dict, semana: str, banco_ate: datetime | None, planos: dict | None = None) -> tuple[list, dict]:
    """(linhas do histórico no formato do motor, relatório por semana). Só entram as semanas ANTERIORES à gerada (o
    histórico "de antes da semana": gerar de novo a mesma semana não pode contá-la, S40 em 25/09: 415 tarefas viraram
    reprogramadas). `planos` = {semana: (origem, {chaves})}: o plano publicado das semanas que não fecharam no banco.
    Semana anterior à 1ª do banco não entra: o banco é a fonte, e antes dele só há o arquivo do PC."""
    planos = planos or {}
    usadas, rel = {}, {}
    pendentes = set(semanas_sem_fechar(do_banco, banco_ate, semana))
    for w in sorted(set(do_banco) | pendentes):
        if w >= semana:
            continue
        no_banco = do_banco.get(w)
        if w in pendentes:
            if w in planos:
                origem, ks = planos[w]
                usadas[w] = set(ks)
                rel[w] = {"origem": origem, "chaves": len(ks), "no_banco": len(no_banco or ())}
            elif no_banco:
                usadas[w] = set(no_banco)
                rel[w] = {"origem": ORIGEM_ANDAMENTO, "chaves": len(no_banco), "no_banco": len(no_banco),
                          "aviso": True}
            else:
                rel[w] = {"origem": "fora do banco e sem plano publicado", "chaves": 0, "no_banco": 0, "aviso": True}
            continue
        usadas[w] = set(no_banco)
        rel[w] = {"origem": ORIGEM_BANCO, "chaves": len(no_banco), "no_banco": len(no_banco)}
    por_chave = defaultdict(set)
    for w, ks in usadas.items():
        for k in ks:
            por_chave[k].add(w)
    linhas = []
    for k in sorted(por_chave):
        ws = sorted(por_chave[k])
        linhas.append([k, ws[0], ws[-1], len(ws), ",".join(ws)])
    semanas = sorted(w for w, r in rel.items() if r["chaves"])
    return linhas, {"semana_gerada": semana, "banco_ate": banco_ate.isoformat(timespec="minutes") if banco_ate else None,
                    "semanas": rel, "tarefas": len(linhas), "de": semanas[0] if semanas else None,
                    "ate": semanas[-1] if semanas else None,
                    "avisos": [f"{w}: {r['origem']}" for w, r in sorted(rel.items()) if r.get("aviso")]}


def escrever(linhas: list, destino: Path) -> None:
    """Uma aba (o motor lê a primeira), cabeçalho na linha 1, como o próprio motor grava (pandas, aba Sheet1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(CABECALHO)
    for l in linhas:
        ws.append(l)
    wb.save(destino)


def detalhe(rel: dict) -> str:
    """Uma linha para a rodada: de onde veio e o que é aviso."""
    s = rel["semanas"]
    n = sum(1 for r in s.values() if r["chaves"])
    txt = (f"do banco (nexus_programacao · fato_programacao): {n} semanas, {rel['de']} a {rel['ate']}, "
           f"{rel['tarefas']} tarefas")
    outras = [f"{w} pelo {r['origem']}" for w, r in sorted(s.items()) if r["origem"] not in (ORIGEM_BANCO,)]
    return txt + ("; " + "; ".join(outras) if outras else "")


# ── a leitura do banco e a rodada ──────────────────────────────────────────────────────────────────────────────
def _base(config) -> str:
    from ..cadastro.ligacoes import BASE_API
    return str(config.get("GRIDCO_DB_API") or BASE_API).rstrip("/")


def conferir_leitura(linhas: list, atualizacao: list) -> None:
    """A leitura é a gravação inteira? A carga grava, junto do fato, a contagem e o sha das linhas (`atualizacao`). Uma
    leitura parcial que volte com 200 (página que falhou calada, leitura no meio de uma gravação) daria um histórico
    menor sem ninguém ver: aqui ela vira erro, e a geração usa a reserva e diz por quê. Medido em 09/10: 14.276 de
    14.276 linhas, sha igual, 0,4 s."""
    a = (atualizacao or [{}])[0]
    try:
        n = int(float(str(a.get("linhas")).strip()))
    except (TypeError, ValueError):
        n = None
    if n is not None and n != len(linhas):
        raise HistoricoErro(f"o banco devolveu {len(linhas)} de {n} linhas do fato")
    sha = str(a.get("sha_linhas") or "").strip()
    if sha and P.sha_linhas(P.mesclar_semanas(linhas, [])) != sha:
        raise HistoricoErro("as linhas lidas não são as gravadas (o sha não bate)")


def ler_do_banco(config, sessao=None) -> tuple[list, datetime | None]:
    """(linhas do fato_programacao, até quando o banco sabe). Só GET: a API do banco tem leitura aberta. Nos testes, sem
    `NEXUS_PCM_HISTORICO_TESTE` ({"linhas": [...], "gerado_em": "..."}), não vai à rede: falha como banco fora do ar."""
    if config.get("TESTING"):
        teste = config.get("NEXUS_PCM_HISTORICO_TESTE")
        if teste is None:
            raise HistoricoErro("sem banco nos testes")
        if isinstance(teste, Exception):
            raise HistoricoErro(str(teste))
        return list(teste.get("linhas") or []), _quando(teste.get("gerado_em"))
    import requests
    from ..dados import livros
    s = sessao or requests.Session()
    base = _base(config)
    try:
        abas = livros.abas(base, s, P.LIVRO)
        if P.ABA not in abas:
            raise HistoricoErro(f"o banco não tem {P.LIVRO} · {P.ABA}")
        linhas = livros._linhas(base, s, abas[P.ABA])
        atual = livros._linhas(base, s, abas["atualizacao"]) if "atualizacao" in abas else []
    except HistoricoErro:
        raise
    except Exception as ex:      # noqa: BLE001 — rede, 500, JSON quebrado: a geração usa a reserva e diz por quê
        raise HistoricoErro(f"{type(ex).__name__}: {str(ex)[:160]}") from ex
    if not linhas:
        raise HistoricoErro(f"{P.LIVRO} · {P.ABA} veio vazio")
    conferir_leitura(linhas, atual)
    ate = max((d for d in (_quando(a.get("gerado_em")) for a in atual) if d), default=None)
    if ate is None:          # sem a linha da gravação, vale a última leitura que está no fato
        ate = max((d for d in (_quando(l.get("lido_em")) for l in linhas) if d), default=None)
    return linhas, ate


def planos_publicados(semanas: list[str], reserva: dict | None = None, origem: Path | None = None,
                      ler_repo=None) -> dict:
    """{semana: (origem, {chaves})} das semanas pedidas: a reserva do Nexus (curada: a W41 de lá é a S41 que foi ao
    campo), a planilha do repositório do PCM (`ler_repo(nome)` -> bytes ou None: o que foi ao App) e a da pasta do PCM
    (só leitura)."""
    da_reserva = semanas_da_reserva(reserva)
    out = {}
    for w in semanas:
        if da_reserva.get(w):
            out[w] = (ORIGEM_RESERVA, da_reserva[w])
            continue
        nome = f"Programação Semana {segunda(w).isocalendar()[1]:02d}.xlsx"
        dados = None
        if ler_repo:
            try:
                dados = ler_repo(nome)
            except Exception:       # noqa: BLE001 — sem rede: segue para a pasta do PCM e, na falta, o banco com aviso
                dados = None
        if dados and planilha_da_semana(io.BytesIO(dados), w):
            ks = chaves_da_planilha(io.BytesIO(dados))
            if ks:
                out[w] = (ORIGEM_REPO + f" ({nome})", ks)
                continue
        p = Path(origem) / nome if origem else None
        if p and p.exists():
            ks = chaves_da_planilha(p)
            if ks:
                out[w] = (ORIGEM_OFICIAL + f" ({p.name})", ks)
    return out


def materializar(config, destino: Path, semana: str, reserva: dict | None = None, origem: Path | None = None,
                 sessao=None, ler_repo=None) -> dict:
    """Escreve na rodada o histórico do banco, no formato do motor. Devolve o relatório (sem nome de pessoa: só
    chaves de OS e código, que nem vão no relatório)."""
    linhas, ate = ler_do_banco(config, sessao)
    do_banco, cont = semanas_do_banco(linhas)
    if not do_banco:
        raise HistoricoErro("nenhuma semana com chave no fato do banco")
    planos = planos_publicados(semanas_sem_fechar(do_banco, ate, semana), reserva, origem, ler_repo)
    hist, rel = montar(do_banco, semana, ate, planos)
    if not hist:
        raise HistoricoErro(f"nenhuma semana antes de {semana} no banco")
    escrever(hist, destino)
    rel.update(cont, fonte="banco", detalhe=None)
    rel["detalhe"] = detalhe(rel)
    return rel
