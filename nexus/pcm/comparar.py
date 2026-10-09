"""Compara duas gerações da MESMA semana, linha a linha: a oficial (a do PCM) e a do Nexus.

Chave da linha: OS + código do equipamento + tarefa (a mesma tarefa aparece em mais de um ativo da OS). Compara equipe,
dia, hora, duração, reprogramada, nº de vezes e RPN; e as pendentes. É a mesma conta do sombra/comparar.py do plano,
testada nas duas gerações da semana 40 (417 reprogramadas por engano, +1 vez nas 1.011 que estavam nas duas).
"""
import re
from collections import Counter

import openpyxl

CAMPOS = ["Equipe", "Dia", "Hora Início", "Duração (h)", "Reprogramada", "Nº vezes programada", "RPN/Prioridade"]


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return ("%.2f" % v).rstrip("0").rstrip(".")
    return str(v).strip()


def ler(caminho) -> tuple[dict, Counter]:
    """(linhas por chave, pendentes por chave) de uma Programação Semana XX.xlsx."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    linhas, pend, dup = {}, Counter(), 0
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            cab = next(it, None)
            if not cab or "OSs ID" not in cab:
                continue
            ix = {c: i for i, c in enumerate(cab)}
            for r in it:
                if not r or r[ix["OSs ID"]] in (None, ""):
                    continue
                chave = (_txt(r[ix["OSs ID"]]),
                         _txt(r[ix["Código Equipamento"]]) if "Código Equipamento" in ix else "",
                         _txt(r[ix["Tarefa"]]) if "Tarefa" in ix else "")
                if ws.title == "_Pendentes":
                    pend[chave] += 1
                    continue
                if ws.title.startswith("_"):
                    continue
                reg = {c: _txt(r[ix[c]]) if c in ix else "" for c in CAMPOS}
                reg["Equipe"] = reg["Equipe"] or ws.title
                # "Quarta-feira (01/10) [NOTURNO]": a marcação vira campo próprio
                reg["Marcação"] = " ".join(re.findall(r"\[[^\]]+\]", reg["Dia"]))
                reg["Dia"] = re.sub(r"\s*\[[^\]]+\]", "", reg["Dia"]).strip()
                if chave in linhas:
                    dup += 1
                    chave = chave + ("#%d" % dup,)
                linhas[chave] = reg
    finally:
        wb.close()
    return linhas, pend


def comparar(oficial, nexus) -> dict:
    A, pa = ler(oficial)
    B, pb = ler(nexus)
    comuns = sorted(set(A) & set(B))
    so_a, so_b = sorted(set(A) - set(B)), sorted(set(B) - set(A))
    dif, exemplos = Counter(), {}
    for k in comuns:
        for c in CAMPOS + ["Marcação"]:
            if A[k][c] != B[k][c]:
                dif[c] += 1
                exemplos.setdefault(c, {"os": k[0], "oficial": A[k][c], "nexus": B[k][c]})
    iguais = sum(1 for k in comuns if all(A[k][c] == B[k][c] for c in CAMPOS + ["Marcação"]))
    return {
        "linhas_a": len(A), "linhas_b": len(B), "comuns": len(comuns), "iguais": iguais,
        "so_a": len(so_a), "so_b": len(so_b),
        "os_so_a": sorted({k[0] for k in so_a})[:15], "os_so_b": sorted({k[0] for k in so_b})[:15],
        "diferencas": dict(dif.most_common()), "exemplos": exemplos,
        "pendentes_a": sum(pa.values()), "pendentes_b": sum(pb.values()), "pendentes_iguais": pa == pb,
        "identicas": not so_a and not so_b and not dif and pa == pb,
    }


def resumo(caminho) -> dict:
    linhas, pend = ler(caminho)
    return {"linhas": len(linhas), "pendentes": sum(pend.values()),
            "equipes": len({r["Equipe"] for r in linhas.values()})}


# ── a conferência da semana gerada (antes de alguém publicar) ─────────────────────────────────────────────────
# A capacidade e a jornada são as da tela Semana (semana.py), que são as do gerador: 7 h por equipe e dia (80% das
# 8,8 h líquidas), MPA na janela da noite fora da conta, zeladoria terceirizada fora da conta e da jornada.
CAPACIDADE_H = 7.0
JORNADA = ("07:00", "17:00")
# o almoço do gerador (programacao_v7: 12:00 às 13:12): o bloco que o atravessa termina 72 min depois
ALMOCO = (12 * 60, 13 * 60 + 12)
# o dia útil do gerador em minutos (USEFUL_DAY_MIN = int((17:00 - 07:00 - 72 min) × 0,80)); as 7 h da tela arredondam
DIA_DO_MOTOR_MIN = 422
_DIA_RE = re.compile(r"^\s*([^\s(]+)[^(]*\((\d{2})/(\d{2})\)")
_HORA_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})(?:\s*\(D\+(\d+)\))?")
_NOMES_DIA = ("segunda", "terca", "quarta", "quinta", "sexta")


def _minutos(h) -> int | None:
    """"16:50" -> 1010; "03:15 (D+1)" -> 1635 (o motor escreve a hora do dia seguinte assim)."""
    m = _HORA_RE.match(str(h or ""))
    return None if not m else int(m.group(1)) * 60 + int(m.group(2)) + 1440 * int(m.group(3) or 0)


def _no_relogio(ini, fim) -> int | None:
    """Os minutos de trabalho de um bloco de dia pelo relógio: do início ao fim, menos o almoço que ele atravessa."""
    a, b = _minutos(ini), _minutos(fim)
    if a is None or b is None or b < a:
        return None
    return b - a - (ALMOCO[1] - ALMOCO[0] if a < ALMOCO[0] and b > ALMOCO[0] else 0)


def _sem_acento(s) -> str:
    import unicodedata
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def linhas_da_planilha(caminho) -> tuple[list[dict], list[dict]]:
    """(blocos agendados, pendentes) de uma Programação Semana XX.xlsx, cada um como {coluna: valor}. Abas que começam
    com "_" não são de equipe (`_Pendentes`, `_Resumo`, `_Qualidade`...)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    blocos, pend = [], []
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            cab = next(it, None)
            if not cab or "OSs ID" not in cab:
                continue
            for r in it:
                if not r or r[cab.index("OSs ID")] in (None, ""):
                    continue
                d = dict(zip(cab, r))
                if ws.title == "_Pendentes":
                    pend.append(d)
                elif not ws.title.startswith("_"):
                    d.setdefault("Equipe", ws.title)
                    d["Equipe"] = d["Equipe"] or ws.title
                    blocos.append(d)
    finally:
        wb.close()
    return blocos, pend


def conferir(caminho, semana: str) -> dict:
    """O que olhar numa semana gerada antes de publicar: quantos blocos, tarefas, OS, pendentes e equipes; se todo bloco
    cai num dia útil DESTA semana (o dia por extenso bate com o dd/mm); horas por equipe e dia contra as 7 h (fora da
    conta: a janela da noite, MPA ou [NOTURNO], e a zeladoria terceirizada), pela coluna "Duração (h)" (a soma que a
    tela Semana e o App mostram) e pelo relógio dos blocos (contra os 422 min do motor), com os blocos cuja duração não
    cabe no horário; blocos de dia que começam fora da jornada (07:00-17:00, ou no dia seguinte, "D+1"); e o que aparece
    repetido (a mesma tarefa em mais de um bloco, a mesma OS em mais de uma equipe). Só contagens e números de OS."""
    from datetime import date, timedelta
    m = re.fullmatch(r"(\d{4})-W(\d{2})", semana)
    seg = date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1)
    dias = {(seg + timedelta(days=i)).strftime("%d/%m"): i for i in range(5)}
    blocos, pend = linhas_da_planilha(caminho)
    fora_semana, horas, fora_jornada, flags = [], Counter(), Counter(), Counter()
    noite, horas_relogio, curtos = Counter(), Counter(), []
    por_tarefa, os_equipes, partidas = Counter(), {}, set()
    # [EXCEDE HH]: a corretiva que o motor forçou num dia já cheio. W42 de 09/10: a OS 14331 (4 tarefas de 6 h, a duração
    # estimada no Fracttal; o total de horas da OS é 1,5 h por tarefa) pôs 24,5 h numa quarta de uma equipe
    excede = {}
    for b in blocos:
        dia = str(b.get("Dia") or "")
        mm = _DIA_RE.match(dia)
        idx = dias.get(f"{mm.group(2)}/{mm.group(3)}") if mm else None
        if idx is None or not _sem_acento(mm.group(1)).startswith(_NOMES_DIA[idx]):
            fora_semana.append({"os": _txt(b.get("OSs ID")), "dia": dia})
            continue
        marcas = set()
        for f in re.findall(r"\[([^\]]+)\]", dia):
            # as marcações do motor vêm juntas ("NOTURNO/EXCEDE JANELA MPA"); a da zeladoria tem a barra no nome
            for x in f.replace("ZELADORIA/PARALELO", "ZELADORIA (PARALELO)").split("/"):
                marcas.add(re.sub(r"^\+\d+D ", "+ND ", x.strip()))
        flags.update(marcas)
        if "EXCEDE HH" in marcas:
            e = excede.setdefault(_txt(b.get("OSs ID")), {"os": _txt(b.get("OSs ID")), "equipe": str(b.get("Equipe")),
                                                         "dia": _NOMES_DIA[idx][:3], "tarefas": 0, "horas": 0.0})
            e["tarefas"] += 1
            e["horas"] = round(e["horas"] + _num(b.get("Duração (h)")), 2)
        tipo, paralelo = _sem_acento(b.get("Tipo")), _sem_acento(b.get("Paralelo (terceirizada)"))
        chave_dia = (str(b.get("Equipe")), idx)
        # a janela da noite é da MPA, e o motor põe nela também a corretiva com "MPA" no texto ([NOTURNO]): W42 de
        # 09/10, uma corretiva de MPA das 15:45 à 00:00 contava 8 h no dia de uma equipe que tinha 5,55 h de dia
        if tipo == "mpa" or "NOTURNO" in marcas:
            noite[chave_dia] += _num(b.get("Duração (h)"))
        elif paralelo != "sim":
            horas[chave_dia] += _num(b.get("Duração (h)"))
            h = str(b.get("Hora Início") or "")
            if "(D+" in h or (h[:5] and (h[:5] < JORNADA[0] or h[:5] >= JORNADA[1])):
                fora_jornada[str(b.get("Tipo") or "sem tipo")] += 1
            # a duração da coluna contra o relógio do bloco: na W42 de 09/10 o último MPS do dia saía 16:18-16:50
            # (32 min) com "1,3 h" na coluna, sem [PARCIAL], e o dia "passava" de 7 h só na soma da coluna
            rel = _no_relogio(b.get("Hora Início"), b.get("Hora Fim"))
            horas_relogio[chave_dia] += rel / 60 if rel is not None else _num(b.get("Duração (h)"))
            if rel is not None and rel + 1 < round(_num(b.get("Duração (h)")) * 60) and "PARCIAL" not in marcas:
                curtos.append({"os": _txt(b.get("OSs ID")), "equipe": str(b.get("Equipe")), "dia": _NOMES_DIA[idx][:3],
                               "inicio": h, "fim": str(b.get("Hora Fim") or ""), "duracao_h": _num(b.get("Duração (h)"))})
        k = (_txt(b.get("OSs ID")), _txt(b.get("Código Equipamento")), _txt(b.get("Tarefa")))
        por_tarefa[k] += 1
        if "PARCIAL" in marcas:
            partidas.add(k)
        os_equipes.setdefault(_txt(b.get("OSs ID")), set()).add(str(b.get("Equipe")))
    estouros = sorted(((eq, i, round(h, 2)) for (eq, i), h in horas.items() if h > CAPACIDADE_H + 1e-9),
                      key=lambda x: -x[2])
    repetidas = {k: n for k, n in por_tarefa.items() if n > 1}
    # a tarefa que não cabe no resto do dia é PARTIDA pelo motor (um pedaço [PARCIAL] e o resto no dia seguinte): dois
    # blocos da mesma tarefa são esperados aí; repetida sem PARCIAL é a que merece olhar
    sem_parcial = sorted(k for k in repetidas if k not in partidas)
    os_multi = {o: sorted(e) for o, e in os_equipes.items() if len(e) > 1}
    equipes_com = {str(b.get("Equipe")) for b in blocos}
    nomes = ("seg", "ter", "qua", "qui", "sex")
    return {
        "blocos": len(blocos), "tarefas": len(por_tarefa), "os": len(os_equipes),
        "pendentes": len(pend), "pendentes_os": len({_txt(p.get("OSs ID")) for p in pend}),
        "equipes_com_bloco": len(equipes_com),
        "fora_da_semana": len(fora_semana), "fora_da_semana_exemplos": fora_semana[:5],
        "dias_de_equipe": len(horas), "horas_dia_max": round(max(horas.values()), 2) if horas else 0,
        "acima_da_capacidade": len(estouros),
        "acima_da_capacidade_exemplos": [{"equipe": e, "dia": nomes[i], "horas": h} for e, i, h in estouros[:8]],
        "acima_da_capacidade_pelo_relogio": sum(1 for h in horas_relogio.values() if h * 60 > DIA_DO_MOTOR_MIN + 1),
        "horas_dia_max_pelo_relogio": round(max(horas_relogio.values()), 2) if horas_relogio else 0,
        "duracao_maior_que_o_horario": len(curtos), "duracao_maior_que_o_horario_exemplos": curtos[:5],
        "horas_noite_max": round(max(noite.values()), 2) if noite else 0,
        "fora_da_jornada": sum(fora_jornada.values()), "fora_da_jornada_por_tipo": dict(fora_jornada.most_common()),
        "marcacoes": dict(flags.most_common()),
        "tarefas_em_mais_de_um_bloco": len(repetidas), "blocos_dessas_tarefas": sum(repetidas.values()),
        "repetidas_sem_parcial": len(sem_parcial), "repetidas_sem_parcial_os": sorted({k[0] for k in sem_parcial})[:10],
        "os_em_mais_de_uma_equipe": len(os_multi), "os_em_mais_de_uma_equipe_exemplos": dict(list(os_multi.items())[:5]),
        "reprogramadas": sum(1 for b in blocos if _sem_acento(b.get("Reprogramada")) == "sim"),
        "excede_hh_os": sorted(excede.values(), key=lambda e: -e["horas"]),
    }


def entre_semanas(anterior, atual) -> dict:
    """O que continua de uma semana para a outra, pela chave do motor (OS + código do equipamento) das tarefas
    AGENDADAS: quantas da semana nova já estavam agendadas na anterior, quantas são novas e quantas da anterior não
    voltaram (feitas, canceladas ou fora da semana nova)."""
    def chaves(caminho):
        return {f"{_txt(b.get('OSs ID'))}|{_txt(b.get('Código Equipamento'))}" for b in linhas_da_planilha(caminho)[0]}
    a, b = chaves(anterior), chaves(atual)
    return {"anterior": len(a), "atual": len(b), "continuam": len(a & b), "novas": len(b - a), "sairam": len(a - b)}
