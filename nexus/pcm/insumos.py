"""Etapa 3: os insumos do motor do PCM moram no Nexus, e não em arquivo.

Hoje o motor (programacao_v7.py, cópia idêntica em motor/) lê sete arquivos da pasta do PCM. Cinco deles passam a
ser dado do Nexus (o pedido do Levi em 30/09/2026: "matar arquivos em excel e termos nossa base estruturada"):

| Arquivo de hoje                      | No Nexus                                         |
|--------------------------------------|--------------------------------------------------|
| Observacoes_Semana.txt               | observações POR SEMANA, editáveis na tela        |
| Feriados/FERIADOS ... 2026.xlsx      | feriados: nacionais/estaduais e municipais       |
| Lista_Prioridades_GridCo.xlsx        | tabela de prioridades                            |
| Planilha Confiabilidade R00.xlsx     | tabela de tempos por categoria                   |
| Historico_Programacoes.xlsx          | histórico das semanas programadas                |

A AUXILIAR sai pelo cadastro (etapa 4) e o duracoes_aprendidas.json é derivado; os dois seguem vindo da pasta.

Como o motor não muda: a cada geração, o Nexus ESCREVE estes arquivos na pasta da rodada, no formato exato que o
motor lê (mesma aba, mesma linha de cabeçalho, mesma posição de coluna). Ninguém edita mais os arquivos.

Durante a sombra, o "Importar da pasta do PCM" traz para o Nexus o que a pasta do PCM tem hoje: as duas gerações partem
do mesmo dado e a comparação continua justa. Na virada, para de importar e os arquivos morrem.

O HISTÓRICO não se importa mais (08/10/2026, Levi: "PUXE O HISTÓRICO"): ele sai do banco a cada geração
(`historico_banco.py`). O arquivo do PC guarda a união de todas as gerações de cada semana (a W41 com as duas de 02/10)
e importá-lo trazia isso de volta por cima. O que está guardado aqui fica como RESERVA: vale se o banco não responder, e
dá o plano publicado da semana que o banco ainda não fechou (a W41 daqui é a S41 que foi ao campo, corrigida em 05/10).

Onde fica: um arquivo JSON na pasta de trabalho do PCM (fora do OneDrive), com cópia do anterior a cada gravação.
É o mesmo passo do cadastro (ensaio local primeiro, a API de dados depois): a troca é só neste módulo.
"""
import hashlib
import json
import os
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl

_BRT = timezone(timedelta(hours=-3))
_trava = threading.Lock()
VERSAO = 1

# ── o que o motor lê de cada planilha (programacao_v7.py) ───────────────────────────────────────────────────────
# chave: (arquivo, aba, linha do cabeçalho 1-based). O motor lê com pandas header=linha-1.
PLANILHAS = {
    "prioridades": ("Lista_Prioridades_GridCo.xlsx", "📋 Lista de Prioridades", 8),
    "confiabilidade": ("Planilha Confiabilidade R00.xlsx", "Resumo por Categoria", 4),
    # o histórico é lido pela 1ª aba, cabeçalho na linha 1; o motor reescreve com o pandas (aba Sheet1). Não se importa
    # mais (vem do banco, `historico_banco.py`); o guardado aqui é a reserva
    "historico": ("Historico_Programacoes.xlsx", None, 1),
}
# O que o "Importar da pasta do PCM" traz. O histórico ficou fora em 08/10/2026: o arquivo do PC guarda a união de todas
# as gerações de cada semana, e importá-lo punha a W41 com as duas gerações de 02/10 (1.286 tarefas, 167 que nunca foram
# ao campo) por cima da reserva corrigida.
IMPORTADAS = ("prioridades", "confiabilidade")
# Sem estes o motor não roda. O histórico tem o banco como fonte e esta reserva como reserva.
OBRIGATORIOS = ("prioridades", "confiabilidade", "feriados")
OBSERVACOES = "Observacoes_Semana.txt"
FERIADOS_PASTA = "Feriados"
FERIADOS_ARQ = "FERIADOS ESTADUAIS, MUNICIPAIS E NACIONAIS 2026.xlsx"
FERIADOS_ABA = "Feriados 2026"
# Bloco da esquerda (nacionais e estaduais) e da direita (municipais), separados por uma coluna vazia, como na
# planilha do PCM. O motor acha a coluna municipal pelo nome: "Tipo.1", "MUNICIPIO", "Data.1".
FER_ESQ = ["Tipo", "Estado", "Data", "Mês Feriado", "Feriado"]
FER_DIR = ["Tipo", "Estado", "MUNICIPIO", "Data", "Mês Feriado", "Feriado"]

NOMES = {
    "observacoes": "Observações da semana",
    "feriados": "Feriados",
    "prioridades": "Lista de Prioridades",
    "confiabilidade": "Confiabilidade (tempo por categoria)",
    "historico": "Histórico das programações",
}


class InsumoErro(ValueError):
    pass


# ── valores: datas viram texto no JSON e voltam a ser data na planilha ─────────────────────────────────────────

def _para_json(v):
    if isinstance(v, datetime):
        return {"$dt": v.isoformat()}
    if isinstance(v, date):
        return {"$d": v.isoformat()}
    return v


def _de_json(v):
    if isinstance(v, dict):
        if "$dt" in v:
            return datetime.fromisoformat(v["$dt"])
        if "$d" in v:
            return date.fromisoformat(v["$d"])
    return v


def _linha(r) -> list:
    """Linha sem as células vazias do fim (a largura da planilha varia com a formatação, não com o dado)."""
    r = list(r)
    while r and r[-1] is None:
        r.pop()
    return [_para_json(c) for c in r]


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()[:12]


def _agora() -> str:
    return datetime.now(_BRT).isoformat(timespec="seconds")


def _quando_arquivo(p: Path) -> str:
    return datetime.fromtimestamp(p.stat().st_mtime, _BRT).isoformat(timespec="seconds")


# ── armazém ────────────────────────────────────────────────────────────────────────────────────────────────

def caminho(pasta_trabalho: Path) -> Path:
    return Path(pasta_trabalho) / "insumos.json"


def carregar(pasta_trabalho: Path) -> dict:
    p = caminho(pasta_trabalho)
    if not p.exists():
        return {"versao": VERSAO, "observacoes": {}}
    dados = json.loads(p.read_text(encoding="utf-8"))
    dados.setdefault("observacoes", {})
    return dados


def salvar(pasta_trabalho: Path, dados: dict) -> None:
    """Gravação atômica, com a versão anterior guardada ao lado (insumos.json.anterior)."""
    p = caminho(pasta_trabalho)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")
    if p.exists():
        anterior = p.with_suffix(".json.anterior")
        anterior.unlink(missing_ok=True)
        p.replace(anterior)
    tmp.replace(p)


# ── leitura dos arquivos de hoje (importar) ────────────────────────────────────────────────────────────────

def ler_planilha(arquivo: Path, aba: str | None, linha_cab: int) -> dict:
    """A aba como o motor a vê: as linhas antes do cabeçalho (título e notas), o cabeçalho e as linhas de dado.
    Guardar as linhas de antes mantém o cabeçalho na mesma linha, que é por onde o motor (pandas header=N) o acha."""
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        if aba is not None and aba not in wb.sheetnames:
            raise InsumoErro(f"{arquivo.name}: falta a aba '{aba}'")
        ws = wb[aba] if aba is not None else wb.worksheets[0]
        todas = [_linha(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if len(todas) < linha_cab:
        raise InsumoErro(f"{arquivo.name}: não tem a linha {linha_cab} do cabeçalho")
    linhas = todas[linha_cab:]
    while linhas and not linhas[-1]:
        linhas.pop()
    return {"aba": aba, "preambulo": todas[:linha_cab - 1], "colunas": todas[linha_cab - 1], "linhas": linhas}


def ler_feriados(arquivo: Path) -> dict:
    """Os dois blocos da aba viram duas listas de registros. O bloco municipal começa depois da coluna vazia."""
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        if FERIADOS_ABA not in wb.sheetnames:
            raise InsumoErro(f"{arquivo.name}: falta a aba '{FERIADOS_ABA}'")
        rows = [list(r) for r in wb[FERIADOS_ABA].iter_rows(values_only=True)]
    finally:
        wb.close()
    if len(rows) < 2:
        raise InsumoErro(f"{arquivo.name}: aba sem cabeçalho")
    cab = [str(c).strip() if c is not None else "" for c in rows[1]]
    n_esq = len(FER_ESQ)
    if cab[:n_esq] != FER_ESQ or cab[n_esq + 1:n_esq + 1 + len(FER_DIR)] != FER_DIR:
        raise InsumoErro(f"{arquivo.name}: o cabeçalho da linha 2 mudou ({', '.join(c for c in cab if c)})")
    gerais, municipais = [], []
    for r in rows[2:]:
        r = r + [None] * (n_esq + 1 + len(FER_DIR) - len(r))
        esq, dir_ = r[:n_esq], r[n_esq + 1:n_esq + 1 + len(FER_DIR)]
        if any(c is not None for c in esq):
            gerais.append(dict(zip(["tipo", "estado", "data", "mes", "feriado"], [_para_json(c) for c in esq])))
        if any(c is not None for c in dir_):
            municipais.append(dict(zip(["tipo", "estado", "municipio", "data", "mes", "feriado"],
                                       [_para_json(c) for c in dir_])))
    titulo = rows[0][0] if rows and rows[0] else None
    return {"titulo": titulo, "gerais": gerais, "municipais": municipais}


def importar(pasta_trabalho: Path, origem: Path, semana: str | None = None) -> dict:
    """Traz para o Nexus o que está hoje na pasta do PCM. As observações entram na semana pedida."""
    origem = Path(origem)
    resumo = {}
    with _trava:
        dados = carregar(pasta_trabalho)
        resumo["historico"] = "não importado: vem do banco a cada geração (a reserva do Nexus fica como está)"
        for chave in IMPORTADAS:
            arq, aba, cab = PLANILHAS[chave]
            p = origem / arq
            if not p.exists():
                resumo[chave] = "não está na pasta"
                continue
            tab = ler_planilha(p, aba, cab)
            tab["meta"] = {"arquivo": arq, "sha": _sha(p), "arquivo_em": _quando_arquivo(p), "importado_em": _agora()}
            dados[chave] = tab
            resumo[chave] = f"{len(tab['linhas'])} linhas"
        pf = origem / FERIADOS_PASTA / FERIADOS_ARQ
        if pf.exists():
            fer = ler_feriados(pf)
            fer["meta"] = {"arquivo": f"{FERIADOS_PASTA}/{FERIADOS_ARQ}", "sha": _sha(pf),
                           "arquivo_em": _quando_arquivo(pf), "importado_em": _agora()}
            dados["feriados"] = fer
            resumo["feriados"] = f"{len(fer['gerais'])} nacionais e estaduais, {len(fer['municipais'])} municipais"
        else:
            resumo["feriados"] = "não está na pasta"
        po = origem / OBSERVACOES
        if semana and po.exists():
            texto = po.read_text(encoding="utf-8")
            dados["observacoes"][semana] = {"texto": texto, "atualizado_em": _agora(),
                                            "origem": {"arquivo": OBSERVACOES, "sha": _sha(po),
                                                       "arquivo_em": _quando_arquivo(po)}}
            resumo["observacoes"] = f"{len([l for l in texto.splitlines() if l.strip()])} linhas, na semana {semana}"
        salvar(pasta_trabalho, dados)
    return resumo


def salvar_observacoes(pasta_trabalho: Path, semana: str, texto: str, autor: str = "",
                       origem: dict | None = None) -> None:
    with _trava:
        dados = carregar(pasta_trabalho)
        dados["observacoes"][semana] = {"texto": texto.replace("\r\n", "\n"), "atualizado_em": _agora(),
                                        "autor": autor}
        if origem:
            dados["observacoes"][semana]["origem"] = origem
        salvar(pasta_trabalho, dados)


# ── o plano publicado vai para a reserva (09/10/2026) ───────────────────────────────────────────────────────
HIST_COLUNAS = ["task_key", "first_week", "last_week", "count", "weeks"]


def registrar_plano_publicado(pasta_trabalho: Path, semana: str, chaves, origem: str) -> dict:
    """Põe o plano que foi ao campo na reserva do histórico (o formato do motor): cada chave do plano ganha a semana, e
    a chave que estava nela e saiu (a semana publicada de novo) perde. É o "plano publicado" que o histórico usa para a
    semana que o banco ainda não fechou (`historico_banco.planos_publicados`): sem ele, a W43 gerada durante a W42
    contaria a W42 pelo retrato do meio da semana, sem as tarefas que rolaram (na W41 foram 247). Devolve o resumo."""
    chaves = {str(k) for k in chaves if k}
    with _trava:
        dados = carregar(pasta_trabalho)
        tab = dados.get("historico") or {"aba": None, "preambulo": [], "colunas": list(HIST_COLUNAS), "linhas": []}
        ix = {c: i for i, c in enumerate(tab.get("colunas") or [])}
        if "task_key" not in ix or "weeks" not in ix:
            raise InsumoErro("a reserva do histórico não tem as colunas task_key e weeks")
        por_chave = {}
        for r in tab.get("linhas") or []:
            if len(r) > max(ix["task_key"], ix["weeks"]) and r[ix["task_key"]]:
                ws = {w.strip() for w in str(r[ix["weeks"]] or "").split(",") if w.strip()}
                por_chave.setdefault(str(r[ix["task_key"]]), set()).update(ws)
        antes = sum(1 for ws in por_chave.values() if semana in ws)
        for ws in por_chave.values():
            ws.discard(semana)
        for k in chaves:
            por_chave.setdefault(k, set()).add(semana)
        linhas = []
        for k in sorted(por_chave):
            ws = sorted(por_chave[k])
            if ws:
                linhas.append([k, ws[0], ws[-1], len(ws), ",".join(ws)])
        tab.update(colunas=list(HIST_COLUNAS), linhas=linhas)
        planos = (tab.setdefault("meta", {})).setdefault("planos_publicados", {})
        planos[semana] = {"origem": origem, "chaves": len(chaves), "antes": antes, "registrado_em": _agora()}
        dados["historico"] = tab
        salvar(pasta_trabalho, dados)
    return {"semana": semana, "chaves": len(chaves), "antes": antes}


# ── feriados do ano seguinte (09/10/2026: a W53 de 2026 vai de 28/12 a 01/01/2027) ─────────────────────────────
def pascoa(ano: int) -> date:
    """O domingo de Páscoa (algoritmo de Meeus/Jones/Butcher, calendário gregoriano)."""
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l_ = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l_) // 451
    mes, dia = divmod(h + l_ - 7 * m + 114, 31)
    return date(ano, mes, dia + 1)


# Os feriados que andam com a Páscoa, pelos dias depois dela: Carnaval (-47), Paixão (-2), Corpus Christi (+60). O
# estadual do ES vem sem nome na planilha ("FERIADO ESTADUAL"): é Nossa Senhora da Penha, 8 dias depois da Páscoa
# (13/04/2026). Os municipais ficam na mesma data: a planilha não diz qual seria móvel.
MOVEIS_NACIONAIS = (-47, -2, 60)
MOVEIS_ESTADUAIS = {"ES": (8,)}


def _data(v):
    v = _de_json(v)
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def projetar_feriados(fer: dict, ano: int) -> dict:
    """Os feriados de `ano` projetados dos de `ano - 1`: data fixa repetida; móvel (MOVEIS_*) pela Páscoa de `ano`.
    Devolve {"gerais": [...], "municipais": [...]} no mesmo formato da planilha."""
    base, nova = pascoa(ano - 1), pascoa(ano)
    out = {"gerais": [], "municipais": []}
    for bloco in ("gerais", "municipais"):
        for r in fer.get(bloco) or []:
            d = _data(r.get("data"))
            if not d or d.year != ano - 1:
                continue
            tipo, uf = str(r.get("tipo") or "").strip().upper(), str(r.get("estado") or "").strip().upper()
            desvio = (d - base).days
            movel = ((tipo == "NACIONAL" and desvio in MOVEIS_NACIONAIS)
                     or (tipo == "ESTADUAL" and desvio in MOVEIS_ESTADUAIS.get(uf, ())))
            if movel:
                n = nova + timedelta(days=desvio)
            else:
                try:
                    n = d.replace(year=ano)
                except ValueError:           # 29/02 sem ano bissexto
                    continue
            novo = dict(r, data=_para_json(datetime.combine(n, datetime.min.time())), mes=n.month)
            out[bloco].append(novo)
    return out


def com_feriados_ate(fer: dict, ano_final: int) -> tuple[dict, list[int]]:
    """Os feriados com os anos que faltam até `ano_final` projetados (o ano que a planilha já tem não se projeta: o dado
    de verdade vence). Devolve (feriados, anos projetados)."""
    anos = {d.year for b in ("gerais", "municipais") for d in (_data(r.get("data")) for r in fer.get(b) or []) if d}
    if not anos:
        return fer, []
    out = {**fer, "gerais": list(fer.get("gerais") or []), "municipais": list(fer.get("municipais") or [])}
    projetados = []
    ano = max(anos)
    while ano < ano_final:
        p = projetar_feriados(out, ano + 1)
        out["gerais"] += p["gerais"]
        out["municipais"] += p["municipais"]
        ano += 1
        projetados.append(ano)
    return out, projetados


def observacoes(pasta_trabalho: Path, semana: str) -> str:
    return (carregar(pasta_trabalho)["observacoes"].get(semana) or {}).get("texto", "")


def observacoes_efetivas(pasta_trabalho: Path, semana: str) -> tuple[str, str | None]:
    """(o texto que vale para a semana, a semana de onde ele foi herdado ou None).

    O padrão é a última programação (05/10/2026, Levi: "quero como padrão marcado o que já estava marcado na última
    programação semanal"). Semana sem observação salva herda, da semana salva mais recente antes dela, só os dias por
    usina (`@usina`): é restrição de acesso da usina e vale até alguém mudar, como no arquivo único do PC do PCM. OS
    fora e OS com dia fixo são decisão da semana e não passam. Semana salva, mesmo vazia, vale o que foi salvo.
    A tela e o motor (`materializar`) leem daqui: o motor recebe os mesmos dias que a tela mostra."""
    obs = carregar(pasta_trabalho)["observacoes"]
    if semana in obs:
        return (obs[semana] or {}).get("texto", ""), None
    antes = sorted(s for s in obs if s < semana)
    if not antes:
        return "", None
    de = antes[-1]
    linhas = [l.strip() for l in (obs[de] or {}).get("texto", "").splitlines() if l.strip().lower().startswith("@usina")]
    return ("\n".join(linhas) + "\n", de) if linhas else ("", None)


# ── escrita para o motor (materializar) ─────────────────────────────────────────────────────────────────────

def escrever_planilha(tab: dict, destino: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab["aba"] or "Sheet1"
    for r in tab["preambulo"] + [tab["colunas"]] + tab["linhas"]:
        ws.append([_de_json(c) for c in r])
    wb.save(destino)


def escrever_feriados(fer: dict, destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = FERIADOS_ABA
    ws.append([fer.get("titulo")])
    ws.append(FER_ESQ + [None] + FER_DIR)
    g, m = fer["gerais"], fer["municipais"]
    for i in range(max(len(g), len(m))):
        esq = [_de_json(g[i].get(k)) for k in ("tipo", "estado", "data", "mes", "feriado")] if i < len(g) else [None] * 5
        dir_ = ([_de_json(m[i].get(k)) for k in ("tipo", "estado", "municipio", "data", "mes", "feriado")]
                if i < len(m) else [None] * 6)
        ws.append(esq + [None] + dir_)
    wb.save(destino)


def materializar(pasta_trabalho: Path, pasta_rodada: Path, semana: str) -> list[dict]:
    """Escreve na pasta da rodada os arquivos que o motor lê, a partir do que está no Nexus. Devolve o carimbo de
    cada um (de quando é o dado), para a rodada registrar de onde a semana partiu."""
    dados = carregar(pasta_trabalho)
    falta = [NOMES[k] for k in OBRIGATORIOS if k not in dados]
    if falta:
        raise InsumoErro("falta no Nexus: " + ", ".join(falta) + ". Importe da pasta do PCM.")
    carimbos = []
    for chave, (arq, _aba, _cab) in PLANILHAS.items():
        if chave == "historico":
            # a reserva sai com outro nome e sem carimbo: o histórico que o motor lê sai do banco na geração
            # (`historico_banco`), e esta cópia só vale se o banco não responder (aí a geração a volta para antes da
            # semana e carimba de onde veio)
            if dados.get(chave):
                escrever_planilha(dados[chave], pasta_rodada / "historico_do_nexus.xlsx")
            continue
        escrever_planilha(dados[chave], pasta_rodada / arq)
        carimbos.append(_carimbo(chave, dados[chave]))
    # o ano seguinte ao da semana vai projetado (a W53 de 2026 termina em 01/01/2027; a planilha do PCM só tem 2026)
    fim_semana = date.fromisocalendar(int(semana[:4]), int(semana[6:]), 5)
    fer, projetados = com_feriados_ate(dados["feriados"], fim_semana.year + 1)
    escrever_feriados(fer, pasta_rodada / FERIADOS_PASTA / FERIADOS_ARQ)
    carimbo = _carimbo("feriados", dados["feriados"])
    if projetados:
        orig = dados["feriados"]
        n = len(fer["gerais"]) + len(fer["municipais"]) - len(orig["gerais"]) - len(orig["municipais"])
        carimbo["detalhe"] += (f"; {', '.join(map(str, projetados))} projetado ({n} datas: as fixas repetidas; "
                               "Carnaval, Paixão, Corpus Christi e o estadual do ES pela Páscoa)")
    carimbos.append(carimbo)
    texto, herdada_de = observacoes_efetivas(pasta_trabalho, semana)
    obs = dados["observacoes"].get(herdada_de or semana) or {}
    (pasta_rodada / OBSERVACOES).write_text(texto, encoding="utf-8")
    carimbos.append({"nome": NOMES["observacoes"], "atualizado": (obs.get("atualizado_em") or "")[:16],
                     "detalhe": ("dias por usina herdados da semana " + herdada_de) if herdada_de else
                     ("vazias" if not texto.strip() else "da semana " + semana)})
    return carimbos


def _carimbo(chave: str, item: dict) -> dict:
    m = item.get("meta") or {}
    return {"nome": NOMES[chave], "sha": m.get("sha", ""), "atualizado": (m.get("importado_em") or "")[:16],
            "detalhe": "importado de " + m.get("arquivo", "") if m.get("arquivo") else "editado no Nexus"}


# ── estado para a tela ─────────────────────────────────────────────────────────────────────────────────────

def estado(pasta_trabalho: Path, origem: Path | None, semana: str) -> list[dict]:
    """Um item por insumo: se está no Nexus, de quando, e se o arquivo da pasta do PCM mudou depois da importação
    (na sombra, o Nexus precisa partir do mesmo dado que o PCM)."""
    dados = carregar(pasta_trabalho)
    itens = []
    for chave in ("prioridades", "confiabilidade", "historico", "feriados"):
        item = dados.get(chave)
        if chave == "historico":
            itens.append(_estado_historico(item))
            continue
        if not item:
            itens.append({"nome": NOMES[chave], "ok": False, "obrigatorio": True, "aviso": False,
                          "detalhe": "ainda não está no Nexus: importe da pasta do PCM"})
            continue
        m = item.get("meta") or {}
        if chave == "feriados":
            qtd = f"{len(item['gerais'])} nacionais e estaduais, {len(item['municipais'])} municipais"
            arq = (Path(origem) / FERIADOS_PASTA / FERIADOS_ARQ) if origem else None
        else:
            qtd = f"{len(item['linhas'])} linhas"
            arq = (Path(origem) / PLANILHAS[chave][0]) if origem else None
        mudou = bool(arq and arq.exists() and m.get("sha") and _sha(arq) != m["sha"])
        detalhe = f"no Nexus, {qtd}, importado em {_curto(m.get('importado_em'))}"
        if mudou:
            detalhe += f". O arquivo da pasta mudou em {_curto(_quando_arquivo(arq))}: importe de novo"
        itens.append({"nome": NOMES[chave], "ok": True, "obrigatorio": True, "aviso": mudou, "detalhe": detalhe})
    obs = dados["observacoes"].get(semana)
    texto, herdada_de = observacoes_efetivas(pasta_trabalho, semana)
    n = len([l for l in texto.splitlines() if l.strip() and not l.strip().startswith("#")])
    if herdada_de:
        detalhe = f"{n} {'usina' if n == 1 else 'usinas'} com dia, herdadas da {herdada_de} (ainda não salvas nesta semana)"
    else:
        detalhe = f"{n} para a semana {semana}" if obs else f"nenhuma para a semana {semana}"
    itens.append({"nome": NOMES["observacoes"], "ok": True, "obrigatorio": False, "aviso": False, "detalhe": detalhe})
    return itens


def _estado_historico(item: dict | None) -> dict:
    """O histórico sai do banco na hora de gerar. A tela não lê o banco a cada abertura (o fato tem ~14 mil linhas, ~4 s
    de GET): diz de onde vem e o que a reserva guardada aqui tem. Sem reserva não é falta: é só a geração sem rede de
    proteção (ela para, com o motivo, se o banco também não responder)."""
    base = "do banco a cada geração (nexus_programacao · fato_programacao)"
    if not item:
        return {"nome": NOMES["historico"], "ok": True, "obrigatorio": False, "aviso": False,
                "detalhe": base + "; sem reserva no Nexus"}
    ix = {c: i for i, c in enumerate(item.get("colunas") or [])}
    i_w = ix.get("weeks", 4)
    semanas = sorted({w.strip() for r in item.get("linhas") or [] if len(r) > i_w
                      for w in str(r[i_w] or "").split(",") if w.strip()})
    ate = f", até a {semanas[-1]}" if semanas else ""
    n = f"{len(item.get('linhas') or []):,}".replace(",", ".")
    return {"nome": NOMES["historico"], "ok": True, "obrigatorio": False, "aviso": False,
            "detalhe": f"{base}; reserva no Nexus: {n} linhas{ate}"}


def _curto(iso: str | None) -> str:
    return f"{iso[8:10]}/{iso[5:7]} {iso[11:16]}" if iso and len(iso) >= 16 else ""
