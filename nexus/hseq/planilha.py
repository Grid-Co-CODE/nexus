"""O relatório dos extintores em Excel (Segurança · HSEQ).

A TST, 19/10/2026, depois do relatório por cliente em PDF: "e excel tbm? pq às vezes o cliente pede". O mesmo recorte do
PDF (a rota `extintores_planilha` usa o `_do_relatorio` da torre: os filtros da tela, a usina ou o cliente da linha), em
planilha para o cliente filtrar e somar:
- aba "Extintores": UMA linha por extintor, na ordem da tabela da tela (o cliente, a usina e o mais grave primeiro:
  `extintores.por_cliente`), com as mesmas contas (situação, sem atualização, status da TST). A conferência vai como
  data de verdade (o Excel ordena e filtra por ela); as validades vão como a etiqueta diz ("08/2026", "2027", "Sem
  data"), porque a etiqueta é o dado: o dia em que ela vence também vai, à parte, como data;
- aba "Não reconhecidos": os que a equipe não achou na usina, quando o recorte os leva (a regra do PDF);
- em cima da tabela, o título, quando foi gerado e os filtros, como no PDF.
Sem cor de fundo nem fórmula: o cliente abre em qualquer programa de planilha. A pessoa vai pelo nome (como no PDF),
nunca pelo HMAC.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import extintores as EXT

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATA = "DD/MM/YYYY"
CARGA = {"CARREGADO": "carregado", "SEM_CARGA": "sem carga", "SOBRECARGA": "sobrecarga"}
# (cabeçalho, largura, valor) de cada coluna da aba Extintores
COLUNAS = (
    ("Cliente", 18, lambda x, u: u["cliente"]),
    ("Usina", 28, lambda x, u: x["usina"]),
    ("Responsável", 22, lambda x, u: "" if u["responsavel"] == EXT.SEM_VALOR["responsavel"] else u["responsavel"]),
    ("Técnicos", 30, lambda x, u: ", ".join(u["tecnicos"])),
    ("Extintor", 10, lambda x, u: x["ext"]),
    ("Código", 28, lambda x, u: x["codigo"]),
    ("Classe", 12, lambda x, u: x.get("classe") or ""),
    ("Peso (kg)", 10, lambda x, u: x.get("peso")),
    ("Local", 20, lambda x, u: x.get("local") or ""),
    ("Posição", 12, lambda x, u: (x.get("posicao") or "").lower()),
    ("Ativo no Fracttal", 24, lambda x, u: x.get("ativo") or ""),
    ("Situação", 16, lambda x, u: EXT.SITUACOES[x["situacao"]][0]),
    ("Validade da recarga", 12, lambda x, u: x["val2"]),
    ("Recarga vence em", 12, lambda x, u: x.get("fim2")),
    ("Validade do hidrostático", 12, lambda x, u: x["val3"]),
    ("Hidrostático vence em", 12, lambda x, u: x.get("fim3")),
    ("Última conferência", 12, lambda x, u: x["conferencia"]),
    ("Sem atualização há mais de 30 dias", 12, lambda x, u: "sim" if x["sem_atualizacao"] else "não"),
    ("Origem da conferência", 14, lambda x, u: x.get("origem_conferencia") or ""),
    ("Conferido por", 22, lambda x, u: x.get("conferido_por") or ""),
    ("Carga", 12, lambda x, u: CARGA.get(x.get("carga") or "", "")),
    ("Itens com NÃO", 22, lambda x, u: ", ".join(x.get("itens_nao") or [])),
    ("Status da TST", 14, lambda x, u: (EXT.STATUS_TST.get(x.get("status")) or ("sem conferência",))[0]),
    ("Motivo do status", 40, lambda x, u: x.get("motivo") or ""),
)
COLUNAS_NR = (
    ("Cliente", 18, lambda x: EXT.valor_do_filtro(x, "cliente")),
    ("Usina", 28, lambda x: x["usina"]),
    ("Extintor", 10, lambda x: x["ext"]),
    ("Código", 28, lambda x: x["codigo"]),
    ("Local", 20, lambda x: x.get("local") or ""),
    ("Não reconhecido em", 12, lambda x: x["nao_reconhecido"]),
    ("Marcado por", 22, lambda x: x.get("nao_reconhecido_por") or ""),
    ("Conferência antes", 12, lambda x: x["conferencia"]),
)


def _aba(ws, titulo: str, topo: list[str], colunas, linhas) -> None:
    """O título e as linhas de cima, o cabeçalho (negrito, congelado, com o filtro do Excel) e as linhas."""
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=14)
    for t in topo:
        ws.append([t])
    ws.append([])
    cab = len(topo) + 3
    ws.append([c[0] for c in colunas])
    for i, (_rot, larg, _f) in enumerate(colunas, 1):
        cel = ws.cell(row=cab, column=i)
        cel.font = Font(bold=True)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = larg
    for valores in linhas:
        ws.append(valores)
        for cel in ws[ws.max_row]:
            if isinstance(cel.value, datetime):
                cel.number_format = DATA + " HH:MM"
            elif isinstance(cel.value, date):
                cel.number_format = DATA
    ws.freeze_panes = ws.cell(row=cab + 1, column=1)
    ws.auto_filter.ref = f"A{cab}:{get_column_letter(len(colunas))}{max(ws.max_row, cab)}"


def gerar(lista, *, filtros_tela: list[str], hoje: date, gerado_em: str, nao_reconhecidos=()) -> bytes:
    """O .xlsx: a aba Extintores e, se o recorte os leva, a aba Não reconhecidos. `lista` = os extintores já
    filtrados (os da tabela da tela)."""
    topo = [f"Gerado em {gerado_em} · situação no dia {hoje:%d/%m/%Y}",
            "Filtros: " + (" · ".join(filtros_tela) if filtros_tela else "nenhum (todos os extintores)")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Extintores"
    linhas = [[f(x, u) for _r, _l, f in COLUNAS]
              for c in EXT.por_cliente(lista) for u in c["usinas"] for x in u["extintores"]]
    _aba(ws, "Relatório de extintores", topo + [f"{len(linhas)} extintor{'es' if len(linhas) != 1 else ''}"],
         COLUNAS, linhas)
    if nao_reconhecidos:
        nr = wb.create_sheet("Não reconhecidos")
        _aba(nr, "Extintores não reconhecidos em campo",
             topo + ["A equipe não achou estes extintores na usina; ficam fora das contas até a TST conferir o cadastro."],
             COLUNAS_NR, [[f(x) for _r, _l, f in COLUNAS_NR] for x in nao_reconhecidos])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── a planilha de todas as pendências da Mesa (a TST, 24/10/2026: "um resumo com todas as pendências (...) e ter a opção
#    de baixar") ──────────────────────────────────────────────────────────────────────────────────────────────────────
def _sem(x, campo):
    """O valor do grupo na planilha: o "Sem ..." da tela vira vazio (o filtro do Excel junta os vazios)."""
    v = EXT.valor_do_filtro(x, campo)
    return "" if v == EXT.SEM_VALOR.get(campo) else v


def _hora(dt):
    """Data e hora de Brasília sem o fuso (o Excel não guarda fuso)."""
    return dt.replace(tzinfo=None) if dt else None


def _vence(x):
    """O vencimento que pesa: a recarga ou o hidrostático, o que vence (ou venceu) antes."""
    datas = [d for d in (x.get("fim2"), x.get("fim3")) if d]
    return min(datas) if datas else None


_ONDE = (("Cliente", 18, lambda x: _sem(x, "cliente")), ("Usina", 26, lambda x: x.get("usina") or ""),
         ("Responsável UFV", 22, lambda x: _sem(x, "responsavel")))
_EXT = _ONDE + (("Extintor", 10, lambda x: x["ext"]), ("Código", 28, lambda x: x["codigo"]),
                ("Local", 20, lambda x: x.get("local") or ""), ("Validade da recarga", 12, lambda x: x["val2"]),
                ("Validade do hidrostático", 12, lambda x: x["val3"]), ("Vence em", 12, _vence),
                ("Última conferência", 12, lambda x: x["conferencia"]))
_LUVAS = _ONDE + (("Equipe", 16, lambda x: x.get("equipe") or ""),
                  ("Última verificação", 12, lambda x: (x.get("ultima") or {}).get("dia")),
                  ("Há quantos dias", 10, lambda x: x.get("dias")))
ABAS_PENDENCIAS = {
    "pts": ("PTs paradas", "PTs esperando o De acordo há mais de 2 h", _ONDE[:2] + (
        ("Responsável UFV", 22, lambda x: _sem(x, "responsavel")), ("PT", 10, lambda x: x["numero"]),
        ("OS", 10, lambda x: x.get("os") or ""), ("Tarefa", 30, lambda x: x.get("tarefa") or ""),
        ("Ativo", 22, lambda x: x.get("ativo") or ""), ("Pedida por", 22, lambda x: x.get("solicitante") or ""),
        ("Criada em", 16, lambda x: _hora(x.get("criada"))),
        ("Esperando há (h)", 10, lambda x: round((x.get("idade_min") or 0) / 60, 1)))),
    "vencidos": ("Extintores vencidos", "A recarga ou o teste hidrostático venceu", _EXT),
    "nr": ("Não reconhecidos", "A equipe não achou o extintor na usina", _ONDE + (
        ("Extintor", 10, lambda x: x["ext"]), ("Código", 28, lambda x: x["codigo"]),
        ("Local", 20, lambda x: x.get("local") or ""), ("Não reconhecido em", 12, lambda x: x["nao_reconhecido"]),
        ("Marcado por", 22, lambda x: x.get("nao_reconhecido_por") or ""))),
    "sem_luvas": ("Usinas sem luvas", "A última ronda diária respondeu Não à pergunta das luvas isolantes", _LUVAS + (
        ("Sem luvas desde", 12, lambda x: x.get("desde")),)),
    "perto": ("Perto de vencer", "A recarga ou o hidrostático vence nos próximos 30 dias", _EXT),
    "sem_atual": ("Sem atualização", "Conferido há mais de 30 dias, ou nunca (sem os vencidos e os perto de vencer)",
                  _EXT + (("Há quantos dias", 10, lambda x: x.get("dias_conferencia")),)),
    "luvas_semv": ("Luvas sem verificação", "Sem verificação das luvas nas rondas dos últimos 7 dias", _LUVAS),
    "ronda": ("Nova ronda pedida", "A Segurança pediu uma nova ronda de extintores, do zero, e a equipe ainda não a "
                                   "refez", _ONDE + (("Pedida em", 16, lambda x: _hora(_pedida(x))),
                                                     ("Extintores refeitos", 10, lambda x: x["feitos"]),
                                                     ("Extintores da usina", 10, lambda x: x["total"]))),
}


def _pedida(x):
    try:
        return datetime.fromisoformat(str(x.get("pedida_em") or ""))
    except ValueError:
        return None


def gerar_pendencias(listas: dict, cartoes: list, *, chave: str, filtros_tela: list[str], hoje: date, gerado_em: str,
                     erros: dict) -> bytes:
    """O .xlsx da visão geral: a aba Resumo (uma linha por responsável ou cliente, os quatro de agora e os de atenção)
    e uma aba por pendência, uma linha por caso, com os mesmos filtros da tela. A fonte que não foi lida diz no topo."""
    from ..torres.hseq import AGIR, ATENCAO
    topo = [f"Gerado em {gerado_em} · situação no dia {hoje:%d/%m/%Y}",
            "Filtros: " + (" · ".join(filtros_tela) if filtros_tela else "nenhum (todas as pendências)")]
    falhou = {"pts": "as PT", "ext": "os extintores", "epi": "as rondas das luvas"}
    topo += [f"Não consegui ler {falhou[k]} agora: os números delas ficam fora" for k in falhou if erros.get(k)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    grupo = "Responsável UFV" if chave == "responsavel" else "Cliente"
    cols = ((grupo, 26, None),) + tuple((rot, 12, None) for _k, rot in AGIR + ATENCAO)
    _aba(ws, "Pendências da Segurança", topo + ["Os quatro primeiros pedem ação agora; os outros, atenção."], cols,
         [[c["nome"]] + [c[k] for k, _rot in AGIR + ATENCAO] for c in cartoes])
    for k, (titulo, explica, colunas) in ABAS_PENDENCIAS.items():
        aba = wb.create_sheet(titulo)
        linhas = [[f(x) for _r, _l, f in colunas] for x in listas.get(k) or []]
        _aba(aba, titulo, topo + [explica, f"{len(linhas)} {'caso' if len(linhas) == 1 else 'casos'}"], colunas, linhas)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

