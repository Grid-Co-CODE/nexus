"""Extintores no Nexus (Segurança · HSEQ, Levi, 09/10/2026): "ver a tabela dos extintores com uma visão do que está
atrasado, perto de atrasar e sem atualização a mais de 30 dias. Preciso da visão por supervisor".

As regras (`nexus/hseq/extintores.py`) num dia fixo, 09/10/2026, e a tela com o banco falso do Campo
(test_campo_visao.py): Altair (pelo de-para do Fracttal) e Brodowski 1 (pelo código do ativo) são da SP Norte 01, região
Sudeste 03 (Supervisora Campo), gestor Beltrano Supervisor; Coração 1 (pelo código da usina) é da Sul 01, com a vaga de
supervisor aberta, gestor Ciclano Chefe; Irecê 2 não está no cadastro (sem região e sem gestor).

Desde 18/10/2026 (Levi: a tabela "agrupado por Cliente", "a coluna do responsável desse usina e o ou os técnicos da
usina", os filtros "cliente, usina, responsável UFV, equipe e o campo para digitar" e o PDF só unitário), o
`com_cadastro` põe o Responsável O&M nas fichas (Altair e Brodowski 1: Supervisora Campo; Coração 1: Coordenador Campo)
e os técnicos da tela Técnicos por usina do PCM (Altair: dois; Brodowski 1: só um apoio agendado; Coração 1: um)."""
import json
import re
from datetime import date, datetime

import pytest
from test_campo_visao import BRT, CHAVE, CHAVE_CADASTRO, USINAS, _aba, _pessoas, banco  # noqa: F401  (banco é fixture)

from nexus.cadastro.cifra import Cofre
from nexus.campo import visao
from nexus.campo.ligacao_cadastro import codigo_da_pessoa
from nexus.hseq import extintores as EXT
from nexus.pcm import atribuicoes as ATR

HOJE = date(2026, 10, 9)


def _ext(codigo, usina, val2, val3, conf, ativo="", cb="", carga="CARREGADO", nao=None, hmac=None, **kw):
    """Uma linha do livro `extintores_app_campo · Extintores`, no contrato com o App (`EXT.COLUNAS`)."""
    d = {"Código": codigo, "Usina": usina, "Código da usina": cb, "Tipo de ativo": "CABN", "Ativo": ativo,
         "Local": "Cabine", "Posição": "Interno", "Classe": "ABC pó", "Peso (kg)": 6, "Validade da recarga": val2,
         "Validade do hidrostático": val3, "Última conferência": conf,
         "Origem da conferência": "App" if conf else None, "Conferido por (HMAC)": hmac,
         "Carga": carga if conf else None, "Itens NÃO": nao, "Fotos": 1 if conf else None,
         "Tem informação adicional": "não", "Status": None, "Motivo": None, "Origem do cadastro": "Planilha da TST"}
    d.update(kw)
    return d


LIVRO = [
    # Altair: vencido há 39 dias; vence em 22 dias e conferido há 69; em dia (o NÃO no difusor só pesa no status da TST)
    _ext("ALT100-INFC1-PPCI-EXT01", "Thopen - Altair 1 - SP", "08/2026", "Sem selo", "2026-10-05",
         hmac=codigo_da_pessoa(CHAVE, "tec1@exemplo.test")),
    _ext("ALT100-INFC1-PPCI-EXT02", "Thopen - Altair 1 - SP", "10/2026", "Sem selo", "2026-08-01"),
    _ext("ALT100-INFC1-PPCI-EXT03", "Thopen - Altair 1 - SP", "12/2027", "2026", "2026-10-01", nao="difusor"),
    # Brodowski 1, pelo código do ativo: sem data na recarga e nunca conferido
    _ext("BWK100-INFC1-PPCI-EXT01", "Thopen - Brodowski 1 - SP", "Sem data", "Sem selo", None, ativo="BWK100-SKID1"),
    # Coração 1, pelo código da usina: o hidrostático venceu (a recarga não) e está sem carga
    _ext("COR100-INFC1-PPCI-EXT01", "Thopen - Coração 1 - SC", "03/2027", "09/2026", "2026-03-01", cb="COR100",
         carga="SEM_CARGA"),
    # fora do cadastro do Nexus
    _ext("IRC200-INFC1-PPCI-EXT01", "Thopen - Irecê 2 - BA", "11/2026", "Sem selo", "2026-10-08", cb="IRC200"),
]


@pytest.fixture
def com_livro(banco, monkeypatch):
    monkeypatch.setattr(visao, "_agora", lambda: datetime(2026, 10, 9, 10, 0, tzinfo=BRT))
    _aba(banco, EXT.LIVRO, EXT.ABA, LIVRO)
    visao.limpar()
    return banco


# os técnicos de campo da SP Norte 01 (1, 2) e da SC Oeste 01 (4) ganham nome (fictício), cifrado como no banco
NOMES = {1: "Técnico Alfa", 2: "Eletricista Beta", 4: "Técnico Gama"}
RESPONSAVEL = {1: 92, 2: 92, 3: 93}          # Altair e Brodowski 1: Supervisora Campo; Coração 1: Coordenador Campo


def _atribuicao(usina_id, pessoa_id, de=""):
    return {"decidida_em": "2026-10-01T10:00:00-03:00", "data_id": 20261001, "usina_id": usina_id,
            "usina_fracttal": "", "pessoa_id": pessoa_id, "pessoa_hmac": "", "papel": "executa", "na_usina": "sim",
            "de_data_id": de, "ate_data_id": "", "origem": "tela", "decidida_por_hmac": "", "dados_cifrados": "",
            "assinatura": ""}


@pytest.fixture
def com_cadastro(com_livro, app, tmp_path):
    """O Responsável O&M nas fichas das usinas e os técnicos da tela Técnicos por usina do PCM (o livro como arquivo,
    `NEXUS_PCM_ATRIBUICOES_ARQUIVO`): Altair com dois, Brodowski 1 só com um apoio que começa em 20/10 (ainda não
    conta) e Coração 1 com um."""
    cofre = Cofre(CHAVE_CADASTRO)
    pessoas = [dict(p, sensivel_cifrado=cofre.cifrar(json.dumps({"nome": NOMES[p["pessoa_id"]],
                                                                 "nome_padrao": NOMES[p["pessoa_id"]]}),
                                                     f"banco/pessoas/{p['pessoa_id']}"))
               if p["pessoa_id"] in NOMES else p for p in _pessoas()]
    _aba(com_livro, "cadastro_nexus", "pessoas", pessoas)
    _aba(com_livro, "cadastro_nexus", "usinas", [dict(u, responsavel_om_id=RESPONSAVEL.get(u["usina_id"]))
                                                 for u in USINAS])
    arq = tmp_path / "atribuicoes.json"
    arq.write_text(json.dumps({ATR.ABA: [_atribuicao(1, 1), _atribuicao(1, 2), _atribuicao(2, 2, de=20261020),
                                         _atribuicao(3, 4)]}), encoding="utf-8")
    app.config["NEXUS_PCM_ATRIBUICOES_ARQUIVO"] = str(arq)
    ATR.limpar()
    visao.limpar()
    yield com_livro
    ATR.limpar()


def _texto(html):
    return " ".join(re.sub(r"<[^>]+>", " ", html).split())


def _linhas(html):
    """Os códigos dos extintores na tabela, na ordem."""
    return re.findall(r'title="([A-Z]{3}\d{3}-INFC1-PPCI-EXT\d+)"', html)


# ── as regras ─────────────────────────────────────────────────────────────────────────────────────────────────────
def test_a_etiqueta_vale_ate_o_ultimo_dia_do_mes_e_so_o_ano_ate_31_12():
    assert EXT.validade_fim("02/2028") == date(2028, 2, 29) and EXT.validade_fim("9/2026") == date(2026, 9, 30)
    assert EXT.validade_fim("2027") == EXT.validade_fim("2027.0") == date(2027, 12, 31)       # a API não garante o tipo
    for v in ("Sem data", "Sem selo", "", None, "13/2026", "00/2026", "12/1999", "31/12/2026"):
        assert EXT.validade_fim(v) is None, v


def test_status_da_tst_e_a_regra_do_app_inclusive_na_virada_do_mes():
    """A grade de casos do App (test_v247_extintores_servidor.py roda o `_ext_status` dele nos mesmos)."""
    limpo = set()
    assert EXT.status_tst("CARREGADO", limpo, "12/2027", "Sem selo", HOJE) == (
        "OK", "Nada fora do lugar e validade em dia")
    assert EXT.status_tst("CARREGADO", limpo, "Sem data", "Sem selo", HOJE) == (
        "OK_SEM_VALIDADE", "Checklist limpo, mas falta a validade")
    assert EXT.status_tst("SOBRECARGA", limpo, "12/2027", "", HOJE) == ("CRITICO", "sobrecarga")
    assert EXT.status_tst("Sem carga", limpo, "12/2027", "", HOJE) == ("CRITICO", "sem carga")
    assert EXT.status_tst("CARREGADO", {"manometro", "pintura"}, "12/2027", "", HOJE) == ("CRITICO", "manômetro: Não")
    assert EXT.status_tst("CARREGADO", {"pintura"}, "12/2027", "", HOJE) == ("ATENCAO", "pintura: Não")
    assert EXT.status_tst("CARREGADO", limpo, "10/2026", "", HOJE) == ("ATENCAO", "recarga vence em 31/10")
    # as DUAS etiquetas: o hidrostático vencido deixa crítico mesmo com a recarga em dia (o caso de Marabá 2)
    assert EXT.status_tst("CARREGADO", limpo, "12/2027", "12/2025", HOJE) == (
        "CRITICO", "teste hidrostático vencido em 12/2025")
    # a virada: vence hoje (0 dia) ainda é atenção; no dia seguinte, crítico
    assert EXT.status_tst("CARREGADO", limpo, "09/2026", "", date(2026, 9, 30))[0] == "ATENCAO"
    assert EXT.status_tst("CARREGADO", limpo, "09/2026", "", date(2026, 10, 1)) == (
        "CRITICO", "recarga vencida em 09/2026")
    assert EXT.status_tst("CARREGADO", limpo, "12/2027", "2026", date(2026, 12, 1)) == (
        "ATENCAO", "teste hidrostático vence em 31/12")


def test_o_tempo_vai_em_dias_ate_59_e_depois_em_meses_e_anos():
    """Na 1ª conferência com o cadastro da TST a tabela dizia "venceu há 557 d" e "vence em 1179 d"."""
    assert [EXT.tempo(n) for n in (0, 22, -39, 59, 60, 161, 365, 557, 715, 730, 1179)] == [
        "0 d", "22 d", "39 d", "59 d", "2 meses", "5 meses", "12 meses", "18 meses", "23 meses", "2 anos", "3 anos"]


def test_os_itens_nao_aceitam_a_chave_e_o_rotulo_com_acento():
    assert EXT.itens_nao("manometro; difusor") == {"manometro", "difusor"}
    assert EXT.itens_nao("Manômetro, Sinalização") == {"manometro", "sinalizacao"}
    assert EXT.itens_nao(None) == set()


def test_situacoes_exclusivas_e_a_sem_atualizacao_a_parte():
    s = {x["Código"][-5:] + x["Código"][:3]: EXT.situacao(x, HOJE) for x in LIVRO}
    alt1, alt2, alt3 = s["EXT01ALT"], s["EXT02ALT"], s["EXT03ALT"]
    assert (alt1["situacao"], alt1["dias_vence"], alt1["vence_qual"], alt1["sem_atualizacao"]) == (
        "atrasado", -39, "recarga", False)
    assert (alt2["situacao"], alt2["dias_vence"], alt2["sem_atualizacao"], alt2["dias_conferencia"]) == (
        "perto", 22, True, 69)
    # em dia: o NÃO no difusor deixa o status da TST em atenção, mas não mexe na validade
    assert (alt3["situacao"], alt3["status"], alt3["motivo"]) == ("em_dia", "ATENCAO", "difusor: Não")
    bwk = s["EXT01BWK"]
    assert (bwk["situacao"], bwk["sem_atualizacao"], bwk["conferencia"], bwk["status"]) == (
        "sem_validade", True, None, "")
    cor = s["EXT01COR"]
    assert (cor["situacao"], cor["vence_qual"], cor["dias_vence"]) == ("atrasado", "hidrostático", -9)
    assert cor["status"] == "CRITICO" and cor["motivo"] == "sem carga · teste hidrostático vencido em 09/2026"
    # o que a tabela mostra ao lado do status: só a carga e os itens (a validade já tem coluna)
    assert (cor["motivo_checklist"], alt3["motivo_checklist"], alt1["motivo_checklist"]) == ("sem carga", "difusor: Não", "")
    # 30 dias é o limite dos dois lados: vence em 30 dias = perto; conferido há 30 dias ainda não está sem atualização
    lim = EXT.situacao(_ext("X", "Y", "11/2026", "", "2026-09-09"), date(2026, 10, 31))
    assert (lim["situacao"], lim["dias_vence"], lim["dias_conferencia"], lim["sem_atualizacao"]) == (
        "perto", 30, 52, True)
    assert EXT.situacao(_ext("X", "Y", "12/2026", "", "2026-09-09"), date(2026, 10, 9))["sem_atualizacao"] is False
    assert EXT.situacao(_ext("X", "Y", "12/2026", "", "2026-09-08"), date(2026, 10, 9))["sem_atualizacao"] is True
    # a recarga sem data com o hidrostático vencido é atrasado, não "sem validade" (sabe-se que venceu)
    assert EXT.situacao(_ext("X", "Y", "Sem data", "01/2026", "2026-10-01"), HOJE)["situacao"] == "atrasado"


# ── a tela ────────────────────────────────────────────────────────────────────────────────────────────────────────
def test_sem_o_livro_a_tela_diz_que_os_extintores_ainda_nao_chegaram(banco, logado):
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "Os extintores ainda não chegaram ao banco do Nexus" in html and "extintores_app_campo" in html
    assert "cn-numeros" not in html and "<table>" not in html
    assert "Como cada extintor entra" in html and 'class="hs-legenda"' in html     # a explicação das regras fica


def test_a_faixa_conta_o_que_o_pedido_quer_e_cada_numero_filtra_a_tabela(com_livro, logado):
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    t = _texto(html)
    assert "Extintores 6 em 4 usinas" in t
    assert "Atrasados 2 recarga ou teste hidrostático vencido" in t
    assert "Perto de vencer 1 vencem em até 30 dias" in t
    assert "Sem atualização 3 conferido há mais de 30 dias, ou nunca" in t
    assert "Sem validade 1 recarga sem data na etiqueta" in t
    for f in EXT.FILTROS:
        assert f'href="?f={f}"' in html, f


def test_por_responsavel_um_cartao_por_responsavel_da_usina(com_cadastro, logado):
    """Os cartões seguiram o filtro novo (18/10/2026): um por Responsável O&M da ficha da usina, no lugar dos de região
    de campo e de gestor de contrato; o cartão leva à tabela filtrada por ele."""
    html = logado.get("/t/hseq/extintores?modo=responsaveis").get_data(as_text=True)
    assert 'aria-current="page">Por responsável</a>' in html and "<table" not in html.split('id="tabela"')[1]
    assert "Região de campo" not in html and "Gestor de contrato" not in html
    # Supervisora Campo: Altair e Brodowski 1 (4 extintores); antes do Coordenador Campo (mesmos atrasados, mais perto
    # de vencer); Irecê 2, fora do cadastro, no cartão "Sem responsável O&M", por último
    assert html.index("<b>Supervisora Campo</b>") < html.index("<b>Coordenador Campo</b>") < html.index(
        "<b>Sem responsável O&amp;M</b>")
    corpo = _texto(html[html.index("<b>Supervisora Campo</b>"):html.index("<b>Coordenador Campo</b>")])
    assert "Responsável UFV" in corpo and "2 usinas · 1 equipe" in corpo
    assert "1 extintor atrasado" in corpo and "1 vence em até 30 dias" in corpo
    assert "1 em dia 1 sem validade 4 total" in corpo and "2 sem atualização há mais de 30 dias" in corpo
    assert "usinas sem responsável no cadastro" in _texto(html[html.index("<b>Sem responsável O&amp;M</b>"):])
    assert 'href="?responsavel=Supervisora+Campo#tabela"' in html
    tab = logado.get("/t/hseq/extintores?responsavel=Supervisora+Campo&ver=extintor").get_data(as_text=True)
    assert sorted(_linhas(tab)) == ["ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02", "ALT100-INFC1-PPCI-EXT03",
                                    "BWK100-INFC1-PPCI-EXT01"]
    # o filtro da faixa sempre abre a tabela
    assert 'aria-current="page">Tabela</a>' in logado.get(
        "/t/hseq/extintores?modo=responsaveis&f=atrasado").get_data(as_text=True)


def test_a_tabela_poe_o_mais_grave_primeiro_e_filtra_pela_situacao(com_livro, logado):
    html = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor").get_data(as_text=True)
    # atrasado (o vencido há mais tempo primeiro), perto, sem validade, em dia (o que vence antes primeiro)
    assert _linhas(html) == ["ALT100-INFC1-PPCI-EXT01", "COR100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02",
                             "BWK100-INFC1-PPCI-EXT01", "IRC200-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT03"]
    t = _texto(html[html.index("<table>"):html.index("</table>")])
    assert "08/2026 venceu há 39 d" in t and "10/2026 vence em 22 d" in t and "nunca conferido" in t
    # o status não repete a validade (está nas colunas ao lado): só a carga e os itens; o motivo inteiro, no title
    assert "Crítico sem carga" in t and 'title="sem carga · teste hidrostático vencido em 09/2026"' in html
    assert "Atenção difusor: Não" in t
    # a usina vai sem a região ao lado (Levi, 11/10: "fica estranho"); a região está nos cartões e no filtro
    assert "Sudeste 03" not in t and "Supervisora Campo" not in t and "sem região" not in t
    assert "Thopen - Irecê 2 - BA" in t                                        # fora do cadastro: o nome do App
    cab = _texto(html[html.index("<thead>"):html.index("</thead>")])
    assert cab == "Situação Última conferência Usina Extintor Local Recarga Hidrostático Status da TST"
    assert "05/10/2026 há 4 d · App Técnico Silva" in t                          # quem conferiu, pelo nome
    assert codigo_da_pessoa(CHAVE, "tec1@exemplo.test") not in html
    atr = logado.get("/t/hseq/extintores?f=atrasado&ver=extintor").get_data(as_text=True)          # o filtro abre a tabela
    assert _linhas(atr) == ["ALT100-INFC1-PPCI-EXT01", "COR100-INFC1-PPCI-EXT01"]
    sem = logado.get("/t/hseq/extintores?f=sem_atualizacao&ver=extintor").get_data(as_text=True)    # o mais antigo primeiro
    assert _linhas(sem) == ["BWK100-INFC1-PPCI-EXT01", "COR100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02"]
    eq = logado.get("/t/hseq/extintores?ver=extintor&equipe=SC+Oeste+01").get_data(as_text=True)
    assert _linhas(eq) == ["COR100-INFC1-PPCI-EXT01"] and '<option value="SC Oeste 01" selected>' in eq
    busca = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor&q=skid1").get_data(as_text=True)
    assert _linhas(busca) == ["BWK100-INFC1-PPCI-EXT01"]


def test_so_os_cinco_filtros_em_cascata_e_logo_acima_da_tabela(com_nao_reconhecido, logado):
    """Levi, 18/10/2026: "Quero os filtros que estão antes de extintores não reconhecidos, acima da tabela" e "De
    filtros quero apenas cliente, usina, responsável UFV, equipe e o campo para digitar!"."""
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    ini = html.index('<form class="cn-filtros"')
    form = html[ini:html.index("</form>", ini)]
    assert re.findall(r'<select class="gc-campo" id="ext-(\w+)" name="(\w+)"', form) == [
        ("cliente", "cliente"), ("usina", "usina"), ("responsavel", "responsavel"), ("equipe", "equipe")]
    assert _texto(form).count("Responsável UFV") == 1 and 'type="search" name="q"' in form
    assert not re.search(r'name="(regiao|regiao_campo|gestor)"', html)
    # a ordem na tela: a faixa, a caixa dos não reconhecidos, os filtros e a tabela
    assert html.index("cn-numeros") < html.index('id="nao-reconhecidos"') < html.index('id="tabela"') < html.index(
        '<table class="cn-expande ext-clientes">')
    # em cascata: com o cliente escolhido, a lista de usinas traz só as dele (e a escolhida fica, mesmo sem resultado)
    outro = logado.get("/t/hseq/extintores?cliente=Outro+Cliente").get_data(as_text=True)
    assert re.findall(r'<option value="([^"]*)"', outro[outro.index('id="ext-usina"'):outro.index('id="ext-responsavel"')]) == [
        "", "Coração 1"]
    vazio = logado.get("/t/hseq/extintores?cliente=Outro+Cliente&usina=Altair").get_data(as_text=True)
    assert '<option value="Altair" selected>' in vazio and "Nenhum extintor com esses filtros." in vazio
    # o "Sem cliente" (Irecê 2, fora do cadastro) é uma opção, por último, e acha a usina
    cli = re.findall(r'<option value="([^"]*)"', html[html.index('id="ext-cliente"'):html.index('id="ext-usina"')])
    assert cli == ["", "Outro Cliente", "Thopen", "Sem cliente"]
    sem = logado.get("/t/hseq/extintores?cliente=Sem+cliente&ver=extintor").get_data(as_text=True)
    assert _linhas(sem) == ["IRC200-INFC1-PPCI-EXT01"]


def test_ninguem_abre_filtrado_pelo_login(com_livro, cliente):
    """Até 18/10/2026 o Supervisor de Campo abria na região dele; o filtro de região saiu, e com ele o padrão do login
    (sem o filtro na tela, a pessoa não teria como desfazer)."""
    with cliente.session_transaction() as s:
        s.update(logado=True, admin=False, usuario={"email": "x@exemplo.test", "nome": "X", "perfil": "Supervisor"},
                 supervisor_padrao={"pessoa_id": 93, "nome": "Coordenador Campo", "papel": "supervisor_campo",
                                    "regioes": ["Sul 01"]})
    html = cliente.get("/t/hseq/extintores?ver=extintor").get_data(as_text=True)
    form = html[html.index('<form class="cn-filtros"'):]
    assert len(_linhas(html)) == 6 and " selected>" not in form[:form.index("</form>")]


def test_coluna_da_conta_faltando_a_tela_nao_mostra_numero_pela_metade(com_livro, logado):
    _aba(com_livro, EXT.LIVRO, EXT.ABA, [{k: v for k, v in x.items() if k != "Validade da recarga"} for x in LIVRO])
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "O livro dos extintores chegou sem a coluna Validade da recarga" in html
    assert "cn-numeros" not in html and "<table>" not in html


def test_a_tela_aparece_pronta_no_menu_da_torre(com_livro, logado):
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "campo-nativa" in html and 'href="/t/hseq/extintores"' in html
    assert logado.get("/t/hseq/riscos").status_code == 200                       # as outras seguem no placeholder

# ── por usina e dia (Levi, 09/10: "queria agrupado por usina e dia!") ─────────────────────────────────────────────
DIA_DO_GRUPO = [
    # dois de Altair conferidos no mesmo dia viram UMA linha (e o terceiro, noutro dia, outra)
    _ext("ALT100-INFC1-PPCI-EXT04", "Thopen - Altair 1 - SP", "02/2026", "Sem selo", "2026-10-05", carga="SOBRECARGA"),
    _ext("ALT100-INFC1-PPCI-EXT05", "Thopen - Altair 1 - SP", "Sem data", "Sem selo", "2026-10-05"),
]


def test_por_usina_e_dia_junta_os_extintores_da_mesma_conferencia():
    g = {(x["usina"], x["conferencia"]): x for x in EXT.por_usina_dia(
        [dict(EXT.situacao(e, HOJE), usina=e["Usina"], codigo=e["Código"], ext=e["Código"][-5:])
         for e in LIVRO + DIA_DO_GRUPO])}
    alt = g[("Thopen - Altair 1 - SP", date(2026, 10, 5))]
    assert alt["qtd"] == 3 and [x["ext"] for x in alt["extintores"]] == ["EXT04", "EXT01", "EXT05"]
    # a recarga mais próxima é a mais antiga (vencida há mais tempo): 02/2026, e não 08/2026
    assert (alt["situacao"], alt["recarga"], alt["recarga_fim"], alt["recarga_ext"]) == (
        "atrasado", "02/2026", date(2026, 2, 28), "EXT04")
    assert (alt["atrasado"], alt["sem_validade"], alt["status"]) == (2, 1, "CRITICO")
    assert alt["status_qtd"]["CRITICO"] == 2 and alt["status_qtd"]["OK_SEM_VALIDADE"] == 1
    # a mesma usina noutro dia é outra linha; o grupo sem nenhuma recarga com data diz "Sem data"
    assert g[("Thopen - Altair 1 - SP", date(2026, 8, 1))]["qtd"] == 1
    bwk = g[("Thopen - Brodowski 1 - SP", None)]
    assert (bwk["recarga"], bwk["recarga_fim"], bwk["status"]) == ("Sem data", None, "")
    ordem = [(x["usina"][9:13], x["conferencia"]) for x in EXT.por_usina_dia(
        [dict(EXT.situacao(e, HOJE), usina=e["Usina"], codigo=e["Código"], ext=e["Código"][-5:])
         for e in LIVRO + DIA_DO_GRUPO])]
    assert ordem[0] == ("Alta", date(2026, 10, 5))           # o atrasado com a recarga mais antiga primeiro


def _abre(html, inicio):
    """O texto da linha (a do cliente ou a da usina) em que está `inicio`."""
    i = html.index(inicio)
    return _texto(html[html.rindex("<tr ", 0, i):html.index("</tr>", i)])


def test_a_tabela_abre_por_cliente_e_o_cliente_abre_as_usinas(com_cadastro, logado):
    """Levi, 18/10/2026: "Quero agrupado por Cliente, abrindo o cliente terá as usinas desse cliente. É importante ter a
    coluna do responsável desse usina e o ou os técnicos da usina!". A usina é UMA linha (todos os dias dela) e abre os
    extintores."""
    _aba(com_cadastro, EXT.LIVRO, EXT.ABA, LIVRO + DIA_DO_GRUPO)
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)                  # a tabela é o padrão
    assert 'aria-current="page">Tabela</a>' in html and 'aria-current="page">Cliente</a>' in html
    assert 'href="?ver=extintor#tabela"' in html and "Por usina e dia" not in html
    # UMA tabela: a linha do cliente e as das usinas nas mesmas colunas (aninhadas, não cabiam a 1440 px)
    cab = _texto(html[html.index("<thead>"):html.index("</thead>")])
    # a última coluna é a dos três pontos, sem nome (Levi, 24/10/2026: "tira esse nome ações de uma coluna que não existe")
    assert cab == ("Situação Última conferência Cliente e usina Responsável Técnicos Quantidade de extintores "
                   "Recarga mais próxima Status da TST")
    assert '<th aria-label="Ações"></th>' in html
    # 3 clientes (Thopen, Outro Cliente e "Sem cliente", a Irecê 2 fora do cadastro) e 4 usinas, tudo fechado
    assert html.count('<tr class="cn-linha ext-cli"') == 3 and html.count('<tr class="cn-linha ext-us"') == 4
    assert len(re.findall(r'<tr class="cn-linha ext-us"[^>]* hidden>', html)) == 4
    assert len(re.findall(r'<tr class="cn-detalhe" data-cli="\d" hidden data-url="[^"]+">', html)) == 4
    ordem = re.findall(r'<span class="ext-seta" aria-hidden="true"></span>(?:<b>)?([^<]+)<', html)
    assert ordem == ["Thopen", "Altair", "Brodowski 1", "Outro Cliente", "Coração 1", "Sem cliente",
                     "Thopen - Irecê 2 - BA"]
    thopen = _abre(html, "<b>Thopen</b>")
    assert thopen.startswith("Atrasado Thopen 2 usinas 1 com extintor atrasado 6 2 atrasados · 1 perto de vencer · "
                             "2 sem validade · 1 em dia")
    assert "02/2026 venceu há 7 meses" in thopen and "Crítico 2 críticos" in thopen
    # a usina, uma linha só: a última conferência é a mais recente, com o responsável e os técnicos (um embaixo do outro)
    alt = _abre(html, "</span>Altair</td>")
    assert alt.startswith("Atrasado 05/10/2026 há 4 d · App Altair Supervisora Campo Eletricista Beta Técnico Alfa 5 "
                          "2 atrasados · 1 perto de vencer · 1 sem validade · 1 em dia")
    assert '<span class="ext-tecs"><span>Eletricista Beta</span><span>Técnico Alfa</span></span>' in html
    # o apoio agendado para 20/10 ainda não conta; a usina fora do cadastro não tem responsável nem técnico
    assert "Brodowski 1 Supervisora Campo sem técnico" in _abre(html, "</span>Brodowski 1</td>")
    assert "Coração 1 Coordenador Campo Técnico Gama" in _abre(html, "</span>Coração 1</td>")
    assert "Thopen - Irecê 2 - BA sem responsável sem técnico" in _abre(html, "</span>Thopen - Irecê 2 - BA</td>")
    # os extintores NÃO vêm na página (19/10/2026): a usina os busca ao abrir, de todos os dias, com a conferência de
    # cada um e sem repetir a usina
    assert _linhas(html) == [] and "ext-sub cn-fotos-grupo" not in html
    sub = _pedaco(logado, html, "Altair")
    assert _linhas(sub) == ["ALT100-INFC1-PPCI-EXT04", "ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02",
                            "ALT100-INFC1-PPCI-EXT05", "ALT100-INFC1-PPCI-EXT03"]
    assert "01/08/2026" in _texto(sub) and "Altair" not in _texto(sub) and "sobrecarga" in _texto(sub)
    # o filtro da faixa vale para o agrupamento: só os atrasados entram na conta
    atr = logado.get("/t/hseq/extintores?f=atrasado").get_data(as_text=True)
    assert atr.count('<tr class="cn-linha') == 4                          # Thopen › Altair e Outro Cliente › Coração 1
    assert "Altair Supervisora Campo Eletricista Beta Técnico Alfa 2 2 atrasados" in _abre(atr, "</span>Altair</td>")
    # com o cliente escolhido, o cliente já abre; com a usina, a usina também
    cli = logado.get("/t/hseq/extintores?cliente=Thopen").get_data(as_text=True)
    assert cli.count('aria-expanded="true"') == 1 and cli.count('<tr class="cn-linha ext-us"') == 2
    assert not re.search(r'<tr class="cn-linha ext-us"[^>]* hidden>', cli) and len(re.findall(r'<tr class="cn-detalhe" data-cli="1" hidden', cli)) == 2
    # com a usina (uma só), os extintores dela já vêm na página
    us = logado.get("/t/hseq/extintores?usina=Altair").get_data(as_text=True)
    assert us.count('aria-expanded="true"') == 2 and us.count('<tr class="cn-detalhe" data-cli="1">') == 1
    assert len(_linhas(us)) == 5 and "data-url" not in us


def _pedaco(cliente, html, usina):
    """O pedaço que a linha da usina busca ao abrir (o data-url da linha de detalhe dela)."""
    i = html.index(f"</span>{usina}</td>")
    url = re.search(r'data-url="([^"]+)"', html[i:]).group(1).replace("&amp;", "&")
    r = cliente.get(url)
    assert r.status_code == 200
    return r.get_data(as_text=True)


def test_a_pagina_leva_so_o_resumo_e_a_usina_busca_os_extintores_ao_abrir(com_cadastro, logado):
    """O orçamento de Extintores (19/10/2026: o Raio-X do servidor pegou a tela 21 s montando a página, que levava todos
    os extintores escondidos, 897 KB; Levi: "primeiro mostra a informação e só quando a pessoa clica que começa a
    carregar"). A página leva só o cliente e a usina, com todos os campos de hoje; os extintores vêm por usina."""
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert _linhas(html) == [] and not re.search(r"<img[^>]+src=", html.split('id="tabela"')[1])
    # a linha da usina segue com os campos de antes: situação, conferência, usina, responsável, técnicos, quantidade,
    # recarga, status da TST e o relatório
    assert _abre(html, "</span>Altair</td>").startswith(
        "Atrasado 05/10/2026 há 4 d · App Altair Supervisora Campo Eletricista Beta Técnico Alfa 3 1 atrasado · "
        "1 perto de vencer · 1 em dia 08/2026 venceu há 39 d Crítico 1 crítico · 2 atenção")
    # o pedaço: só os extintores daquela usina (e do cliente dela), com os filtros da tela e o da faixa
    sub = _pedaco(logado, html, "Altair")
    assert _linhas(sub) == ["ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02", "ALT100-INFC1-PPCI-EXT03"]
    assert "<html" not in sub and "Brodowski" not in sub
    atr = logado.get("/t/hseq/extintores?f=atrasado").get_data(as_text=True)
    assert _linhas(_pedaco(logado, atr, "Altair")) == ["ALT100-INFC1-PPCI-EXT01"]
    # sem a usina, ou sem login, o pedaço não sai
    assert logado.get("/t/hseq/extintores/usina").status_code == 404
    assert logado.application.test_client().get("/t/hseq/extintores/usina?usina=Altair").status_code == 302


def test_com_800_extintores_a_pagina_fica_abaixo_de_200_kb(com_livro, logado):
    """O tamanho do real (19/10/2026: ~800 extintores em ~90 usinas; eram 897 KB). Meta do desenho: abaixo de 200 KB."""
    livro = [_ext(f"U{u:02d}100-INFC1-PPCI-EXT{e:02d}", f"Usina de teste {u:02d}", "08/2026", "2027", "2026-10-05")
             for u in range(90) for e in range(9)]
    _aba(com_livro, EXT.LIVRO, EXT.ABA, livro)
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "Extintores 810 em 90 usinas" in _texto(html) and _linhas(html) == []
    assert len(html.encode()) < 200 * 1024, len(html.encode())


def test_sem_ler_os_tecnicos_a_coluna_diz_que_nao_leu(com_cadastro, logado, monkeypatch):
    """A lista do PCM fora do ar: "não lido", e não "sem técnico" (seria dizer que a usina não tem técnico)."""
    monkeypatch.setattr(ATR, "ler", lambda cfg, forcar=False: {"decisoes": [], "erro": "ConnectionError: x",
                                                                 "lido_em": None})
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "Os técnicos de cada usina não foram lidos agora." in html
    assert "não lido" in _abre(html, "</span>Altair</td>") and "sem técnico" not in html


# ── o relatório em PDF (Levi, 09/10, com as observações da TST) ──────────────────────────────────────────────────
def _pdf_texto(corpo):
    pypdf = pytest.importorskip("pypdf")
    import io
    return " ".join(" ".join((p.extract_text() or "").split()) for p in pypdf.PdfReader(io.BytesIO(corpo)).pages)


def test_relatorio_pdf_por_usina_e_dia_com_o_filtro_da_tst(com_livro, logado):
    r = logado.get("/t/hseq/extintores/relatorio.pdf")
    assert r.status_code == 200 and r.mimetype == "application/pdf" and r.data[:4] == b"%PDF"
    assert "inline" in r.headers["Content-Disposition"] and "extintores-todos-2026-10-09.pdf" in r.headers[
        "Content-Disposition"]
    t = _pdf_texto(r.data)
    assert "Relatório de extintores" in t and "Todos os extintores · com as fotos ao lado de cada extintor" in t
    assert "Altair" in t and "Coração 1" in t and "Thopen - Irecê 2 - BA" in t
    assert "Sem foto no Nexus" in t                                   # o padrão é com fotos: o quadro diz que falta
    assert "Recarga mais próxima" in t and "venceu há 39 d" in t
    # só os críticos: o ALT EXT01 (recarga vencida) e o COR EXT01 (sem carga e hidrostático vencido)
    crit = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos").data)
    assert "Só os críticos" in crit and "ALT100-INFC1-PPCI-EXT01" in crit and "COR100-INFC1-PPCI-EXT01" in crit
    assert "EXT02" not in crit and "IRC200" not in crit and "Brodowski" not in crit
    # sem fotos: uma tabela por usina e dia, sem o quadro da foto; os filtros da tela valem e o PDF diz quais
    sem = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?fotos=0&fotos=0&equipe=SC+Oeste+01").data)
    assert "sem fotos" in sem and "Sem foto no Nexus" not in sem and "Filtros: Equipe: SC Oeste 01" in sem
    assert "Coração 1" in sem and "Altair" not in sem


def test_relatorio_pdf_poe_a_foto_que_o_app_enviou_ao_lado_do_extintor(com_livro, logado, app, tmp_path):
    from PIL import Image
    pasta = tmp_path / "hseq" / "extintores" / "fotos"
    pasta.mkdir(parents=True)
    Image.new("RGB", (1200, 900), (200, 30, 30)).save(pasta / "ALT100-INFC1-PPCI-EXT01.jpg")
    app.config["NEXUS_DADOS"] = str(tmp_path)
    r = logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos")
    pypdf = pytest.importorskip("pypdf")
    import io
    paginas = pypdf.PdfReader(io.BytesIO(r.data)).pages
    fotos = [im for p in paginas for im in p.images if im.name.endswith(".jpg")]     # o logo do topo é PNG
    assert len(fotos) == 1 and max(fotos[0].image.size) <= 600         # reduzida (a de 1200 px não entra inteira)
    assert "Sem foto no Nexus" in _pdf_texto(r.data)                   # o outro crítico continua sem foto


def test_o_pdf_usa_a_copia_de_600_px_feita_uma_vez_por_foto(com_livro, logado, app, tmp_path):
    """14/10/2026, com as 90 primeiras fotos do App: o PDF de todos os extintores reduzia as 90 a cada pedido (65 s no
    servidor). A cópia de 600 px é feita uma vez e refeita só quando o App manda outra foto; o PDF sai em binário."""
    import io
    import os
    from PIL import Image
    from nexus.hseq import fotos as FOT
    cod = "ALT100-INFC1-PPCI-EXT01"
    pasta = tmp_path / "hseq" / "extintores" / "fotos"
    pasta.mkdir(parents=True)
    Image.new("RGB", (1200, 1600), (200, 30, 30)).save(pasta / f"{cod}.jpg", quality=82)
    app.config["NEXUS_DADOS"] = str(tmp_path)
    r = logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos")
    copia = pasta / "pdf" / f"{cod}.jpg"
    assert r.status_code == 200 and copia.is_file() and Image.open(copia).size == (450, 600)
    assert b"/ASCII85Decode" not in r.data and b"/DCTDecode" in r.data   # binário: sem o ASCII85 do reportlab
    antes = copia.stat().st_mtime_ns
    logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos")
    assert copia.stat().st_mtime_ns == antes                            # o 2º PDF não refaz a cópia
    dados = copia.read_bytes()
    assert FOT.reduzir(dados, 600, 75)[0] is dados                      # e ela entra no PDF sem ser decodificada
    assert sorted(p.name for p in pasta.iterdir()) == [f"{cod}.jpg", "pdf"]   # a miniatura e a tela não veem a cópia
    assert [p.name for p in (pasta / "pdf").iterdir()] == [f"{cod}.jpg"]      # e nenhum .parcial esquecido
    assert FOT.disponiveis(app.config) == {cod: ""}
    # o App mandou outra foto (a inteira fica mais nova que a cópia): a cópia é refeita
    Image.new("RGB", (1600, 900), (40, 160, 90)).save(pasta / f"{cod}.jpg", quality=82)
    os.utime(pasta / f"{cod}.jpg", ns=(antes + 10**9, antes + 10**9))
    larg, alt = Image.open(io.BytesIO(FOT.foto_pdf(app.config, cod))).size
    assert larg == 600 and alt < larg and Image.open(copia).size == (larg, alt)
    # sem a foto, a cópia velha não aparece
    (pasta / f"{cod}.jpg").unlink()
    assert FOT.foto_pdf(app.config, cod) is None


def test_sem_o_livro_o_relatorio_volta_para_a_tela(banco, logado):
    r = logado.get("/t/hseq/extintores/relatorio.pdf")
    assert r.status_code == 302 and r.headers["Location"].endswith("/t/hseq/extintores")


# ── o relatório de cada linha (PDF e Excel), o filtro de cliente e a foto do extintor ─────────────────────────────
def _baixar(html, rotulo="PDF", de="usina"):
    """Os endereços do menu dos três pontos de cada linha (24/10/2026: o menu é um só e o clique põe nele os filtros da
    linha, `data-args`): os das usinas (`de="usina"`) ou os dos clientes (`de="cliente"`). A mesma conta do JavaScript
    da tela."""
    args = [a.replace("&amp;", "&") for a in re.findall(r'<button class="hs-tres"[^>]*data-args="([^"]*)"', html)]
    args = [a for a in args if ("usina=" in a) == (de == "usina")]
    if rotulo == "PDF":
        return ["/t/hseq/extintores/relatorio.pdf?baixar=1" + ("&" + a if a else "") for a in args]
    return ["/t/hseq/extintores/relatorio.xlsx" + ("?" + a if a else "") for a in args]


def test_o_pdf_de_cada_usina_traz_a_usina_inteira(com_livro, logado):
    """O relatório só unitário (Levi, 18/10/2026: "O relatório quero a opção de ser unitário (...), pode tirar esse
    RELATÓRIO EM PDF que tira em massa!"): a tela não tem mais o formulário do PDF de todos; o PDF de cada usina traz a
    usina inteira, todos os dias."""
    _aba(com_livro, EXT.LIVRO, EXT.ABA, LIVRO + DIA_DO_GRUPO)
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "Relatório em PDF" not in html and "Gerar PDF" not in html and 'name="status"' not in html
    links = _baixar(html)
    assert len(links) == 4 and not any("dia=" in h for h in links)                  # uma por usina
    alt = next(h for h in links if "usina=Altair" in h)
    assert "baixar=1" in alt
    r = logado.get(alt)
    assert r.status_code == 200 and r.headers["Content-Disposition"].startswith("attachment;")
    assert 'filename="extintores-altair-2026-10-09.pdf"' in r.headers["Content-Disposition"]
    t = _pdf_texto(r.data)
    assert all(f"ALT100-INFC1-PPCI-EXT0{n}" in t for n in range(1, 6)) and "Coração" not in t     # os 3 dias dela
    assert "Usina: Altair" in t and "Dia da conferência" not in t
    # a usina nunca conferida também baixa, e os filtros da tela e o da faixa vão junto no link
    bwk = next(h for h in links if "usina=Brodowski" in h)
    assert "BWK100-INFC1-PPCI-EXT01" in _pdf_texto(logado.get(bwk).data)
    atr = _baixar(logado.get("/t/hseq/extintores?f=atrasado&cliente=Thopen").get_data(as_text=True))
    assert atr and all("f=atrasado" in h and "cliente=Thopen" in h for h in atr)
    # o link de antes, com o dia, ainda filtra aquele dia
    dia = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?usina=Altair&dia=2026-10-05&baixar=1").data)
    assert "ALT100-INFC1-PPCI-EXT04" in dia and "ALT100-INFC1-PPCI-EXT02" not in dia


def test_o_pdf_de_cada_cliente_traz_as_usinas_dele(com_livro, logado):
    """A TST, 19/10/2026: "tem como eu baixar um relatório dos extintores por cliente?". A linha do cliente tem o PDF
    dele (todas as usinas do cliente, com os filtros da tela)."""
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    cab = _texto(html[html.index("<thead>"):html.index("</thead>")])
    assert cab.endswith("Status da TST")
    clientes = _baixar(html, de="cliente")
    assert len(clientes) == 3                                       # Thopen, Outro Cliente e Sem cliente
    th = next(h for h in clientes if "cliente=Thopen" in h)
    r = logado.get(th)
    assert r.status_code == 200 and r.headers["Content-Disposition"].startswith("attachment;")
    assert 'filename="extintores-thopen-2026-10-09.pdf"' in r.headers["Content-Disposition"]
    t = _pdf_texto(r.data)
    assert "Cliente: Thopen" in t and "Altair" in t and "Brodowski" in t and "Coração" not in t and "Irecê" not in t
    sem = _pdf_texto(logado.get(next(h for h in clientes if "cliente=Sem+cliente" in h)).data)
    assert "IRC200-INFC1-PPCI-EXT01" in sem and "Altair" not in sem
    # com um filtro na tela, ele vai junto (e o cliente da linha vence o do filtro)
    eq = _baixar(logado.get("/t/hseq/extintores?equipe=SP+Norte+01").get_data(as_text=True), de="cliente")
    assert eq == ["/t/hseq/extintores/relatorio.pdf?baixar=1&equipe=SP+Norte+01&cliente=Thopen"]


def _planilha(corpo):
    import io
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(corpo))


def test_o_excel_do_cliente_e_da_usina_tem_um_extintor_por_linha(com_cadastro, logado):
    """A TST, 19/10/2026: "e excel tbm? pq às vezes o cliente pede". O mesmo recorte do PDF, em planilha: uma linha por
    extintor, na ordem da tabela, com a conferência como data e o filtro do Excel no cabeçalho."""
    from nexus.hseq import planilha as PLA
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert len(_baixar(html, "Excel", "cliente")) == 3 and len(_baixar(html, "Excel")) == 4
    url = next(h for h in _baixar(html, "Excel", "cliente") if "cliente=Thopen" in h)
    r = logado.get(url)
    assert r.status_code == 200 and r.mimetype == PLA.MIME
    assert r.headers["Content-Disposition"] == 'attachment; filename="extintores-thopen-2026-10-09.xlsx"'
    wb = _planilha(r.data)
    assert wb.sheetnames == ["Extintores"]                                  # sem não reconhecido no livro
    ws = wb["Extintores"]
    assert ws["A1"].value == "Relatório de extintores" and ws["A3"].value == "Filtros: Cliente: Thopen"
    assert ws["A4"].value == "4 extintores"
    cab = [c.value for c in ws[6]]
    assert cab[:6] == ["Cliente", "Usina", "Responsável", "Técnicos", "Extintor", "Código"]
    assert ws.freeze_panes == "A7" and ws.auto_filter.ref.startswith("A6:")
    linhas = [dict(zip(cab, (c.value for c in l))) for l in ws.iter_rows(min_row=7)]
    # a ordem da tabela: Altair (o mais grave primeiro dentro dela) e depois Brodowski 1
    assert [x["Código"] for x in linhas] == ["ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02",
                                              "ALT100-INFC1-PPCI-EXT03", "BWK100-INFC1-PPCI-EXT01"]
    alt = linhas[0]
    assert (alt["Cliente"], alt["Usina"], alt["Responsável"], alt["Técnicos"]) == (
        "Thopen", "Altair", "Supervisora Campo", "Eletricista Beta, Técnico Alfa")
    assert (alt["Situação"], alt["Validade da recarga"], alt["Status da TST"]) == ("Atrasado", "08/2026", "Crítico")
    assert alt["Recarga vence em"].date() == date(2026, 8, 31) and alt["Última conferência"].date() == date(2026, 10, 5)
    assert alt["Conferido por"] == "Técnico Silva" and alt["Sem atualização há mais de 30 dias"] == "não"
    assert ws.cell(row=7, column=cab.index("Última conferência") + 1).number_format == "DD/MM/YYYY"
    bwk = linhas[3]
    assert bwk["Última conferência"] is None and bwk["Status da TST"] == "sem conferência" and bwk["Técnicos"] is None
    assert not any(codigo_da_pessoa(CHAVE, "tec1@exemplo.test") in str(c.value or "") for l in ws.iter_rows() for c in l)
    # o da usina: só ela
    us = _planilha(logado.get(next(h for h in _baixar(html, "Excel") if "usina=Cora" in h)).data)["Extintores"]
    assert [c.value for c in us["F"][6:]] == ["COR100-INFC1-PPCI-EXT01"]


def test_o_excel_leva_os_nao_reconhecidos_numa_aba_a_parte(com_nao_reconhecido, logado):
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    wb = _planilha(logado.get(next(h for h in _baixar(html, "Excel") if "usina=Altair" in h)).data)
    assert wb.sheetnames == ["Extintores", "Não reconhecidos"]
    nr = wb["Não reconhecidos"]
    cab = [c.value for c in nr[6]]
    assert cab[:4] == ["Cliente", "Usina", "Extintor", "Código"]
    assert [c.value for c in nr[7]][:4] == ["Thopen", "Altair", "EXT02", "ALT100-INFC1-PPCI-EXT02"]
    assert "ALT100-INFC1-PPCI-EXT02" not in [c.value for c in wb["Extintores"]["F"]]      # fora das contas
    # o filtro da faixa tira a aba (a regra do PDF)
    atr = _planilha(logado.get("/t/hseq/extintores/relatorio.xlsx?usina=Altair&f=atrasado").data)
    assert atr.sheetnames == ["Extintores"]


def test_duas_usinas_com_o_mesmo_nome_em_clientes_diferentes_nao_se_juntam(com_livro, logado):
    """O cadastro tem duas "Linhares 1" (ES), uma de cada cliente (19/10/2026: pelo nome, 13 extintores da Thopen
    apareciam na Axis). A usina se separa pelo ID do cadastro; o link da usina leva também o cliente."""
    from test_campo_visao import _MOB
    outra = {"usina_id": 6, "nome": "Altair", "codigo": "OUTR-ALT200", "status": "OPERAÇÃO", "equipe_id": 20,
             "data_mobilizacao": _MOB, "uf": "SC", "cidade": "Altair", "cliente_id": 2, "cluster": "SC oeste"}
    _aba(com_livro, "cadastro_nexus", "usinas", USINAS + [outra])
    _aba(com_livro, EXT.LIVRO, EXT.ABA, LIVRO + [
        _ext("ALT200-INFC1-PPCI-EXT01", "Outro - Altair - SC", "11/2026", "Sem selo", "2026-10-08", cb="ALT200")])
    visao.limpar()
    html = logado.get("/t/hseq/extintores").get_data(as_text=True)
    assert "Extintores 7 em 5 usinas" in _texto(html)
    assert "Outro Cliente 2 usinas" in _abre(html, "<b>Outro Cliente</b>")
    assert "Thopen 2 usinas" in _abre(html, "<b>Thopen</b>")
    links = _baixar(html, "Excel")
    th = next(h for h in links if "usina=Altair" in h and "cliente=Thopen" in h)
    ou = next(h for h in links if "usina=Altair" in h and "cliente=Outro+Cliente" in h)
    assert [c.value for c in _planilha(logado.get(th).data)["Extintores"]["F"][6:]] == [
        "ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02", "ALT100-INFC1-PPCI-EXT03"]
    assert [c.value for c in _planilha(logado.get(ou).data)["Extintores"]["F"][6:]] == ["ALT200-INFC1-PPCI-EXT01"]
    # o PDF da usina também
    pdf = _pdf_texto(logado.get(next(h for h in _baixar(html) if "cliente=Outro+Cliente" in h and "Altair" in h)).data)
    assert "ALT200-INFC1-PPCI-EXT01" in pdf and "ALT100" not in pdf


def test_sem_o_livro_o_excel_volta_para_a_tela(banco, logado):
    r = logado.get("/t/hseq/extintores/relatorio.xlsx")
    assert r.status_code == 302 and r.headers["Location"].endswith("/t/hseq/extintores")


def test_filtro_por_cliente_vale_na_tela_no_pdf_e_no_epi(com_livro, logado):
    html = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor").get_data(as_text=True)
    assert '<select class="gc-campo" id="ext-cliente" name="cliente"' in html
    assert '<option value="Outro Cliente">Outro Cliente</option>' in html and '<option value="Thopen">Thopen</option>' in html
    outro = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor&cliente=Outro+Cliente").get_data(as_text=True)
    assert _linhas(outro) == ["COR100-INFC1-PPCI-EXT01"] and '<option value="Outro Cliente" selected>' in outro
    pdf = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?cliente=Thopen").data)
    assert "Filtros: Cliente: Thopen" in pdf and "Coração" not in pdf and "Altair" in pdf
    epi =logado.get("/t/hseq/epi?modo=tabela&cliente=Outro+Cliente").get_data(as_text=True)
    assert '<select class="gc-campo" id="cn-cliente" name="cliente"' in epi and "Coração 1" in epi and "Altair" not in epi


def _foto_jpeg():
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (2400, 1800), (40, 160, 90)).save(buf, "JPEG")
    return buf.getvalue()


def _enviar(cliente, app, codigo="ALT100-INFC1-PPCI-EXT01", dia="2026-10-08", ts=None, dados=None, assinatura=None):
    import io
    import time
    from nexus.hseq import fotos as FOT
    dados = _foto_jpeg() if dados is None else dados
    ts = int(time.time()) if ts is None else ts
    sig = assinatura or FOT.assinatura(FOT.chave(app.config), codigo, dia, ts, dados)
    return cliente.post("/t/hseq/extintores/foto", data={"codigo": codigo, "dia": dia, "ts": str(ts),
                                                         "assinatura": sig, "foto": (io.BytesIO(dados), "f.jpg")},
                        content_type="multipart/form-data")


def test_o_app_envia_a_foto_assinada_e_ela_aparece_na_tabela_e_no_pdf(com_livro, app, cliente, logado, tmp_path):
    import time
    app.config["NEXUS_DADOS"] = str(tmp_path)
    # sem login (é o App), mas só com a assinatura certa, na hora certa e com imagem de verdade
    anon = app.test_client()
    assert _enviar(anon, app, assinatura="0" * 64).status_code == 401
    assert _enviar(anon, app, ts=int(time.time()) - 3600).status_code == 401
    assert _enviar(anon, app, dados=b"isto nao e imagem").status_code == 400
    assert _enviar(anon, app, codigo="../../etc").status_code == 400
    r = _enviar(anon, app)
    assert r.status_code == 200 and r.get_json()["ok"] is True
    guardada = tmp_path / "hseq" / "extintores" / "fotos" / "ALT100-INFC1-PPCI-EXT01.jpg"
    from PIL import Image
    assert guardada.is_file() and max(Image.open(guardada).size) == 1600                 # reduzida a 1600 px
    # a foto de um dia mais velho não substitui a guardada
    assert "mais nova" in _enviar(anon, app, dia="2026-09-01").get_json()["mensagem"]
    # ver a foto pede login; a tabela mostra a miniatura e a caixa que amplia
    assert anon.get("/t/hseq/extintores/foto/ALT100-INFC1-PPCI-EXT01.jpg").status_code == 302
    assert logado.get("/t/hseq/extintores/foto/ALT100-INFC1-PPCI-EXT01.jpg").mimetype == "image/jpeg"
    assert logado.get("/t/hseq/extintores/foto/ALT100-INFC1-PPCI-EXT02.jpg").status_code == 404
    # a tabela: o botão "Foto" com o dia, sem miniatura (a imagem só no clique; 19/10/2026), na usina aberta e por extintor
    for url in ("/t/hseq/extintores?ver=extintor", "/t/hseq/extintores?usina=Altair"):
        html = logado.get(url).get_data(as_text=True)
        assert html.count('class="cn-foto ext-foto-bt"') == 1 and "foto de 08/10/2026" in html and ">Foto 08/10</a>" in html
        assert not re.search(r"<img[^>]+extintores/foto/", html) and 'class="cn-caixa-foto"' in html
    pedaco = _pedaco(logado, logado.get("/t/hseq/extintores").get_data(as_text=True), "Altair")
    assert ">Foto 08/10</a>" in pedaco and "<img" not in pedaco
    pypdf = pytest.importorskip("pypdf")
    import io
    paginas = pypdf.PdfReader(io.BytesIO(logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos").data)).pages
    assert len([im for p in paginas for im in p.images if im.name.endswith(".jpg")]) == 1


def test_a_miniatura_le_a_foto_inteira_uma_vez_so(com_livro, app, logado, tmp_path, monkeypatch):
    """19/10/2026: cada miniatura lia a foto inteira do disco antes de olhar a memória e, sem a memória (cada deploy),
    decodificava a inteira. Agora: memória, depois a cópia de 160 px no disco, e a inteira só na 1ª vez."""
    from PIL import Image
    from nexus.hseq import fotos as FOT
    cod = "ALT100-INFC1-PPCI-EXT01"
    pasta = tmp_path / "hseq" / "extintores" / "fotos"
    pasta.mkdir(parents=True)
    Image.new("RGB", (1200, 1600), (200, 30, 30)).save(pasta / f"{cod}.jpg", quality=82)
    app.config["NEXUS_DADOS"] = str(tmp_path)
    lidas = []
    original = FOT.foto
    monkeypatch.setattr(FOT, "foto", lambda cfg, c: lidas.append(c) or original(cfg, c))
    FOT._MINIS.clear()
    r = logado.get(f"/t/hseq/extintores/foto/{cod}.jpg?mini=1")
    assert r.status_code == 200 and max(Image.open(__import__("io").BytesIO(r.data)).size) == FOT.MINI_PX
    assert lidas == [cod] and (pasta / "mini" / f"{cod}.jpg").is_file()
    logado.get(f"/t/hseq/extintores/foto/{cod}.jpg?mini=1")
    assert lidas == [cod]                                              # da memória
    FOT._MINIS.clear()                                                  # o reinício de um deploy
    assert logado.get(f"/t/hseq/extintores/foto/{cod}.jpg?mini=1").status_code == 200 and lidas == [cod]   # do disco
    assert FOT.mini(app.config, "ALT100-INFC1-PPCI-EXT09") is None and lidas == [cod]   # sem foto: nem tenta ler


def test_a_lista_das_fotos_fica_guardada_e_a_foto_nova_aparece_na_hora(com_livro, app, tmp_path):
    from nexus.hseq import fotos as FOT
    app.config["NEXUS_DADOS"] = str(tmp_path)
    pasta = tmp_path / "hseq" / "extintores" / "fotos"
    pasta.mkdir(parents=True)
    assert FOT.disponiveis(app.config) == {}
    (pasta / "ALT100-INFC1-PPCI-EXT02.jpg").write_bytes(_foto_jpeg())
    assert FOT.disponiveis(app.config) == {}                            # guardada: não lista a pasta a cada pedido
    assert _enviar(app.test_client(), app).status_code == 200            # a foto que chega renova a lista
    assert set(FOT.disponiveis(app.config)) == {"ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT02"}


def test_sem_a_chave_o_recebimento_de_fotos_fica_desligado(com_livro, app, tmp_path):
    app.config["NEXUS_DADOS"] = str(tmp_path)
    app.config["NEXUS_PESSOA_HMAC"] = ""
    r = app.test_client().post("/t/hseq/extintores/foto", data={"codigo": "ALT100-INFC1-PPCI-EXT01"})
    assert r.status_code == 503 and "desligado" in r.get_json()["mensagem"]


# ── não reconhecido em campo (Levi, 14/10/2026, na Axis Linhares 1: a equipe só achou 4 dos 17 extintores do cadastro,
#    "pode ter um botão no app de campo de 'Não reconhecido' para o extintor e aí apareceria para a pessoa da Segurança") ──
@pytest.fixture
def com_nao_reconhecido(banco, monkeypatch):
    """O livro com as duas colunas do não reconhecido: o EXT02 da Altair não foi achado na ronda de 05/10 (a mesma em que
    o EXT01 foi conferido); a marca do Coração (01/02) é mais velha que a conferência dele (01/03) e não vale."""
    monkeypatch.setattr(visao, "_agora", lambda: datetime(2026, 10, 9, 10, 0, tzinfo=BRT))
    quem = codigo_da_pessoa(CHAVE, "tec1@exemplo.test")
    marcas = {"ALT100-INFC1-PPCI-EXT02": "2026-10-05", "COR100-INFC1-PPCI-EXT01": "2026-02-01"}
    livro = [{**x, "Não reconhecido em": marcas.get(x["Código"]),
              "Não reconhecido por (HMAC)": quem if x["Código"] in marcas else None} for x in LIVRO]
    _aba(banco, EXT.LIVRO, EXT.ABA, livro)
    visao.limpar()
    return banco


def test_nao_reconhecido_sai_das_contas_e_vai_para_a_caixa_da_tst(com_nao_reconhecido, logado):
    html = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor").get_data(as_text=True)
    caixa = html[html.index('id="nao-reconhecidos"'):html.index("</section>")]
    assert _linhas(caixa) == ["ALT100-INFC1-PPCI-EXT02"]                 # a marca velha do Coração não vale
    t = _texto(caixa)
    assert "1 extintor não reconhecido em campo" in t and "é lá que a Segurança aceita (dá baixa) ou pede uma nova ronda" in t
    assert "05/10/2026 há 4 d Técnico Silva" in t                           # quando e quem marcou, pelo nome
    assert "01/08/2026" in t                                                # a conferência de antes da marca
    assert codigo_da_pessoa(CHAVE, "tec1@exemplo.test") not in html
    # fora de TODAS as contas: a tabela, o total da faixa (5, não 6) e o perto de vencer (o EXT02 vencia em 22 dias)
    assert _linhas(html[html.index("</section>"):]) == ["ALT100-INFC1-PPCI-EXT01", "COR100-INFC1-PPCI-EXT01",
                                                        "BWK100-INFC1-PPCI-EXT01", "IRC200-INFC1-PPCI-EXT01",
                                                        "ALT100-INFC1-PPCI-EXT03"]
    faixa = _texto(html[html.index("cn-numeros"):html.index('id="nao-reconhecidos"')])
    assert "Extintores 5" in faixa and "Perto de vencer 0" in faixa and "Não reconhecidos 1" in faixa
    sem = logado.get("/t/hseq/extintores?f=sem_atualizacao&ver=extintor").get_data(as_text=True)
    assert "ALT100-INFC1-PPCI-EXT02" not in _linhas(sem[sem.index("</section>"):])
    # os filtros da tela valem na caixa: o cliente do Coração não tem não reconhecido, e a caixa some
    outro = logado.get("/t/hseq/extintores?cliente=Outro+Cliente").get_data(as_text=True)
    assert 'id="nao-reconhecidos"' not in outro and "Não reconhecidos" not in outro


def test_a_linha_do_nao_reconhecido_leva_a_tabela_filtrada_na_usina(com_nao_reconhecido, logado):
    """Levi, 18/10/2026: "Em extintores não reconhecidos quero que ao clicar na linha já encaminhe para a tabela
    filtrando aquela usina". A linha (e o nome da usina, pelo teclado) leva à tabela por cliente com a usina no filtro,
    o cliente e a usina já abertos; a situação, a visão e o agrupamento de antes saem."""
    html = logado.get("/t/hseq/extintores?f=sem_atualizacao&ver=extintor&cliente=Thopen").get_data(as_text=True)
    caixa = html[html.index('id="nao-reconhecidos"'):html.index("</section>")]
    assert '<tr class="ext-ir" data-href="?cliente=Thopen&amp;usina=Altair#tabela"' in caixa
    assert '<a class="cn-link" href="?cliente=Thopen&amp;usina=Altair#tabela">Altair</a>' in caixa
    us = logado.get("/t/hseq/extintores?cliente=Thopen&usina=Altair").get_data(as_text=True)
    assert '<option value="Altair" selected>' in us and us.count('aria-expanded="true"') == 2
    tab = us[us.index('<table class="cn-expande ext-clientes">'):]
    # na usina, o não reconhecido vem em cima dos outros, para a Segurança decidir ali (24/10/2026), e segue fora das contas
    assert _linhas(tab) == ["ALT100-INFC1-PPCI-EXT02", "ALT100-INFC1-PPCI-EXT01", "ALT100-INFC1-PPCI-EXT03"]
    assert tab.count('<tr class="hs-nr">') == 1
    assert "2 1 atrasado · 1 em dia · 1 não reconhecido" in _abre(us, "</span>Altair</td>")    # na conta, à parte
    assert _linhas(us[us.index('id="nao-reconhecidos"'):us.index("</section>")]) == ["ALT100-INFC1-PPCI-EXT02"]


def test_sem_as_colunas_novas_ninguem_e_nao_reconhecido(com_livro, logado):
    """O livro do App antes da versão que publica as duas colunas: a tela segue igual, sem a caixa."""
    html = logado.get("/t/hseq/extintores?modo=tabela&ver=extintor").get_data(as_text=True)
    assert 'id="nao-reconhecidos"' not in html and len(_linhas(html)) == 6
    assert EXT.nao_reconhecido({"Não reconhecido em": "2026-10-05"}, date(2026, 10, 6)) is None    # achado depois
    assert EXT.nao_reconhecido({"Não reconhecido em": "2026-10-05"}, date(2026, 10, 5)) == date(2026, 10, 5)
    assert EXT.nao_reconhecido({"Não reconhecido em": "2026-10-05"}, None) == date(2026, 10, 5)
    assert EXT.nao_reconhecido({}, None) is None and EXT.COLUNAS[-2:] == ("Não reconhecido em",
                                                                          "Não reconhecido por (HMAC)")


def test_nao_reconhecido_no_pdf_vai_a_parte_so_no_relatorio_de_todos(com_nao_reconhecido, logado):
    t = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?fotos=0").data)
    assert "Não reconhecidos em campo 1 extintor" in t and "EXT02 · ALT100-INFC1-PPCI-EXT02" in t
    assert "05/10/2026 Técnico Silva" in t and "1 não reconhecidos em campo" in t     # e a casa no resumo
    # só os críticos (sem status da TST, o não reconhecido não entra) e o filtro da faixa: sem a seção
    assert "Não reconhecidos em campo" not in _pdf_texto(
        logado.get("/t/hseq/extintores/relatorio.pdf?status=criticos&fotos=0").data)
    assert "Não reconhecidos em campo" not in _pdf_texto(
        logado.get("/t/hseq/extintores/relatorio.pdf?f=atrasado&fotos=0").data)
    # o "Baixar" da Altair (a usina em que o EXT02 não foi achado) leva a seção; o do Coração, não
    links = _baixar(logado.get("/t/hseq/extintores").get_data(as_text=True))
    alt = _pdf_texto(logado.get(next(h for h in links if "usina=Altair" in h)).data)
    assert "Não reconhecidos em campo 1 extintor" in alt and "ALT100-INFC1-PPCI-EXT02" in alt
    cor = _pdf_texto(logado.get(next(h for h in links if "usina=Cora" in h)).data)
    assert "Não reconhecidos em campo" not in cor
    # o link de antes, com o dia: só os marcados naquele dia
    dia01 = _pdf_texto(logado.get("/t/hseq/extintores/relatorio.pdf?usina=Altair&dia=2026-10-01&baixar=1").data)
    assert "Não reconhecidos em campo" not in dia01
