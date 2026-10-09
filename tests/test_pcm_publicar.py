"""Full Nexus no PCM (Levi, 09/10/2026: "faça as adaptações necessárias para que isso seja possível, semana que vem já
quero full nexus sem falta").

O que estes testes seguram: a publicação no App pelo caminho do PC do PCM (a planilha e as observações no repositório do
PCM, conferidas relendo, sem tocar nos ajustes da semana em curso), o plano publicado na reserva do histórico, o que barra
uma publicação, a tela de publicar (só administrador, com confirmação), o plano publicado lido do repositório do PCM
(também no servidor), os feriados do ano seguinte projetados (a W53 de 2026 vai a 01/01/2027), as observações
importadas do repositório e a conferência que aponta a OS forçada além das horas do dia. Nenhum teste vai ao GitHub nem
ao Fracttal: um GitHub falso em memória e o motor falso dos testes da geração."""
import base64
import hashlib
import io
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytest

from nexus.pcm import comparar as C
from nexus.pcm import geracao as G
from nexus.pcm import historico_banco as HB
from nexus.pcm import insumos as I
from nexus.pcm import publicar as PB
from test_pcm_geracao import CAB, config, origem, planilha  # noqa: F401 (fixtures)
from test_pcm_insumos import ler_feriados

_BRT = timezone(timedelta(hours=-3))
QUINTA_W41 = datetime(2026, 10, 8, 20, 0, tzinfo=_BRT)       # a W41 de teste ainda não acabou
SABADO_ANTES = datetime(2026, 10, 3, 10, 0, tzinfo=_BRT)      # nem começou


class Resp:
    def __init__(self, status, corpo=None, conteudo=b""):
        self.status_code, self._corpo, self.content = status, corpo, conteudo

    def json(self):
        return self._corpo

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class GitHubFalso:
    """A API de conteúdo do GitHub em memória: {caminho: bytes}, um commit por PUT."""

    def __init__(self, arquivos=None, estragar_leitura=False):
        self.arquivos = dict(arquivos or {})
        self.commits = []
        self.estragar = estragar_leitura
        self.tokens = set()

    def _caminho(self, url):
        from urllib.parse import unquote
        return unquote(url.split("/contents/", 1)[1])

    def get(self, url, headers=None, params=None, timeout=None):
        self.tokens.add((headers or {}).get("Authorization"))
        if url.endswith("/commits"):
            caminho = (params or {}).get("path")
            cs = [c for c in self.commits if c["caminho"] == caminho]
            return Resp(200, [{"sha": c["sha"], "commit": {"author": {"name": "Pessoa Teste", "date": "2026-10-08T23:00:00Z"},
                                                           "message": c["mensagem"]}} for c in reversed(cs)])
        caminho = self._caminho(url)
        if caminho not in self.arquivos:
            return Resp(404, {"message": "Not Found"})
        dados = self.arquivos[caminho]
        if self.estragar and self.commits:
            dados = dados + b"x"
        return Resp(200, {"sha": hashlib.sha1(dados).hexdigest(), "content": base64.b64encode(dados).decode()})

    def put(self, url, headers=None, json=None, timeout=None):
        self.tokens.add((headers or {}).get("Authorization"))
        caminho = self._caminho(url)
        if caminho in self.arquivos and json.get("sha") != hashlib.sha1(self.arquivos[caminho]).hexdigest():
            return Resp(409, {"message": "sha errado"})
        self.arquivos[caminho] = base64.b64decode(json["content"])
        sha = hashlib.sha1(f"{caminho}{len(self.commits)}".encode()).hexdigest()
        self.commits.append({"caminho": caminho, "mensagem": json["message"], "sha": sha})
        return Resp(201, {"commit": {"sha": sha}})


@pytest.fixture
def agora(monkeypatch):
    monkeypatch.setattr(PB, "_agora", lambda: QUINTA_W41)
    return QUINTA_W41


def _rodada(cfg, origem_, semana="2026-W41"):
    planilha(origem_ / "Programação Semana 41.xlsx")
    r = G.iniciar(cfg, semana)
    st = G.aguardar(cfg, r["id"], 60)
    assert st["estado"] == "ok", st
    return r["id"]


def _com_github(cfg, gh=None):
    gh = gh or GitHubFalso()
    cfg.update(NEXUS_PCM_GITHUB_TOKEN="tok-teste", NEXUS_PCM_GITHUB_SESSAO=gh)
    return gh


# ── publicar ─────────────────────────────────────────────────────────────────────────────────────────────────
def test_publica_a_planilha_e_as_observacoes_e_confere(config, origem, agora):
    rid = _rodada(config, origem)
    gh = _com_github(config, GitHubFalso({"Observacoes_Semana_Atual.txt": b"15448; nao\n"}))
    pub = PB.publicar(config, rid, "pessoa.teste@exemplo.com")
    pasta = G.pasta_trabalho(config) / "geracoes" / rid
    # o mesmo caminho do PC do PCM: a planilha com o nome que o robô procura e as observações da rodada
    assert gh.arquivos["Programação Semana 41.xlsx"] == (pasta / "saida" / "sombra.xlsx").read_bytes()
    assert gh.arquivos["Observacoes_Semana.txt"] == (pasta / "Observacoes_Semana.txt").read_bytes()
    # os ajustes da semana em curso ficam como estavam (são do painel do PCM)
    assert gh.arquivos["Observacoes_Semana_Atual.txt"] == b"15448; nao\n"
    assert [c["caminho"] for c in gh.commits] == ["Programação Semana 41.xlsx", "Observacoes_Semana.txt"]
    # o repositório é público: a mensagem do commit não leva o e-mail de quem publicou
    assert all("Nexus publica" in c["mensagem"] and "@" not in c["mensagem"] for c in gh.commits)
    assert gh.tokens == {"Bearer tok-teste"}
    assert all(a["conferido"] for a in pub["arquivos"]) and pub["quem"] == "pessoa.teste@exemplo.com"
    st = G.ler_status(config, rid)
    assert st["publicacao"]["semana"] == "2026-W41" and (pasta / "publicacao.json").exists()
    # o plano publicado foi para a reserva do histórico: a W41 com as 2 tarefas do motor falso
    assert HB.semanas_da_reserva(I.carregar(G.pasta_trabalho(config))["historico"])["2026-W41"] == {
        "15020|ALT1-INV01", "15021|ALT1-CAB01"}
    assert PB.publicacoes(config)["2026-W41"]["rodada"] == rid


def test_republicar_troca_o_plano_da_semana_na_reserva(tmp_path):
    I.registrar_plano_publicado(tmp_path, "2026-W42", ["1|A", "2|B"], "teste")
    I.registrar_plano_publicado(tmp_path, "2026-W43", ["2|B"], "teste")
    r = I.registrar_plano_publicado(tmp_path, "2026-W42", ["1|A", "3|C"], "teste")
    sem = HB.semanas_da_reserva(I.carregar(tmp_path)["historico"])
    assert sem["2026-W42"] == {"1|A", "3|C"} and sem["2026-W43"] == {"2|B"} and r == {"semana": "2026-W42", "chaves": 2,
                                                                                    "antes": 2}
    linhas = {l[0]: l for l in I.carregar(tmp_path)["historico"]["linhas"]}
    assert linhas["2|B"][1:] == ["2026-W43", "2026-W43", 1, "2026-W43"]       # saiu da W42, ficou na W43


def test_o_que_barra_a_publicacao(config, origem, agora, tmp_path):
    rid = _rodada(config, origem)
    assert PB.motivos_para_nao_publicar(config, rid) == []
    # a semana que já acabou
    assert any("já acabou" in m for m in PB.motivos_para_nao_publicar(config, rid, agora + timedelta(days=4)))
    # a rodada com a foto do Fracttal é prova: o nome nem entra na lista das rodadas
    assert "rodada inválida" in PB.motivos_para_nao_publicar(config, rid + "-foto")[0]
    # bloco fora da semana: a mesma planilha gerada como se fosse a W42
    st = G.ler_status(config, rid)
    st["conferencia"]["fora_da_semana"] = 2
    G._gravar(G.pasta_trabalho(config) / "geracoes" / rid, st)
    assert any("fora da semana" in m for m in PB.motivos_para_nao_publicar(config, rid))
    with pytest.raises(PB.PublicacaoErro):
        PB.publicar(config, rid, "x")
    # a rodada sem a conferência (de antes de 09/10) não se publica: a tela não pode dizer que conferiu
    sem = dict(st)
    sem.pop("conferencia")
    G._gravar(G.pasta_trabalho(config) / "geracoes" / rid, sem)
    assert any("não tem a conferência" in m for m in PB.motivos_para_nao_publicar(config, rid))
    # sem token não publica (e nada vai ao GitHub)
    st["conferencia"]["fora_da_semana"] = 0
    G._gravar(G.pasta_trabalho(config) / "geracoes" / rid, st)
    with pytest.raises(PB.PublicacaoErro, match="sem token"):
        PB.publicar(config, rid, "x")


def test_conferencia_que_falha_nao_marca_publicada(config, origem, agora):
    rid = _rodada(config, origem)
    _com_github(config, GitHubFalso(estragar_leitura=True))
    with pytest.raises(PB.PublicacaoErro, match="não tem o arquivo enviado"):
        PB.publicar(config, rid, "x")
    assert "publicacao" not in G.ler_status(config, rid)
    assert "2026-W41" not in HB.semanas_da_reserva(I.carregar(G.pasta_trabalho(config)).get("historico"))


def test_tela_publicar_so_administrador_e_com_confirmacao(app, logado, config, origem, monkeypatch):
    monkeypatch.setattr(PB, "_agora", lambda: SABADO_ANTES)        # a W41 ainda não começou: o caso de toda semana
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    rid = _rodada(app.config, origem)
    gh = _com_github(app.config, GitHubFalso({"Programação Semana 41.xlsx": b"a do PC do PCM",
                                              "Observacoes_Semana_Atual.txt": b"15448; nao\n"}))
    html = logado.get(f"/t/pcm/gerar?rodada={rid}").get_data(as_text=True)
    assert f'href="/t/pcm/gerar/publicar?rodada={rid}"' in html and "Publicar no App" in html
    html = logado.get(f"/t/pcm/gerar/publicar?rodada={rid}").get_data(as_text=True)
    assert "Programação Semana 41.xlsx" in html and "Já existe uma Programação Semana 41.xlsx" in html
    assert "15448; nao" in html and 'name="substituir"' in html
    # sem as confirmações não publica
    r = logado.post("/t/pcm/gerar/publicar", data={"rodada": rid, "repo_ja_tem": "1", "confirmo": "1"})
    assert r.status_code == 400 and not gh.commits
    r = logado.post("/t/pcm/gerar/publicar", data={"rodada": rid, "repo_ja_tem": "1", "confirmo": "1", "substituir": "1"})
    assert r.status_code == 303 and "feito=publicada" in r.headers["Location"] and len(gh.commits) == 2
    html = logado.get(r.headers["Location"]).get_data(as_text=True)
    assert "Publicada no App" in html and "Semana publicada no repositório do PCM" in html
    assert "já começou" not in html
    # a semana que já começou pede uma confirmação a mais: publicar troca a programação no meio dela
    monkeypatch.setattr(PB, "_agora", lambda: QUINTA_W41)
    html = logado.get(f"/t/pcm/gerar/publicar?rodada={rid}").get_data(as_text=True)
    assert "A semana 2026-W41 já começou" in html and 'name="ciente_comecou"' in html
    r = logado.post("/t/pcm/gerar/publicar", data={"rodada": rid, "repo_ja_tem": "1", "confirmo": "1", "substituir": "1"})
    assert r.status_code == 400 and len(gh.commits) == 2
    # quem não é administrador não publica
    c = app.test_client()
    with c.session_transaction() as s:
        s["logado"], s["admin"] = True, False
        s["usuario"] = {"email": "pessoa.teste@exemplo.com", "nome": "Pessoa Teste", "perfil": "PCM"}
    r = c.post("/t/pcm/gerar/publicar", data={"rodada": rid, "confirmo": "1", "substituir": "1"})
    assert r.status_code == 403 and "Só administrador publica" in r.get_data(as_text=True) and len(gh.commits) == 2


# ── o plano publicado no histórico ───────────────────────────────────────────────────────────────────────────
def _planilha_bytes(dia: str, linhas=((15020, "ALT1-INV01"), (15021, "ALT1-CAB01"))) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SP Leste 02"
    ws.append(CAB)
    for os_, cod in linhas:
        ws.append(["SP Leste 02", dia, os_, "U", cod, "Corretiva", "t", "Não", 1, "07:30", "09:00", 1.5, 1])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_plano_publicado_vem_do_repositorio_do_pcm():
    repo = {"Programação Semana 42.xlsx": _planilha_bytes("Segunda-feira (12/10)"),
            "Programação Semana 43.xlsx": _planilha_bytes("Segunda-feira (20/10)")}      # a W43 de 2025
    planos = HB.planos_publicados(["2026-W42", "2026-W43"], ler_repo=repo.get)
    assert planos["2026-W42"] == (HB.ORIGEM_REPO + " (Programação Semana 42.xlsx)", {"15020|ALT1-INV01", "15021|ALT1-CAB01"})
    # a "Semana 43" do repositório é de 2025 (a segunda era 20/10; em 2026, 20/10 é terça): não serve
    assert "2026-W43" not in planos
    # a reserva vence o repositório; sem rede, segue sem o repositório
    reserva = {"colunas": ["task_key", "first_week", "last_week", "count", "weeks"],
               "linhas": [["9|Z", "2026-W42", "2026-W42", 1, "2026-W42"]]}
    assert HB.planos_publicados(["2026-W42"], reserva, ler_repo=repo.get)["2026-W42"][0] == HB.ORIGEM_RESERVA

    def sem_rede(nome):
        raise ConnectionError("fora")
    assert HB.planos_publicados(["2026-W42"], ler_repo=sem_rede) == {}


# ── feriados do ano seguinte ─────────────────────────────────────────────────────────────────────────────────
def test_pascoa_e_os_feriados_moveis_do_ano_seguinte():
    assert [I.pascoa(a) for a in (2025, 2026, 2027, 2028)] == [date(2025, 4, 20), date(2026, 4, 5), date(2027, 3, 28),
                                                               date(2028, 4, 16)]

    def linha(tipo, uf, d, nome, mun=None):
        r = {"tipo": tipo, "estado": uf, "data": {"$dt": f"{d}T00:00:00"}, "mes": int(d[5:7]), "feriado": nome}
        if mun:
            r["municipio"] = mun
        return r
    fer = {"titulo": "CONTROLE DE FERIADOS",
           "gerais": [linha("NACIONAL", "TODOS", "2026-01-01", "Confraternização Universal"),
                      linha("NACIONAL", "TODOS", "2026-02-17", "Carnaval"),
                      linha("NACIONAL", "TODOS", "2026-04-03", "Paixão de Cristo"),
                      linha("NACIONAL", "TODOS", "2026-06-04", "Corpus Christi (ponto facultativo)"),
                      linha("ESTADUAL", "ES", "2026-04-13", "FERIADO ESTADUAL"),
                      linha("ESTADUAL", "SP", "2026-07-09", "FERIADO ESTADUAL")],
           "municipais": [linha("MUNICIPAL", "CE", "2026-04-13", "Feriado Municipal", "FORTALEZA")]}
    p = I.projetar_feriados(fer, 2027)
    datas = {(r["feriado"], r["estado"]): I._data(r["data"]) for r in p["gerais"]}
    assert datas[("Confraternização Universal", "TODOS")] == date(2027, 1, 1)
    assert datas[("Carnaval", "TODOS")] == date(2027, 2, 9)
    assert datas[("Paixão de Cristo", "TODOS")] == date(2027, 3, 26)
    assert datas[("Corpus Christi (ponto facultativo)", "TODOS")] == date(2027, 5, 27)
    assert datas[("FERIADO ESTADUAL", "ES")] == date(2027, 4, 5)        # Nossa Senhora da Penha: Páscoa + 8
    assert datas[("FERIADO ESTADUAL", "SP")] == date(2027, 7, 9)
    # o municipal fica na data (a planilha não diz qual seria móvel): o aniversário de Fortaleza é 13/04
    assert I._data(p["municipais"][0]["data"]) == date(2027, 4, 13)
    # o ano que a planilha já tem não se projeta (o dado de verdade vence)
    com, anos = I.com_feriados_ate({**fer, "gerais": fer["gerais"] + p["gerais"]}, 2027)
    assert anos == [] and len(com["gerais"]) == 12


def test_a_w53_de_2026_leva_o_feriado_de_2027(tmp_path):
    from test_pcm_insumos import criar_pasta_pcm
    origem_ = criar_pasta_pcm(tmp_path / "pcm")
    trab, rod = tmp_path / "trabalho", tmp_path / "rodada"
    rod.mkdir()
    I.importar(trab, origem_, "2026-W53")
    carimbos = I.materializar(trab, rod, "2026-W53")
    nac, _est, _mun = ler_feriados(rod / "Feriados" / I.FERIADOS_ARQ)
    assert date(2027, 1, 1) in nac
    fer = next(c for c in carimbos if c["nome"] == I.NOMES["feriados"])
    assert "2027" in fer["detalhe"] and "projetado" in fer["detalhe"]


# ── observações do repositório e a OS que estoura as horas ───────────────────────────────────────────────────
def test_importar_observacoes_do_repositorio(app, logado, config):
    texto = "@usina Altair = seg, qua, sex\n13203; não\n15120; qui; MPM; tarde; só: Sistema de Segurança\n"
    app.config.update({k: v for k, v in config.items() if k != "TESTING"},
                      NEXUS_PCM_REPO_TESTE={"Observacoes_Semana.txt": texto.encode("utf-8")}.get)
    r = logado.post("/t/pcm/gerar/importar-repositorio", data={"semana": "2026-W43"})
    assert r.status_code == 303 and "feito=importado_repo" in r.headers["Location"]
    obs = I.carregar(G.pasta_trabalho(app.config))["observacoes"]["2026-W43"]
    assert obs["texto"] == texto and obs["autor"] == "repositório do PCM" and "Observacoes_Semana.txt" in obs["origem"]["arquivo"]
    app.config["NEXUS_PCM_REPO_TESTE"] = {}.get
    assert logado.post("/t/pcm/gerar/importar-repositorio", data={"semana": "2026-W43"}).status_code == 404


def test_conferencia_aponta_a_os_forcada_alem_das_horas(tmp_path):
    """A W42 de 09/10/2026: a OS 14331 com 4 tarefas de 6 h (a duração estimada no Fracttal) forçadas numa quarta."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PR Norte 01"
    ws.append(CAB)
    for i in range(4):
        ws.append(["PR Norte 01", "Quarta-feira (07/10) [EXCEDE HH]", 14331, "U", f"DJMT{i + 1}", "Corretiva", "Teste",
                   "Não", 1, "07:30", "13:30", 6, 9])
    ws.append(["PR Norte 01", "Quinta-feira (08/10)", 900, "U", "X", "MPM", "t", "Não", 1, "07:30", "08:30", 1, 9])
    wb.create_sheet("_Pendentes").append(["Equipe", "OSs ID", "Código Equipamento", "Tarefa", "Motivo"])
    wb.save(tmp_path / "s.xlsx")
    c = C.conferir(tmp_path / "s.xlsx", "2026-W41")
    assert c["excede_hh_os"] == [{"os": "14331", "equipe": "PR Norte 01", "dia": "qua", "tarefas": 4, "horas": 24.0}]
