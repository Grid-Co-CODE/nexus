"""Etapa 3: os insumos do motor do PCM moram no Nexus. A cada geração o Nexus escreve os arquivos no formato que o
motor lê, e o motor (cópia idêntica) não muda.

A prova que importa: o que o motor lê do arquivo que o Nexus gera é IGUAL ao que ele lia do arquivo original. As
leituras abaixo são as do programacao_v7.py (pandas, mesma aba, mesma linha de cabeçalho, mesmas colunas). Os arquivos
são de mentira, no formato dos reais: o repositório é público e não leva dado do PCM. A mesma prova rodou com os
arquivos reais da pasta do PCM em 30/09/2026 e deu tudo igual.
"""
from datetime import date, datetime

import openpyxl
import pandas as pd
import pytest

from nexus.pcm import insumos as I

CAB_PRIO = ["#\nPrioridade", "Atividade (Agrupada)", "Tipo de Manutenção", "Ativo Principal", "Impacto\nna Geração",
            "Urgência", "Qtd. OS\nHistóricas", "MTTR\nMediana (h)", "Taxa\nResolução %", "Tendência\n(Obs.)",
            "Descrição / O que fazer"]
CAB_CONF = ["Categoria", "Tipo de Ativo", "Total OS", "OS c/ Tempo", "Outliers Removidos", "OS Usadas", "Média (h)",
            "Mediana (h)", "Mín–Máx (h)"]


def criar_pasta_pcm(o):
    """A pasta do PCM de mentira, com os cinco arquivos no formato dos de verdade."""
    o.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "📋 Lista de Prioridades"
    ws.append(["LISTA DE ATIVIDADES POR PRIORIDADE"])
    ws.append(["Base: OS de teste"])
    ws.append([])
    ws.append(["Prioridade 1–7: Emergencial"])
    ws.append([])
    ws.append(["Total de Atividades: 2"])
    ws.append([])
    ws.append(CAB_PRIO)
    ws.append([1, "Religamento de Usina", "Corretiva Emergencial", "Usina completa", "Total", "Imediata", 1214, 0.36,
               0.99, None, "Trip de proteção"])
    ws.append([2, "Falha de inversor", "Corretiva", "Inversor", "Parcial", "Alta", 300, 2.5, 0.9, "sobe", "Trocar"])
    wb.create_sheet("📊 Resumo por Grupo").append(["não lido pelo motor"])
    wb.save(o / "Lista_Prioridades_GridCo.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo por Categoria"
    ws.append(["Análise de Manutenção Corretiva"])
    ws.append(["Outliers removidos pelo método IQR"])
    ws.append([])
    ws.append(CAB_CONF)
    ws.append(["Inspeção de Inversor", "Inversor", 71, 58, 6, 52, 1.45, 0.72, "0.09–5.48"])
    ws.append(["Troca de fusível", "String", 40, 30, 2, 28, 0.8, 0.6, "0.2–2.0"])
    ws.append(["TOTAL", None, 111, 88, 8, 80, None, None, None])
    wb.save(o / "Planilha Confiabilidade R00.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["task_key", "first_week", "last_week", "count", "weeks"])
    ws.append(["6566|IRC100-INFC1", "2026-W20", "2026-W21", 2, "2026-W20,2026-W21"])
    ws.append(["15021|ALT1-CAB01", "2026-W39", "2026-W40", 2, "2026-W39,2026-W40"])
    wb.save(o / "Historico_Programacoes.xlsx")

    (o / "Feriados").mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feriados 2026"
    ws.append(["CONTROLE DE FERIADOS"])
    ws.append(I.FER_ESQ + [None] + I.FER_DIR)
    ws.append(["NACIONAL", "TODOS", datetime(2026, 1, 1), 1, "Confraternização", None,
               "MUNICIPAL", "RJ", "CAMPOS DOS GOYTACAZES", datetime(2026, 1, 15), 1, "Feriado Municipal"])
    ws.append(["ESTADUAL", "SP", datetime(2026, 7, 9), 7, "Revolução", None,
               "MUNICIPAL", "MA", "MATÕES", datetime(2026, 1, 20), 1, "Feriado Municipal"])
    ws.append([None] * 6 + ["MUNICIPAL", "CE", "QUIXADÁ", datetime(2026, 10, 2), 10, "Feriado Municipal"])
    wb.create_sheet("MUNICIPIOS + ESTADOS")
    wb.save(o / "Feriados" / I.FERIADOS_ARQ)

    (o / "Observacoes_Semana.txt").write_text("# comentário\n10369; qua\n\nsem: 15021\n", encoding="utf-8")
    return o


# ── as leituras do motor (programacao_v7.py, linhas 168-192, 217-256, 798-813 e 929-938) ─────────────────────

def ler_prioridades(p):
    raw = pd.read_excel(p, sheet_name="📋 Lista de Prioridades", header=7)
    df = raw.iloc[:, [0, 1, 2, 3, 6, 7]].copy()
    df.columns = ["prio", "atividade", "tipo", "ativo_principal", "qtd_hist", "mttr_h"]
    df = df.dropna(subset=["prio"])
    df["prio"] = pd.to_numeric(df["prio"], errors="coerce")
    df["mttr_h"] = pd.to_numeric(df["mttr_h"], errors="coerce")
    return df.dropna(subset=["prio"]).reset_index(drop=True)


def ler_confiab(p):
    df = pd.read_excel(p, sheet_name="Resumo por Categoria", header=3)
    df = df.dropna(subset=["Categoria", "Média (h)"]).copy()
    df = df[df["Categoria"].astype(str).str.upper() != "TOTAL"]
    df["Média (h)"] = pd.to_numeric(df["Média (h)"], errors="coerce")
    return df.dropna(subset=["Média (h)"]).reset_index(drop=True)


def ler_historico(p):
    return {str(r["task_key"]): (r["first_week"], r["last_week"], int(r["count"]), str(r["weeks"]))
            for _, r in pd.read_excel(p).iterrows()}


def ler_feriados(p):
    df = pd.read_excel(p, sheet_name="Feriados 2026", header=1)
    nac, est, mun = set(), {}, {}
    for _, row in df.iterrows():
        tipo, uf, d = str(row.get("Tipo") or "").upper(), str(row.get("Estado") or "").upper(), row.get("Data")
        if pd.isna(d):
            continue
        if tipo == "NACIONAL":
            nac.add(d.date())
        elif tipo == "ESTADUAL" and uf and uf != "TODOS":
            est.setdefault(d.date(), set()).add(uf)
        if "MUNICIP" in str(row.get("Tipo.1") or "").upper():
            mun.setdefault(row["Data.1"].date(), set()).add(row["MUNICIPIO"])
    for _, row in df.iterrows():          # o bloco da direita sozinho numa linha (a da esquerda vazia)
        if pd.isna(row.get("Data")) and "MUNICIP" in str(row.get("Tipo.1") or "").upper():
            mun.setdefault(row["Data.1"].date(), set()).add(row["MUNICIPIO"])
    return nac, est, mun


def so_do_ano(fer, ano):
    """(nacionais, estaduais, municipais) de `ler_feriados` só com as datas do ano."""
    nac, est, mun = fer
    return ({d for d in nac if d.year == ano}, {d: v for d, v in est.items() if d.year == ano},
            {d: v for d, v in mun.items() if d.year == ano})


def ler_obs(p):
    return [(i, l.strip()) for i, l in enumerate(p.read_text(encoding="utf-8").splitlines(), 1)
            if l.strip() and not l.strip().startswith("#")]


@pytest.fixture
def pcm(tmp_path):
    origem = criar_pasta_pcm(tmp_path / "pcm")
    trab, rod = tmp_path / "trabalho", tmp_path / "rodada"
    rod.mkdir()
    return origem, trab, rod


def guardar_reserva_do_arquivo(trab, origem):
    """A reserva do histórico como ela ficou no Nexus (importada do arquivo do PC até 08/10; desde então o histórico
    vem do banco e o importar não a toca)."""
    d = I.carregar(trab)
    d["historico"] = I.ler_planilha(origem / "Historico_Programacoes.xlsx", None, 1)
    I.salvar(trab, d)


def test_ida_e_volta_o_motor_le_igual(pcm):
    origem, trab, rod = pcm
    resumo = I.importar(trab, origem, "2026-W41")
    assert "não importado" in resumo["historico"] and "3 municipais" in resumo["feriados"]
    guardar_reserva_do_arquivo(trab, origem)
    I.materializar(trab, rod, "2026-W41")
    pd.testing.assert_frame_equal(ler_prioridades(origem / "Lista_Prioridades_GridCo.xlsx"),
                                  ler_prioridades(rod / "Lista_Prioridades_GridCo.xlsx"))
    pd.testing.assert_frame_equal(ler_confiab(origem / "Planilha Confiabilidade R00.xlsx"),
                                  ler_confiab(rod / "Planilha Confiabilidade R00.xlsx"))
    assert ler_historico(origem / "Historico_Programacoes.xlsx") == ler_historico(rod / "historico_do_nexus.xlsx")
    fa = ler_feriados(origem / "Feriados" / I.FERIADOS_ARQ)
    fb = ler_feriados(rod / "Feriados" / I.FERIADOS_ARQ)
    # o ano da planilha sai igual; o seguinte vai projetado (09/10/2026: a W53 de 2026 termina em 01/01/2027)
    assert so_do_ano(fb, 2026) == fa and len(fb[2]) == 3 + len(so_do_ano(fb, 2027)[2])   # Quixadá veio, mesmo sozinha
    assert date(2027, 1, 1) in fb[0]
    assert ler_obs(origem / "Observacoes_Semana.txt") == ler_obs(rod / "Observacoes_Semana.txt")


def test_dado_mudado_no_nexus_e_o_que_o_motor_recebe(pcm):
    """Editar no Nexus muda o que o motor lê: a prova não passa por acaso."""
    origem, trab, rod = pcm
    I.importar(trab, origem, "2026-W41")
    d = I.carregar(trab)
    d["confiabilidade"]["linhas"][0][6] = 9.99
    I.salvar(trab, d)
    I.materializar(trab, rod, "2026-W41")
    assert ler_confiab(rod / "Planilha Confiabilidade R00.xlsx")["Média (h)"][0] == 9.99


def test_observacoes_sao_por_semana_e_editaveis(pcm):
    origem, trab, rod = pcm
    I.importar(trab, origem, "2026-W41")
    I.salvar_observacoes(trab, "2026-W42", "10400; seg\r\n", autor="admin")
    assert I.observacoes(trab, "2026-W42") == "10400; seg\n"
    I.materializar(trab, rod, "2026-W42")
    assert (rod / "Observacoes_Semana.txt").read_text(encoding="utf-8") == "10400; seg\n"
    I.materializar(trab, rod, "2026-W43")                    # semana sem observação: arquivo vazio, o motor aceita
    assert (rod / "Observacoes_Semana.txt").read_text(encoding="utf-8") == ""


def test_semana_sem_observacao_entrega_ao_motor_os_dias_herdados(pcm):
    """O padrão da tela é a última programação (05/10/2026): o motor tem de receber os MESMOS dias que a tela mostra."""
    origem, trab, rod = pcm
    I.importar(trab, origem, "2026-W41")
    I.salvar_observacoes(trab, "2026-W41", "@usina Marabá 1 = seg, qua\n13480; não\n")
    carimbos = I.materializar(trab, rod, "2026-W42")
    assert (rod / "Observacoes_Semana.txt").read_text(encoding="utf-8") == "@usina Marabá 1 = seg, qua\n"
    obs = [c for c in carimbos if c["nome"] == I.NOMES["observacoes"]][0]
    assert "herdad" in obs["detalhe"] and "2026-W41" in obs["detalhe"]


def test_sem_importar_nao_gera(pcm):
    _origem, trab, rod = pcm
    with pytest.raises(I.InsumoErro) as erro:
        I.materializar(trab, rod, "2026-W41")
    assert "Importe da pasta do PCM" in str(erro.value) and I.NOMES["historico"] not in str(erro.value)


def test_importar_nao_traz_o_historico_do_pc_por_cima(pcm):
    """08/10/2026 ("PUXE O HISTÓRICO"): o arquivo do PC guarda a união de todas as gerações (a W41 com as duas de 02/10:
    1.286 tarefas, 167 que nunca foram ao campo). O histórico vem do banco; a reserva corrigida que está no Nexus não
    pode ser trocada pelo arquivo do PC num clique de importar."""
    origem, trab, rod = pcm
    I.importar(trab, origem, "2026-W41")
    assert "historico" not in I.carregar(trab)
    d = I.carregar(trab)
    d["historico"] = {"aba": None, "preambulo": [], "colunas": ["task_key", "first_week", "last_week", "count", "weeks"],
                      "linhas": [["1|A", "2026-W41", "2026-W41", 1, "2026-W41"]], "meta": {"correcao": "S41"}}
    I.salvar(trab, d)
    I.importar(trab, origem, "2026-W41")
    assert I.carregar(trab)["historico"]["linhas"] == [["1|A", "2026-W41", "2026-W41", 1, "2026-W41"]]
    # a reserva vai para a rodada com outro nome, sem carimbo (o carimbo do histórico é da geração)
    carimbos = I.materializar(trab, rod, "2026-W42")
    assert (rod / "historico_do_nexus.xlsx").exists() and I.NOMES["historico"] not in [c["nome"] for c in carimbos]


def test_estado_avisa_quando_o_arquivo_da_pasta_mudou(pcm):
    """Na sombra, o Nexus precisa partir do mesmo dado que o PCM: se alguém editou o arquivo depois da importação,
    a tela avisa para importar de novo."""
    origem, trab, _rod = pcm
    I.importar(trab, origem, "2026-W41")
    assert not any(i["aviso"] for i in I.estado(trab, origem, "2026-W41"))
    wb = openpyxl.load_workbook(origem / "Planilha Confiabilidade R00.xlsx")
    wb.active["G5"] = 3.3
    wb.save(origem / "Planilha Confiabilidade R00.xlsx")
    itens = {i["nome"]: i for i in I.estado(trab, origem, "2026-W41")}
    assert itens[I.NOMES["confiabilidade"]]["aviso"] and "importe de novo" in itens[I.NOMES["confiabilidade"]]["detalhe"]
    assert not itens[I.NOMES["prioridades"]]["aviso"]


def test_gravacao_guarda_a_versao_anterior(pcm):
    origem, trab, _rod = pcm
    I.importar(trab, origem, "2026-W41")
    I.salvar_observacoes(trab, "2026-W41", "nova")
    anterior = (trab / "insumos.json.anterior").read_text(encoding="utf-8")
    assert "10369; qua" in anterior and I.observacoes(trab, "2026-W41") == "nova"


@pytest.fixture
def tela(app, pcm):
    origem, trab, _rod = pcm
    app.config.update(NEXUS_PCM_ORIGEM=str(origem), NEXUS_PCM_TRABALHO=str(trab),
                      NEXUS_PCM_FRACTTAL_TESTE={"FRACTTAL_CLIENT_ID": "a", "FRACTTAL_CLIENT_SECRET": "b"})
    return origem, trab


def test_tela_importa_da_pasta_do_pcm(logado, tela):
    _origem, trab = tela
    html = logado.get("/t/pcm/gerar?semana=2026-W41").get_data(as_text=True)
    assert "ainda não está no Nexus" in html and "Importar da pasta do PCM" in html
    resp = logado.post("/t/pcm/gerar/importar", data={"semana": "2026-W41"})
    assert resp.status_code == 303 and "semana=2026-W41" in resp.headers["Location"]
    html = logado.get(resp.headers["Location"]).get_data(as_text=True)
    assert "ainda não está no Nexus" not in html and "Importado da pasta do PCM" in html
    # o histórico não é importado: a tela diz que ele vem do banco
    assert "historico" not in I.carregar(trab) and "do banco a cada geração" in html


def test_tela_edita_as_observacoes_da_semana(logado, tela):
    _origem, trab = tela
    resp = logado.post("/t/pcm/gerar/observacoes", data={"semana": "2026-W42", "texto": "10400; seg\nsem: 15021"})
    assert resp.status_code == 303
    assert I.observacoes(trab, "2026-W42") == "10400; seg\nsem: 15021"
    html = logado.get("/t/pcm/gerar?semana=2026-W42").get_data(as_text=True)
    assert "10400; seg" in html and "2 para a semana 2026-W42" in html


def test_tela_recusa_semana_invalida_nas_observacoes(logado, tela):
    assert logado.post("/t/pcm/gerar/observacoes", data={"semana": "../x", "texto": "a"}).status_code == 400
    assert logado.post("/t/pcm/gerar/importar", data={"semana": "x"}).status_code == 400


def test_cabecalho_dos_feriados_mudado_e_recusado(pcm, tmp_path):
    origem, trab, _rod = pcm
    wb = openpyxl.load_workbook(origem / "Feriados" / I.FERIADOS_ARQ)
    wb.active["B2"] = "UF"
    wb.save(origem / "Feriados" / I.FERIADOS_ARQ)
    with pytest.raises(I.InsumoErro) as erro:
        I.importar(trab, origem, "2026-W41")
    assert "cabeçalho" in str(erro.value)
