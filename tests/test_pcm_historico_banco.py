"""O histórico das programações sai do BANCO (08/10/2026, Levi: "PUXE O HISTÓRICO"), no formato que o motor lê.

A prova que importa é a do motor: o arquivo escrito pelo Nexus, lido com as mesmas linhas do programacao_v7.py
(926-938), dá a contagem certa por chave. Dado de mentira, no formato do fato do banco (a API devolve número como
texto às vezes). Os números reais (W21-W41) foram comparados fora do repositório: ver nexus/pcm/CLAUDE.md.
"""
from datetime import datetime, timedelta, timezone

import openpyxl
import pandas as pd
import pytest

from nexus.pcm import historico_banco as HB

BRT = timezone(timedelta(hours=-3))


def bloco(semana_seg, os_, codigo, lido="2026-10-08T17:05:32-03:00"):
    """Uma linha do fato_programacao como o banco devolve (só as colunas que o histórico usa)."""
    return {"data_id_semana": semana_seg, "os": os_, "codigo_ativo": codigo, "lido_em": lido}


def ler_como_o_motor(p):
    """programacao_v7.py, linhas 926-938: pd.read_excel do arquivo, a chave e a contagem de cada tarefa."""
    hist = {}
    for _, r in pd.read_excel(p).iterrows():
        hist[str(r["task_key"])] = {"first_week": r.get("first_week", ""), "last_week": r.get("last_week", ""),
                                    "count": int(r.get("count", 0)), "weeks": str(r.get("weeks", ""))}
    return hist


def reserva(linhas):
    return {"aba": None, "preambulo": [], "colunas": HB.CABECALHO, "linhas": linhas}


# ── a chave e as semanas ─────────────────────────────────────────────────────────────────────────────────────
def test_chave_e_a_do_motor():
    # programacao_v7.task_key: f"{int(os_id)}|{str(codigo).strip()}"; o Fracttal manda a OS como texto ("7975")
    assert HB.chave_da_tarefa("15071", "PTL300-SSEG1") == "15071|PTL300-SSEG1"
    assert HB.chave_da_tarefa(15071.0, " PTL300-SSEG1 ") == "15071|PTL300-SSEG1"
    assert HB.chave_da_tarefa("15071.0", "X") == "15071|X"
    assert HB.chave_da_tarefa("abc", "X") is None and HB.chave_da_tarefa("1", "") is None


def test_semanas_do_banco_aceitam_texto_e_contam_a_chave_uma_vez():
    linhas = [bloco("20261005", "15071", "A"), bloco(20261005, 15071, "A"),     # a mesma tarefa em dois blocos
              bloco("20260928", "15071", "A"), bloco("20261005", "x", "B")]
    sem, cont = HB.semanas_do_banco(linhas)
    assert sem == {"2026-W41": {"15071|A"}, "2026-W40": {"15071|A"}}
    assert cont == {"blocos": 4, "sem_chave": 1}


def test_fim_da_semana_e_a_segunda_seguinte_em_brasilia():
    assert HB.fim_da_semana("2026-W41") == datetime(2026, 10, 12, tzinfo=BRT)
    assert HB.fim_da_semana("2026-W52") == datetime(2026, 12, 28, tzinfo=BRT)
    assert HB.fim_da_semana("2027-W01") == datetime(2027, 1, 11, tzinfo=BRT)     # a semana ISO, não o ano civil


# ── a montagem ───────────────────────────────────────────────────────────────────────────────────────────────
DO_BANCO = {"2026-W39": {"1|A", "2|B"}, "2026-W40": {"1|A", "3|C"}, "2026-W41": {"1|A", "4|D"}}
FECHADO = datetime(2026, 10, 13, 1, 0, tzinfo=BRT)          # o banco já sabia do fim da W41
NO_MEIO = datetime(2026, 10, 8, 17, 5, tzinfo=BRT)          # gravado numa quinta da W41


def test_so_entram_as_semanas_de_antes_da_gerada():
    """Gerar a mesma semana de novo não pode contá-la (S40 publicada 2x em 25/09: 415 tarefas viraram reprogramadas)."""
    linhas, rel = HB.montar(DO_BANCO, "2026-W41", FECHADO)
    hist = {l[0]: l for l in linhas}
    assert hist["1|A"] == ["1|A", "2026-W39", "2026-W40", 2, "2026-W39,2026-W40"]
    assert "4|D" not in hist and set(rel["semanas"]) == {"2026-W39", "2026-W40"}
    assert [l[0] for l in linhas] == sorted(hist)                  # ordenado: o mesmo dado dá o mesmo arquivo


def test_semana_fechada_no_banco_vale_o_banco():
    linhas, rel = HB.montar(DO_BANCO, "2026-W42", FECHADO, {"2026-W41": ("plano", {"9|Z"})})
    hist = {l[0]: l for l in linhas}
    assert hist["1|A"][3] == 3 and "4|D" in hist and "9|Z" not in hist
    assert rel["semanas"]["2026-W41"]["origem"] == HB.ORIGEM_BANCO and not rel["avisos"]


def test_semana_em_andamento_vale_o_plano_publicado():
    """A W41 do banco de 08/10 era a de uma quinta: à noite o robô do PCM tira da agenda o que rolou (247 tarefas de OS
    vivas) e só refaz a semana quando ela fecha. O plano publicado é o que o motor contaria."""
    planos = {"2026-W41": (HB.ORIGEM_RESERVA, {"1|A", "5|E"})}
    linhas, rel = HB.montar(DO_BANCO, "2026-W42", NO_MEIO, planos)
    hist = {l[0]: l for l in linhas}
    assert "5|E" in hist and "4|D" not in hist and hist["1|A"][3] == 3
    assert rel["semanas"]["2026-W41"] == {"origem": HB.ORIGEM_RESERVA, "chaves": 2, "no_banco": 2}


def test_semana_em_andamento_sem_plano_fica_com_o_banco_e_avisa():
    linhas, rel = HB.montar(DO_BANCO, "2026-W42", NO_MEIO, {})
    assert "4|D" in {l[0] for l in linhas}
    assert rel["semanas"]["2026-W41"]["origem"] == HB.ORIGEM_ANDAMENTO and rel["avisos"]


def test_semana_que_o_banco_ainda_nao_tem_vem_do_plano():
    """Gerar a W43 com o banco parado em 08/10: a W42 não está nele, e o plano publicado dela entra."""
    planos = {"2026-W42": (HB.ORIGEM_OFICIAL, {"6|F"}), "2026-W41": (HB.ORIGEM_RESERVA, {"1|A"})}
    assert HB.semanas_sem_fechar(DO_BANCO, NO_MEIO, "2026-W43") == ["2026-W41", "2026-W42"]
    linhas, rel = HB.montar(DO_BANCO, "2026-W43", NO_MEIO, planos)
    assert "6|F" in {l[0] for l in linhas} and rel["ate"] == "2026-W42"
    _l, sem_plano = HB.montar(DO_BANCO, "2026-W43", NO_MEIO, {})
    assert sem_plano["semanas"]["2026-W42"]["chaves"] == 0 and sem_plano["avisos"]


def test_semana_anterior_ao_banco_nao_entra():
    """O banco começa na W21: a W20 do arquivo do PC não volta pela reserva."""
    linhas, rel = HB.montar(DO_BANCO, "2026-W42", FECHADO, {"2026-W20": ("reserva", {"7|G"})})
    assert "7|G" not in {l[0] for l in linhas} and rel["de"] == "2026-W39"


def test_o_motor_le_o_arquivo_escrito(tmp_path):
    linhas, _rel = HB.montar(DO_BANCO, "2026-W42", FECHADO)
    HB.escrever(linhas, tmp_path / HB.NOME)
    hist = ler_como_o_motor(tmp_path / HB.NOME)
    assert hist["1|A"] == {"first_week": "2026-W39", "last_week": "2026-W41", "count": 3,
                           "weeks": "2026-W39,2026-W40,2026-W41"}
    assert hist["3|C"]["count"] == 1 and len(hist) == 4
    assert openpyxl.load_workbook(tmp_path / HB.NOME).sheetnames == ["Sheet1"]


# ── as fontes do plano publicado ────────────────────────────────────────────────────────────────────────────
def planilha_oficial(caminho, agendadas, pendentes=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SP Leste 02"
    ws.append(["Equipe", "Dia", "OSs ID", "Código Equipamento", "Tarefa"])
    for os_, cod in agendadas:
        ws.append(["SP Leste 02", "Segunda-feira (05/10)", os_, cod, "t"])
    wp = wb.create_sheet("_Pendentes")
    wp.append(["Equipe", "OSs ID", "Código Equipamento", "Tarefa", "Motivo"])
    for os_, cod in pendentes:
        wp.append(["SP Leste 02", os_, cod, "t", "não coube"])
    wb.create_sheet("_Resumo").append(["OSs ID"])
    wb.save(caminho)


def test_planilha_oficial_da_so_as_agendadas(tmp_path):
    """O motor grava no histórico só as linhas agendadas (all_rows), nunca as pendentes."""
    planilha_oficial(tmp_path / "p.xlsx", [(15020, "ALT1-INV01"), ("15021", "ALT1-CAB01 ")], [(15099, "X")])
    assert HB.chaves_da_planilha(tmp_path / "p.xlsx") == {"15020|ALT1-INV01", "15021|ALT1-CAB01"}


def test_plano_da_reserva_vence_a_planilha_oficial(tmp_path):
    """A reserva é curada (a W41 de lá é a S41 que foi ao campo, corrigida); a planilha da pasta é o último arquivo do
    PCM, que pode ter sido regerado."""
    planilha_oficial(tmp_path / "Programação Semana 41.xlsx", [(1, "A")])
    planilha_oficial(tmp_path / "Programação Semana 42.xlsx", [(2, "B")])
    res = reserva([["9|Z", "2026-W41", "2026-W41", 1, "2026-W41"]])
    planos = HB.planos_publicados(["2026-W41", "2026-W42", "2026-W43"], res, tmp_path)
    assert planos["2026-W41"] == (HB.ORIGEM_RESERVA, {"9|Z"})
    assert planos["2026-W42"][1] == {"2|B"} and "Programação Semana 42.xlsx" in planos["2026-W42"][0]
    assert "2026-W43" not in planos


# ── a leitura do banco e a rodada ────────────────────────────────────────────────────────────────────────────
def test_nos_testes_sem_dado_nao_vai_a_rede():
    with pytest.raises(HB.HistoricoErro):
        HB.ler_do_banco({"TESTING": True})
    with pytest.raises(HB.HistoricoErro):
        HB.ler_do_banco({"TESTING": True, "NEXUS_PCM_HISTORICO_TESTE": RuntimeError("500")})


def test_materializar_escreve_e_relata(tmp_path):
    cfg = {"TESTING": True, "NEXUS_PCM_HISTORICO_TESTE": {
        "linhas": [bloco("20260928", "1", "A", "2026-10-08T17:05:32-03:00"),
                   bloco("20261005", "1", "A"), bloco("20261005", "4", "D")],
        "gerado_em": "2026-10-08T17:05:32-03:00"}}
    res = reserva([["1|A", "2026-W41", "2026-W41", 1, "2026-W41"], ["5|E", "2026-W41", "2026-W41", 1, "2026-W41"]])
    rel = HB.materializar(cfg, tmp_path / HB.NOME, "2026-W42", reserva=res)
    hist = ler_como_o_motor(tmp_path / HB.NOME)
    assert hist["1|A"]["count"] == 2 and "5|E" in hist and "4|D" not in hist
    assert rel["semanas"]["2026-W41"]["origem"] == HB.ORIGEM_RESERVA and rel["banco_ate"].startswith("2026-10-08")
    assert "do banco" in rel["detalhe"] and "2026-W41 pelo plano publicado" in rel["detalhe"]


def test_leitura_parcial_do_banco_e_erro():
    """A carga grava a contagem e o sha das linhas na `atualizacao`: leitura que não bate não vira histórico."""
    from nexus.dados import programacao as P
    linhas = [bloco("20261005", "1", "A"), bloco("20261005", "2", "B")]
    gravado = {"linhas": "2", "sha_linhas": P.sha_linhas(P.mesclar_semanas(linhas, []))}
    HB.conferir_leitura(linhas, [gravado])                      # a gravação inteira: passa
    HB.conferir_leitura(linhas, [])                             # sem a linha da gravação: não há com o que conferir
    with pytest.raises(HB.HistoricoErro, match="1 de 2"):
        HB.conferir_leitura(linhas[:1], [gravado])
    with pytest.raises(HB.HistoricoErro, match="sha"):
        HB.conferir_leitura([bloco("20261005", "1", "A"), bloco("20261005", "3", "C")], [gravado])


def test_banco_vazio_e_erro(tmp_path):
    cfg = {"TESTING": True, "NEXUS_PCM_HISTORICO_TESTE": {"linhas": [], "gerado_em": None}}
    with pytest.raises(HB.HistoricoErro):
        HB.materializar(cfg, tmp_path / HB.NOME, "2026-W42")
