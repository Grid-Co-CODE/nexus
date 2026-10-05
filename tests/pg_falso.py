"""A API db_performace (PostgreSQL da T.I.) em memória, para os testes que gravam e leem workbooks.

O sync-xlsx lê o xlsx de verdade (openpyxl), como o servidor faz. `como_texto=True` devolve todo valor como texto, para
provar que a leitura não depende de a API preservar o tipo; `perder=N` some com N linhas na gravação, para provar que a
conferência depois de gravar pega.
"""
import io
import re

import requests


class Resposta:
    def __init__(self, status, corpo):
        self.status_code, self._corpo = status, corpo

    def json(self):
        return self._corpo

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"HTTP {self.status_code}")


class ApiPGFalsa:
    def __init__(self, como_texto=False, perder=0):
        self.workbooks, self.abas = {}, {}
        self.criacoes = self.syncs = 0
        self.como_texto, self.perder = como_texto, perder
        self.tokens = []

    def _id(self, wb, aba):
        if (wb, aba) not in self.abas:
            self.abas[(wb, aba)] = {"id": 100 + len(self.abas), "linhas": []}
        return self.abas[(wb, aba)]

    def get(self, url, params=None, headers=None, timeout=None):
        if url.endswith("/api/workbooks"):
            self.tokens.append((headers or {}).get("Authorization"))
            return Resposta(200, [{"key": k} for k in self.workbooks])
        if url.endswith("/api/sheets"):
            return Resposta(200, [{"id": v["id"], "workbook_key": wb, "sheet_name": aba}
                                  for (wb, aba), v in self.abas.items()])
        m = re.search(r"/api/sheets/(\d+)/rows$", url)
        aba = next(v for v in self.abas.values() if v["id"] == int(m.group(1)))
        ini, lim = int((params or {}).get("offset", 0)), int((params or {}).get("limit", 1000))
        return Resposta(200, aba["linhas"][ini:ini + lim])

    def post(self, url, headers=None, json=None, params=None, files=None, timeout=None):
        self.tokens.append((headers or {}).get("Authorization"))
        if url.endswith("/api/workbooks"):
            self.workbooks[json["key"]] = json
            self.criacoes += 1
            return Resposta(201, {})
        wb = re.search(r"/api/workbooks/([^/]+)/sync-xlsx$", url).group(1)
        assert (params or {}).get("replace") == "true"
        from openpyxl import load_workbook
        livro = load_workbook(io.BytesIO(files["file"][1]))
        for ws in livro.worksheets:
            linhas = list(ws.iter_rows(values_only=True))
            cab, corpo = list(linhas[0]), linhas[1:]
            if ws.title != "atualizacao" and self.perder:
                corpo = corpo[self.perder:]
            conv = (lambda v: None if v is None else str(v)) if self.como_texto else (lambda v: v)
            self._id(wb, ws.title)["linhas"] = [{"headers": cab, "values": [conv(v) for v in l]} for l in corpo]
        self.syncs += 1
        return Resposta(200, {"ok": True})

    def linhas(self, wb, aba):
        return [dict(zip(x["headers"], x["values"])) for x in self.abas.get((wb, aba), {}).get("linhas", [])]
