"""Garantias, os pedidos do Levi de 28/10/2026:
1 - o status geral (Grid, Cliente, Fabricante e Standby), AUTOMÁTICO pelo status: "Atividade Grid: Teste em Campo,
    Aguardando abertura de OS, Aguardando retorno Pré - Operação, Aguardando retorno do Supervisor; Atividade Cliente:
    Aguardando retorno do Cliente; Atividade Fabricante: Aguardando coleta do equipamento defeituoso, Aguardando
    retorno do Fabricante"; "stand by deixa de ser status comum e vira status geral" (escolhido no mesmo seletor);
2 - os avisos do quadro numa exclamação ao lado do título, como em Segurança > Extintores;
3 - sem o botão de voltar para Chamados, no Nexus;
4 - o filtro Status Gerais, e o Responsável UFV (o supervisor da usina) no lugar do Responsável;
5 - Exportar relatório por cliente (um checkbox por cliente), para mandar ao cliente;
6 - Exportar tabela;
7 - sem o Limpar filtros.
E o "Read timed out" das observações no print: o banco da Gridco às vezes leva mais de 20 s (medido em 28/10: 20,2 s
com erro, 21,4 s e 0,8 s); o quadro abre com a última leitura e relê por trás.

Fracttal e banco FALSOS (nada vai à rede). Nomes e números de mentira: o repositório é público."""
import io
import sys
import zipfile

import pytest

from nexus.torres.oscreator import ponte

if ponte.RAIZ_CLONE not in sys.path:
    sys.path.append(ponte.RAIZ_CLONE)

from os_web import acomp_web as aw  # noqa: E402

JWT = "aaa.eyJlbWFpbCI6InRlc3RlQGV4ZW1wbG8uaW52YWxpZCIsImV4cCI6OTk5OTk5OTk5OX0.sig"   # de mentira


# ── 1: o status geral automático ────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("status,geral", [
    ("Testes em campo", "Grid"), ("Aguardando Abertura de OS", "Grid"), ("Aguardando retorno Pré - Operação", "Grid"),
    ("Aguardando retorno do Supervisor", "Grid"), ("Aguardando retorno do Cliente", "Cliente"),
    ("Aguardando coleta do equipamento defeituoso", "Fabricante"), ("Aguardando retorno do Fabricante", "Fabricante"),
    ("Standby", "Standby"), ("", ""), ("Algo da planilha antiga", "")])
def test_o_status_geral_sai_do_status(status, geral):
    assert aw.status_geral(status) == geral


def test_as_listas():
    assert aw.STATUS_GERAL == ("Grid", "Cliente", "Fabricante", "Standby")
    assert set(aw.GERAL_DO_STATUS) == set(aw.STATUS_CHAMADO)              # todo status tem o geral dele


# ── a tela ─────────────────────────────────────────────────────────────────────────────────────────────────────
def _linha(i, folio, cliente="Cliente A", usina="Usina A1", status=1, ativo="Inversor 1.1", marca="Huawei"):
    return {"id": i, "folio": str(folio), "cliente": cliente, "usina": usina, "ativo": ativo, "code": "TST-%d" % i,
            "tipo": "Inversor", "tipo_tarefa": "Administrativa", "note": "", "status_id": status, "status": "",
            "descricao": "[%s] - Acompanhamento de chamado %s" % (ativo, marca), "data": "2026-08-01T13:00:00",
            "data_fim": "", "atribuido_a": "Analista Teste", "id_atribuido": None, "criado_por": "Pessoa Teste"}


LINHAS = [_linha(1, 9101), _linha(2, 9102, usina="Usina A2", marca="STI"), _linha(3, 9103, cliente="Cliente B",
          usina="Usina B1"), _linha(4, 9104, status=3)]


@pytest.fixture
def tela(monkeypatch):
    pytest.importorskip("PyQt6")          # o clone importa o steps/ui.py do OS Creator ao montar a tela de Engenharia
    import api
    import chamados_obs_store as obs_store
    from os_web import criar_app, rotas, rotas_acomp
    with rotas._MEMO_LOCK:
        rotas._MEMO.clear()
    banco = [{"id": "s1", "quando": "2026-10-27T10:00:00-03:00", "os": "9101", "tipo": "status", "texto": "Status: x",
              "quem": "", "email": "", "ativo": "", "extra": {"status": "Aguardando retorno do Cliente"}},
             {"id": "o1", "quando": "2026-10-28T09:00:00-03:00", "os": "9101", "tipo": "obs",
              "texto": "Pedimos ao cliente a declaração de transporte.", "quem": "", "email": "", "ativo": "", "extra": {}},
             {"id": "s2", "quando": "2026-10-27T10:00:00-03:00", "os": "9102", "tipo": "status", "texto": "Status: y",
              "quem": "", "email": "", "ativo": "", "extra": {"status": "Aguardando retorno do Fabricante"}},
             {"id": "s3", "quando": "2026-10-27T10:00:00-03:00", "os": "9103", "tipo": "status", "texto": "Status: z",
              "quem": "", "email": "", "ativo": "", "extra": {"status": "Standby"}}]
    monkeypatch.setattr(api, "_code_to_loc", lambda: {})
    monkeypatch.setattr(api, "list_chamados", lambda **k: [dict(l) for l in LINHAS])
    monkeypatch.setattr(api, "total_chamados", lambda de=None: len(LINHAS))
    monkeypatch.setattr(api, "tickets_os3_em_massa", lambda ids: {i: "SR %d" % (90000 + i) for i in ids})
    monkeypatch.setattr(obs_store, "listar", lambda forcar=False: list(banco))
    monkeypatch.setattr(rotas_acomp, "RESPONSAVEL_UFV",
                        lambda cards: {c["folio"]: {"9101": "Supervisor Um", "9102": "Supervisor Dois"}.get(c["folio"], "")
                                       for c in cards})
    c = criar_app(segredo="teste", testing=True).test_client()
    with c.session_transaction() as s:
        s["jwt"] = JWT
        s["conta"] = {"nome": "Analista Teste", "email": "teste@exemplo.invalid", "perfil": ""}
    yield c, banco, rotas_acomp
    with rotas._MEMO_LOCK:
        rotas._MEMO.clear()


def _quadro(c):
    return c.get("/os/chamados/acompanhamento").get_data(as_text=True)


def test_o_cartao_e_a_linha_levam_o_geral_e_o_responsavel_ufv(tela):
    c, _, _ = tela
    h = _quadro(c)
    card = h.split('href="/os/chamados/acompanhamento/9101"', 1)[1].split(">", 1)[0]
    assert 'data-geral="Cliente"' in card and 'data-ufv="Supervisor Um"' in card
    linha = h.split('data-href="/os/chamados/acompanhamento/9103"', 1)[1].split(">", 1)[0]
    assert 'data-geral="Standby"' in linha and 'data-ufv=""' in linha


def test_os_filtros_status_geral_e_responsavel_ufv_no_lugar_do_responsavel(tela):
    c, _, _ = tela
    h = _quadro(c)
    filtros = h.split('<div class="filtros">', 1)[1].split('<div class="board">', 1)[0]
    assert 'data-f="geral"' in filtros and 'data-f="ufv"' in filtros and 'data-f="equipe"' not in filtros
    geral = filtros.split('data-f="geral"', 1)[1].split("</select>", 1)[0]
    assert [o for o in ("Grid", "Cliente", "Fabricante", "Standby") if '<option value="%s">' % o in geral] == \
        ["Grid", "Cliente", "Fabricante", "Standby"]
    assert '<option value="_sem">Sem status</option>' in geral
    ufv = filtros.split('data-f="ufv"', 1)[1].split("</select>", 1)[0]
    assert '<option value="Supervisor Dois">' in ufv and '<option value="Supervisor Um">' in ufv
    assert "Limpar filtros" not in h and "data-limpar" not in h                                  # o item 7


def test_sem_o_nexus_dizer_o_responsavel_ufv_o_filtro_nao_aparece(tela, monkeypatch):
    c, _, ra = tela
    monkeypatch.setattr(ra, "RESPONSAVEL_UFV", None)
    assert 'data-f="ufv"' not in _quadro(c)


def test_a_tabela_tem_o_geral(tela):
    c, _, _ = tela
    tab = _quadro(c).split('<table class="tq">', 1)[1].split("</thead>", 1)[0]
    assert "<th>Geral</th>" in tab


def test_a_tela_do_chamado_mostra_o_geral_e_o_standby_continua_no_seletor(tela, monkeypatch):
    import api
    c, _, _ = tela
    monkeypatch.setattr(api, "get_os_detalhes", lambda wid: {"folio": 9100 + wid, "notas": "", "code": "TST-1",
                                                              "subtarefas": [], "status_id": 1})
    h = c.get("/os/chamados/acompanhamento/9103").get_data(as_text=True)
    caixa = h.split("data-status-caixa", 1)[1].split("</section>", 1)[0]
    assert '<option value="Standby" selected>Standby</option>' in caixa
    assert "Status geral: <b>Standby</b>" in caixa


# ── 2: os avisos na exclamação ──────────────────────────────────────────────────────────────────────────────────
def test_os_avisos_vao_para_a_exclamacao_ao_lado_do_titulo(tela, monkeypatch):
    import chamados_obs_store as obs_store
    c, _, _ = tela
    monkeypatch.setenv("OS_WEB_CHAMADO_RESP_NOME", "Equipe Teste")              # todas no nome de outra pessoa
    h = _quadro(c)
    tit = h.split('<div class="ac-tit">', 1)[1].split("</div>", 1)[0]
    assert '<details class="ac-aviso">' in tit and "no nome de outra pessoa" in tit
    assert '<p class="aviso"><b>' not in h                                       # nada solto embaixo da faixa
    def falha(forcar=False):
        raise RuntimeError("Read timed out")
    monkeypatch.setattr(obs_store, "listar", falha)
    h = c.get("/os/chamados/acompanhamento?atualizar=1").get_data(as_text=True)
    tit = h.split('<div class="ac-tit">', 1)[1].split("</div>", 1)[0]
    assert "Não consegui ler as observações" in tit


def test_sem_aviso_sem_exclamacao(tela, monkeypatch):
    c, _, _ = tela
    monkeypatch.delenv("OS_WEB_CHAMADO_RESP_NOME", raising=False)
    monkeypatch.delenv("OS_WEB_CHAMADO_RESP_ID", raising=False)
    assert "ac-aviso" not in _quadro(c)


# ── 3: sem o voltar para Chamados, no Nexus ─────────────────────────────────────────────────────────────────────
def test_no_nexus_o_quadro_nao_tem_o_voltar_para_chamados(logado, monkeypatch):
    pytest.importorskip("PyQt6")
    corpo = (b'<div class="ac-tit"><a class="ac-volta" href="/os/chamados">&larr; Chamados</a><h1>Acompanhamento</h1>'
             b'</div><a class="ac-volta" href="/os/chamados/acompanhamento">&larr; Acompanhamento</a>')
    limpo = ponte._SEM_VOLTAR_CHAMADOS.sub(b"", corpo)
    assert b'href="/os/chamados">' not in limpo and b'href="/os/chamados/acompanhamento">' in limpo


# ── 5: o relatório por cliente ──────────────────────────────────────────────────────────────────────────────────
def test_o_relatorio_de_um_cliente_e_um_pdf(tela):
    c, _, _ = tela
    r = c.post("/os/chamados/acompanhamento/relatorio", json={"clientes": ["Cliente A"]})
    assert r.status_code == 200 and r.mimetype == "application/pdf" and r.data[:4] == b"%PDF"
    assert "Cliente A" in r.headers["Content-Disposition"]


def test_varios_clientes_vao_num_zip_com_um_pdf_cada(tela):
    c, _, _ = tela
    r = c.post("/os/chamados/acompanhamento/relatorio", json={"clientes": ["Cliente A", "Cliente B"]})
    assert r.status_code == 200 and r.mimetype == "application/zip"
    nomes = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert len(nomes) == 2 and all(n.endswith(".pdf") for n in nomes)


def test_sem_cliente_nao_exporta(tela):
    c, _, _ = tela
    assert c.post("/os/chamados/acompanhamento/relatorio", json={"clientes": []}).status_code == 400


def test_o_modelo_do_relatorio_fala_com_o_cliente():
    from os_web import relatorio_chamados as rc
    hoje = aw.agora_brt().replace(year=2026, month=10, day=28, hour=12)
    cards = [aw.cartao(_linha(1, 9101), "SR 1", "", hoje, status="Aguardando retorno do Cliente"),
             aw.cartao(_linha(2, 9102, usina="Usina A2", marca="STI"), "SR 2", "", hoje,
                       status="Aguardando retorno do Fabricante"),
             aw.cartao(_linha(5, 9105), "", "", hoje, status=None),
             aw.cartao(_linha(4, 9104, status=3), "SR 4", "", hoje, status="Standby")]
    for x, ufv in zip(cards, ("Supervisor Um", "Supervisor Dois", "Supervisor Um", "")):
        x["resp_ufv"] = ufv
    m = rc.modelo(cards, "Cliente A", hoje, {"9101": ("2026-10-28T09:00:00-03:00", "Pedimos a declaração.")})
    assert m["titulo"] == "Atualização dos chamados de garantia · Cliente A"
    assert m["subtitulo"] == "Gestão de Chamados · O&M Grid Co. · Posição de 28/10/2026"
    # só os em andamento (a 9104 está concluída), do cliente pedido
    assert m["numeros"] == {"andamento": 3, "Cliente": 1, "Fabricante": 1, "Grid": 0, "Standby": 0, "sem": 1}
    assert [u["usina"] for u in m["usinas"]] == ["Usina A1", "Usina A2"]
    assert m["usinas"][0]["responsavel"] == "Supervisor Um" and m["usinas"][0]["total"] == 2
    secoes = {s["titulo"]: s["itens"] for s in m["secoes"]}
    assert list(secoes) == ["Aguardando o cliente (1)", "Aguardando o fabricante (1)", "Em acompanhamento (1)"]
    item = secoes["Aguardando o cliente (1)"][0]
    assert item["ref"] == "Ticket SR 1" and item["onde"] == "Usina A1 · Inversor 1.1"
    assert item["atualizacao"] == "28/10: Pedimos a declaração." and item["dias"] == 88
    assert "Huawei" in m["intro"] and "STI" in m["intro"] and "Usina A1" in m["intro"]
    assert any("cliente" in p for p in m["proximos"]) and any("fabricante" in p for p in m["proximos"])


# ── 6: exportar a tabela ────────────────────────────────────────────────────────────────────────────────────────
def test_exportar_a_tabela_leva_so_as_linhas_filtradas_na_ordem(tela):
    import openpyxl
    c, _, _ = tela
    r = c.post("/os/chamados/acompanhamento/tabela", json={"folios": ["9103", "9101"]})
    assert r.status_code == 200 and r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    cab = [x.value for x in ws[1]]
    assert cab[:8] == ["OS", "Ativo", "Cliente", "Usina", "Responsável UFV", "Fornecedor", "Geral", "Status"]
    assert [ws.cell(row=i, column=1).value for i in (2, 3)] == [9103, 9101] and ws.max_row == 3


# ── 4: o Responsável UFV pelo cadastro do Nexus ─────────────────────────────────────────────────────────────────
def test_o_nexus_diz_o_responsavel_ufv_pela_usina_do_cadastro(monkeypatch):
    from nexus.torres.chamados import ufv

    class Base:
        class lig:
            @staticmethod
            def usina(texto, codigo):
                return ({"Usina A1": 7}.get(texto) or {"TST-9": 8}.get(codigo), "x")

        @staticmethod
        def responsavel_da_usina(uid):
            return {7: "Supervisor Um", 8: "Supervisor Dois"}[uid]
    monkeypatch.setattr(ufv, "_base", lambda: Base)
    ufv._MEMO.clear()
    cards = [{"folio": "1", "usina": "Usina A1", "code": ""}, {"folio": "2", "usina": "Outra", "code": "TST-9"},
             {"folio": "3", "usina": "Nenhuma", "code": ""}]
    assert ufv.responsavel_ufv(cards) == {"1": "Supervisor Um", "2": "Supervisor Dois", "3": ""}


def test_sem_o_cadastro_o_responsavel_ufv_fica_vazio(monkeypatch):
    from nexus.torres.chamados import ufv

    def falha():
        raise RuntimeError("sem banco")
    monkeypatch.setattr(ufv, "_base", falha)
    ufv._MEMO.clear()
    assert ufv.responsavel_ufv([{"folio": "1", "usina": "Usina A1", "code": ""}]) == {}


def test_a_ponte_liga_o_responsavel_ufv_no_clone_rodando_no_app_do_nexus(app, monkeypatch):
    """O nome vem do cofre do cadastro, com a chave do Nexus: a função roda no app do Nexus, mesmo chamada pelo clone."""
    pytest.importorskip("PyQt6")
    import flask

    from nexus.torres.chamados import ufv
    ponte.clone(app)
    from os_web import rotas_acomp
    visto = {}

    class Base:
        class lig:
            @staticmethod
            def usina(texto, codigo):
                visto["app"] = flask.current_app._get_current_object()
                return 7, "x"

        @staticmethod
        def responsavel_da_usina(uid):
            return "Supervisor Um"
    monkeypatch.setattr(ufv, "_base", lambda: Base)
    ufv._MEMO.clear()
    assert rotas_acomp.RESPONSAVEL_UFV.nexus is ufv.responsavel_ufv
    assert rotas_acomp.RESPONSAVEL_UFV([{"folio": "1", "usina": "Usina A1", "code": ""}]) == {"1": "Supervisor Um"}
    assert visto["app"] is app


# ── o banco lento das observações ───────────────────────────────────────────────────────────────────────────────
def test_observacoes_vencidas_saem_na_hora_e_releem_por_tras(monkeypatch):
    pytest.importorskip("PyQt6")
    import chamados_obs_store as obs_store
    lidas = {"n": 0}

    def ler():
        lidas["n"] += 1
        return []
    fila = []
    monkeypatch.setattr(obs_store, "_linhas", ler)
    monkeypatch.setattr(obs_store, "_em_segundo_plano", fila.append)
    obs_store._estado.update(lista=None, lido=0.0, renovando=False)
    obs_store.listar()                                                  # sem nada guardado: lê na frente
    assert lidas["n"] == 1
    obs_store._estado["lido"] -= obs_store.TTL + 1                     # venceu
    obs_store.listar()                                                  # sai na hora, com a de antes
    assert lidas["n"] == 1 and len(fila) == 1
    fila.pop()()
    assert lidas["n"] == 2


def test_releitura_que_falha_fica_com_a_de_antes(monkeypatch):
    pytest.importorskip("PyQt6")
    import chamados_obs_store as obs_store
    obs_store._estado.update(lista=[{"id": "a", "quando": "2026-10-28T09:00:00-03:00", "os": "1", "tipo": "obs",
                                     "texto": "x", "quem": "", "email": "", "ativo": "", "extra": {}}],
                             lido=0.0, renovando=False)

    def falha():
        raise RuntimeError("Read timed out")
    monkeypatch.setattr(obs_store, "_linhas", falha)
    fila = []
    monkeypatch.setattr(obs_store, "_em_segundo_plano", fila.append)
    assert [x["id"] for x in obs_store.listar()] == ["a"]
    fila.pop()()                                                        # a releitura falha: fica a de antes
    assert [x["id"] for x in obs_store.listar()] == ["a"]
