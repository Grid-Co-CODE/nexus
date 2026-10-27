"""A Mesa da Segurança (Levi, 24/10/2026: "Criação de mesa da segurança, como PTs atrasadas, extintores vencidos, usinas
sem luvas"), por grupos (2ª rodada: "quero duas visões: Por responsável e cliente (...) não queria em formato de lista
mas em agrupamentos para o TST saber onde atacar!"): a primeira tela da torre, do setor Segurança, que só lê.

O banco falso do Campo (test_campo_visao.py) com o `com_cadastro` (Altair e Brodowski 1 da Supervisora Campo, Coração 1
do Coordenador Campo; Coração 1 é do Outro Cliente), na hora fixa dos extintores (09/10/2026 10:00): a PT-1 (Altair)
espera há 26 h, a PT-2 há 30 min; o EXT01 da Altair e o do Coração vencidos, o EXT02 da Altair não reconhecido."""
from test_campo_visao import _aba, _pt, banco  # noqa: F401  (banco é fixture)
from test_hseq_extintores import LIVRO, _texto, com_cadastro, com_livro  # noqa: F401  (fixtures)

from nexus.campo import visao
from nexus.hseq import extintores as EXT
from nexus.torres.hseq import TORRE


def test_a_mesa_e_a_primeira_tela_e_do_setor_seguranca():
    assert TORRE.telas[0].id == "mesa" and TORRE.telas[0].nome == "Mesa da Segurança"
    assert "Segurança" in (TORRE.telas[0].do_setor if isinstance(TORRE.telas[0].do_setor, (tuple, list))
                           else (TORRE.telas[0].do_setor,))


def _com_casos(banco_):
    _aba(banco_, "pt_app_campo", "PT", [dict(_pt("PT-1", 0), **{"Criada em": "2026-10-08T11:00:00.000Z"}),
                                       dict(_pt("PT-2", 0), **{"Criada em": "2026-10-09T12:30:00.000Z"})])
    # a coluna em todas as linhas (o banco falso tira o cabeçalho da primeira)
    _aba(banco_, EXT.LIVRO, EXT.ABA, [dict(x, **{"Não reconhecido em": "2026-10-05" if x["Código"] == "ALT100-INFC1-PPCI-EXT02"
                                                 else None}) for x in LIVRO])
    visao.limpar()


def _cartao(html, nome):
    """O texto do cartão de um grupo (até o próximo cartão ou o fim dos grupos)."""
    i = html.index(f"<b>{nome}</b>")
    fim = html.find('<div class="cn-equipe ', i)
    return _texto(html[i:fim if fim > 0 else html.index('class="hs-legenda"', i)])


def test_a_mesa_por_responsavel_mostra_onde_atacar(com_cadastro, logado):
    _com_casos(com_cadastro)
    html = logado.get("/t/hseq/mesa").get_data(as_text=True)
    faixa = _texto(html[html.index('class="kb-kpis'):html.index("</section>")])
    assert "PTs paradas 1" in faixa and "Extintores vencidos 2" in faixa and "Não reconhecidos 1" in faixa
    assert 'aria-current="page">Por responsável</a>' in html
    assert "PT-1" not in html and "EXT02" not in html                    # grupos, não a lista dos casos
    # um cartão por responsável, o que tem mais primeiro: a Supervisora Campo tem a PT, um vencido e o não reconhecido
    assert html.index("<b>Supervisora Campo</b>") < html.index("<b>Coordenador Campo</b>")
    assert _cartao(html, "Supervisora Campo").startswith(
        "Supervisora Campo 3 Responsável UFV 1 PTs paradas 1 extintor vencido em 1 usina 1 não reconhecido 0 usinas sem luvas")
    assert _cartao(html, "Coordenador Campo").startswith("Coordenador Campo 1 Responsável UFV 0 PTs paradas 1 extintor vencido")
    # cada número leva à tela de origem, filtrada pelo responsável
    assert 'href="/t/hseq/extintores?f=atrasado&amp;responsavel=Supervisora+Campo#tabela"' in html
    assert 'href="/t/hseq/extintores?f=nao_reconhecido&amp;responsavel=Supervisora+Campo#tabela"' in html
    assert 'href="/t/hseq/epi?modo=tabela&amp;f=sem_luvas&amp;responsavel=Coordenador+Campo"' in html
    # as PTs paradas, à tabela de APR e PT filtrada (Levi, 26/10/2026: "PTs parados vai levar para APR e PT > Esperando
    # o de acordo > Tabela > Filtra o responsável")
    assert 'href="/t/hseq/apr-pt?modo=tabela&amp;f=parada&amp;responsavel=Supervisora+Campo"' in html
    tabela = logado.get("/t/hseq/apr-pt?modo=tabela&f=parada&responsavel=Supervisora+Campo").get_data(as_text=True)
    assert tabela.count('class="cn-link cn-os"') == 1                          # a PT-1, parada (a PT-2 tem 30 min)
    assert 'class="hs-legenda"' in html and "Como ler a mesa" in _texto(html)
    # só lê: nenhum formulário que grava no conteúdo (o do topo é a troca de tema da casca)
    assert 'method="post"' not in html[html.index('class="campo-nativa"'):].lower()


def test_a_mesa_por_cliente(com_cadastro, logado):
    _com_casos(com_cadastro)
    html = logado.get("/t/hseq/mesa?ver=cliente").get_data(as_text=True)
    assert 'aria-current="page">Por cliente</a>' in html
    assert html.index("<b>Thopen</b>") < html.index("<b>Outro Cliente</b>")
    assert _cartao(html, "Thopen").startswith("Thopen 3 Cliente")
    assert 'href="/t/hseq/extintores?f=atrasado&amp;cliente=Outro+Cliente#tabela"' in html


def test_sem_os_extintores_a_mesa_diz_e_os_outros_seguem(banco, logado):
    html = logado.get("/t/hseq/mesa").get_data(as_text=True)
    assert "Não consegui ler os extintores agora" in html and "PTs paradas" in html
    assert "Extintores vencidos —" in _texto(html[html.index('class="kb-kpis'):html.index("</section>")])


# ── a visão geral (a TST, 24/10/2026: "Adicionar uma opção de visão geral (um resumo com todas as pendências, sendo
#    possível filtrar por cliente, por UFV e por supervisor) e ter a opção de baixar"; Levi: "Os quatro + os de
#    atenção") ──────────────────────────────────────────────────────────────────────────────────────────────────────
def _form(html):
    i = html.index('class="cn-filtros hs-mesa-filtros"')
    return html[i:html.index("</form>", i)]


def test_a_visao_geral_filtra_por_cliente_ufv_e_responsavel_em_cascata(com_cadastro, logado):
    import re
    _com_casos(com_cadastro)
    html = logado.get("/t/hseq/mesa").get_data(as_text=True)
    form = _form(html)
    assert re.findall(r'<select class="gc-campo" id="mesa-(\w+)" name="(\w+)"', form) == [
        ("cliente", "cliente"), ("usina", "usina"), ("responsavel", "responsavel")]
    assert "Responsável UFV" in _texto(form) and "UFV" in _texto(form)
    assert 'href="/t/hseq/mesa/pendencias.xlsx"' in form and ">Baixar planilha</a>" in form
    # o Outro Cliente: só o Coração 1 e o Coordenador Campo nas listas, e os números só dele
    out = logado.get("/t/hseq/mesa?cliente=Outro+Cliente").get_data(as_text=True)
    faixa = _texto(out[out.index('class="kb-kpis'):out.index("</section>")])
    assert "PTs paradas 0" in faixa and "Extintores vencidos 1" in faixa and "Não reconhecidos 0" in faixa
    f2 = _form(out)
    usinas = re.findall(r'<option value="([^"]+)"', f2[f2.index('id="mesa-usina"'):f2.index('id="mesa-responsavel"')])
    resp = re.findall(r'<option value="([^"]+)"', f2[f2.index('id="mesa-responsavel"'):])
    assert usinas == ["Coração 1"] and resp == ["Coordenador Campo"]
    assert "<b>Supervisora Campo</b>" not in out
    # o link do cartão leva também o filtro do alto; a planilha, o mesmo recorte
    assert 'href="/t/hseq/extintores?f=atrasado&amp;cliente=Outro+Cliente&amp;responsavel=Coordenador+Campo#tabela"' in out
    assert 'href="/t/hseq/mesa/pendencias.xlsx?cliente=Outro+Cliente"' in out


def test_os_de_atencao_na_faixa_e_nos_cartoes(com_cadastro, logado):
    _com_casos(com_cadastro)
    html = logado.get("/t/hseq/mesa").get_data(as_text=True)
    i = html.index('<p class="hs-atencao">')
    linha = _texto(html[i:html.index("</p>", i)])
    # o Brodowski nunca foi conferido (sem validade): sem atualização; o Coração também, mas está vencido e não conta
    assert linha.startswith("Atenção 0 perto de vencer 1 sem atualização")
    assert 'href="/t/hseq/extintores?f=sem_atualizacao#tabela"' in html
    sup = _cartao(html, "Supervisora Campo")
    assert "Atenção 1 sem atualização" in sup and "perto de vencer" not in sup
    assert 'href="/t/hseq/extintores?f=sem_atualizacao&amp;responsavel=Supervisora+Campo#tabela"' in html


def test_a_planilha_leva_caso_a_caso_com_os_filtros(com_cadastro, logado):
    from io import BytesIO

    from openpyxl import load_workbook

    from nexus.hseq import planilha as PLA
    _com_casos(com_cadastro)
    r = logado.get("/t/hseq/mesa/pendencias.xlsx?responsavel=Supervisora+Campo")
    assert r.status_code == 200 and r.mimetype == PLA.MIME
    assert r.headers["Content-Disposition"] == 'attachment; filename="pendencias-seguranca-supervisora-campo-2026-10-09.xlsx"'
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Resumo", "PTs paradas", "Extintores vencidos", "Não reconhecidos", "Usinas sem luvas",
                             "Perto de vencer", "Sem atualização", "Luvas sem verificação", "Nova ronda pedida"]
    res = [[c.value for c in linha] for linha in wb["Resumo"].iter_rows()]
    assert "Filtros: Responsável UFV: Supervisora Campo" in [x[0] for x in res]
    cab = next(i for i, x in enumerate(res) if x[0] == "Responsável UFV")
    assert res[cab][:5] == ["Responsável UFV", "PTs paradas", "Extintores vencidos", "Não reconhecidos", "Usinas sem luvas"]
    assert res[cab + 1][:7] == ["Supervisora Campo", 1, 1, 1, 0, 0, 1] and len(res) == cab + 2
    pts = [[c.value for c in linha] for linha in wb["PTs paradas"].iter_rows()]
    cab = next(i for i, x in enumerate(pts) if x[0] == "Cliente")
    assert pts[cab + 1][:4] == ["Thopen", "Altair", "Supervisora Campo", "PT-1"] and pts[cab + 1][-1] == 26.0
    venc = [[c.value for c in linha] for linha in wb["Extintores vencidos"].iter_rows()]
    assert [x[4] for x in venc if x[0] == "Thopen"] == ["ALT100-INFC1-PPCI-EXT01"]
