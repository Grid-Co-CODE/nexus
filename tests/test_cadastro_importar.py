"""Importação do BD_Operações para o ensaio do Nexus.

A planilha de teste é montada aqui com os cabeçalhos REAIS do BD (29/09/2026) e os casos que ele tem de fato:
fórmula e valor digitado na mesma coluna, coordenada em texto, "Pendente" em data, supervisor que só existe nas
listas do Auxiliar. Nenhum dado real de pessoa.
"""
import io

import openpyxl
import pytest

from nexus.cadastro.armazem import ArmazemLocal
from nexus.cadastro.calculos import ERRO
from nexus.cadastro.cifra import Cofre, gerar_chave
from nexus.cadastro.importar import Aba, Celula, comparar, ler_xlsx, montar
from nexus.cadastro.servico import Servico
from nexus.cadastro.tipos import Legado

OPS = ["OPERAÇÃO", "CÓDIGO", "STATUS", "CLIENTE", "POTÊNCIA CONTRATUAL (MWp)", "LATITUDE", "Prazo Contratual (Meses)",
       "Receita Mensal", "Receita Contratual", "CIDADE", "LOCALIZAÇÃO", "PAÍS", "UF", "CLUSTER", "Equipe Cluster",
       "REGIÃO", "Gestor de Contrato", "RESPONSÁVEL O&M", "Técnico O&M", "Data Mobilização", "IDUsina"]
PES = ["Nome", "Nome Padrão", "Cluster", "Cargo", "cod", "CPF", "Telefone Pessoal", "Telefone Corporativo ",
       "endereco", "email", "Data de Admissão", "Supervisor ", "Cliente ", "Cidade", "Status de Contratação"]
F_TECNICO = ('=IFERROR(_xlfn.XLOOKUP(Operacoes[[#This Row],[Equipe Cluster]]&Operacoes[[#Headers],[Técnico O&M]],'
             "'Relação Geral Colaboradores'!E:E,'Relação Geral Colaboradores'!B:B),\"\")")
F_RECEITA = "=3915.51*Operacoes[[#This Row],[POTÊNCIA CONTRATUAL (MWp)]]"


# Mesmo nome de outra usina, outro cliente: é o caso real (2 pares no BD de 29/09, cada par na mesma UF).
OUTRA_BRODOWSKI = ["Brodowski", "ATH100", "OPERAÇÃO", "Athon", 1, None, None, None, None, "Brodowski", None,
                   "Brasil", "SP", None, None, None, None, None, None, None, "UFV-12"]


def _xlsx(inverter=False, extra=()) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    op = wb.create_sheet("Operações")
    op.append(OPS)
    linhas = [
        ["Brodowski", "BRD100", "OPERAÇÃO", "Thopen", 3, " -9,87°", 60, F_RECEITA, "=N2*G2", "Brodowski",
         "=CONCAT(J2)", "Brasil", "SP", "SP Leste", "SP Leste 01", " SUDESTE ", "Davi Rocha", "Carla Souza",
         F_TECNICO, "Pendente", "UFV-10"],
        ["Batatais", None, "A MOBILIZAR", "Thopen", "=9.633/3", None, None, 1000, "=N3*G3", "Batatais",
         "=CONCAT(J3)", "Brasil", "SP", None, "SP Leste 01", "Sudeste", None, "carla souza",
         "Bruno Costa", None, "UFV-11"],
        *extra,
    ]
    for ln in (reversed(linhas) if inverter else linhas):
        op.append(list(ln))
    pe = wb.create_sheet("Relação Geral Colaboradores")
    pe.append(PES)
    pe.append(["Ana Maria Lima", '=LEFT(A2,FIND(" ",A2)-1)', "SP Leste 01", "Técnico O&M", "=C2&D2", 1234567890,
               None, 11912345678, "Rua A, 1, Centro, Brodowski, SP", "Ana.Lima@gridco.com.br", None, "=XLOOKUP(1)",
               "Thopen", "Brodowski", "Ativo"])
    pe.append(["Bruno Costa", "Bruno Costa", "sp leste 01", "Eletricista O&M", "=C3&D3", None, None, None, None,
               None, None, None, "Thopen", None, "Ativo"])
    ax = wb.create_sheet("Auxiliar")
    ax.append(["Faixa Potência", "Gestor de Contrato", "Supervisor"])
    ax.append(["1MWp | 3MWp", "Davi Rocha", "Carla Souza"])
    ax.append(["3MWp | 5MWp", None, None])
    pm = wb.create_sheet("Parametros")
    pm.append(["Status Operacional", "Modalidade", "Status Contrato"])
    pm.append(["Ativo", "Full", "Contrato Assinado"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def srv(tmp_path):
    return Servico(ArmazemLocal(tmp_path / "c.json"), Cofre(gerar_chave()))


@pytest.fixture
def proposta(srv):
    return montar(ler_xlsx(_xlsx()), srv, arquivo="BD_Operacoes.xlsx")


def _usina(carga, id_bd):
    """A usina pelo IDUsina do Excel: o id do Nexus agora é número simples (1, 2, 3)."""
    return next(r for r in carga.entidades["usinas"] if r["valores"]["id_bd"] == id_bd)


def _por_nome(carga, entidade, nome):
    return next(r for r in carga.entidades[entidade] if r["valores"].get("nome") == nome)


def test_ler_marca_o_que_era_formula(proposta):
    assert proposta.origem["abas"]["Operações"] == 2


def test_usina_formula_vira_automatico_e_digitado_vira_sobreposicao(proposta):
    c = proposta.carga
    u10 = _usina(c, "UFV-10")
    u11 = _usina(c, "UFV-11")
    assert u10["valores"]["tecnico_om"] is None                     # era fórmula: fica automático
    bruno = _por_nome(c, "pessoas", "Bruno Costa")
    assert u11["valores"]["tecnico_om"] == bruno["id"]              # digitado: vira a pessoa


def test_receita_por_mwp_vira_preco_e_a_receita_fica_automatica(proposta):
    u10 = _usina(proposta.carga, "UFV-10")
    assert u10["valores"]["preco_mwp"] == 3915.51
    assert u10["valores"]["receita_mensal"] is None


def test_formula_de_conta_fixa_entra_como_o_valor_mostrado():
    # "=9.633/3" na potência: é digitação, não regra. Sem valor em cache (arquivo do openpyxl), fica vazio,
    # e o aviso diz quantas células eram fórmula.
    p = montar(ler_xlsx(_xlsx()), None, arquivo="x.xlsx")
    assert any("POTÊNCIA CONTRATUAL" in a for a in p.avisos)


def test_tipos_do_excel_restaurados(proposta):
    u10 = _usina(proposta.carga, "UFV-10")
    assert u10["valores"]["latitude"] == -9.87
    assert u10["valores"]["data_mobilizacao"] == Legado("Pendente")
    assert u10["valores"]["regiao"] == "SUDESTE"                    # espaço das pontas sai; a maiúscula fica
    ana = _por_nome(proposta.carga, "pessoas", "Ana Maria Lima")
    assert ana["valores"]["cpf"] == "01234567890"
    assert ana["valores"]["email"] == "ana.lima@gridco.com.br"
    assert ana["valores"]["telefone_corporativo"] == "11912345678"


def test_equipe_casa_sem_maiuscula_como_o_procv(proposta):
    # "sp leste 01" (pessoa) e "SP Leste 01" (usina): o PROCV do Excel não diferencia maiúscula; é uma equipe só.
    equipes = proposta.carga.entidades["equipes"]
    assert [e["valores"]["nome"] for e in equipes] == ["SP Leste 01"]
    bruno = _por_nome(proposta.carga, "pessoas", "Bruno Costa")
    assert bruno["valores"]["equipe"] == equipes[0]["id"]


def test_supervisor_e_gestor_do_auxiliar_viram_pessoas_e_as_usinas_apontam_para_elas(proposta):
    c = proposta.carga
    carla = _por_nome(c, "pessoas", "Carla Souza")
    davi = _por_nome(c, "pessoas", "Davi Rocha")
    assert carla["valores"]["vinculo"] == "Supervisor" and davi["valores"]["vinculo"] == "Gestor de contrato"
    u10 = _usina(c, "UFV-10")
    u11 = _usina(c, "UFV-11")
    assert u10["valores"]["responsavel_om"] == carla["id"]
    assert u11["valores"]["responsavel_om"] == carla["id"]          # "carla souza" casa sem maiúscula
    assert u10["valores"]["gestor_contrato"] == davi["id"]


def test_listas_juntam_auxiliar_parametros_e_o_que_esta_nos_dados(proposta):
    ls = proposta.carga.listas
    assert {"OPERAÇÃO", "A MOBILIZAR"} <= set(ls["status_usina"])
    assert ls["faixa_potencia"][:2] == ["1MWp | 3MWp", "3MWp | 5MWp"]
    assert "Full" in ls["modalidade"]
    assert {"SUDESTE", "Sudeste"} <= set(ls["regiao"])


def test_importar_de_novo_nao_muda_ids_nem_cria_versao(srv, proposta):
    srv.aplicar_carga(proposta.carga)
    ids = sorted(r["id"] for r in proposta.carga.entidades["pessoas"])
    versoes = {r.id: r.versao for r in srv.registros("usinas")}
    de_novo = montar(ler_xlsx(_xlsx()), srv, arquivo="BD_Operacoes.xlsx")
    assert sorted(r["id"] for r in de_novo.carga.entidades["pessoas"]) == ids
    assert de_novo.resumo["usinas"]["novos"] == 0 and de_novo.resumo["usinas"]["alterados"] == 0
    srv.aplicar_carga(de_novo.carga)
    assert {r.id: r.versao for r in srv.registros("usinas")} == versoes


def test_resumo_avisa_o_que_foi_editado_no_nexus(srv, proposta):
    srv.aplicar_carga(proposta.carga)
    r = next(x for x in srv.registros("usinas") if x.valores["id_bd"] == "UFV-11")
    f = srv.formulario(r, revelar=True)
    f["nome"] = "Batatais 1"
    assert srv.salvar("usinas", r.id, f, r.versao, quem="admin").ok
    de_novo = montar(ler_xlsx(_xlsx()), srv, arquivo="BD_Operacoes.xlsx")
    assert de_novo.resumo["usinas"]["editados_no_nexus"] == [r.id]


# ── ID numérico e cliente com ID (Levi, 29/09: "reduza as caracteres do ID [...] ID para cliente e ID para usinas") ──

def test_id_vira_numero_na_ordem_da_planilha_e_o_idusina_fica_guardado(proposta):
    us = proposta.carga.entidades["usinas"]
    assert [(u["id"], u["valores"]["id_bd"]) for u in us] == [("1", "UFV-10"), ("2", "UFV-11")]
    # Ana e Bruno (Relação Geral), depois Carla (Supervisor) e Davi (Gestor), que só existiam no Auxiliar.
    assert [p["id"] for p in proposta.carga.entidades["pessoas"]] == ["1", "2", "3", "4"]
    assert [e["id"] for e in proposta.carga.entidades["equipes"]] == ["1"]


def test_cliente_vira_cadastro_com_id_e_separa_usinas_de_mesmo_nome():
    p = montar(ler_xlsx(_xlsx(extra=[OUTRA_BRODOWSKI])), None, arquivo="x.xlsx")
    clientes = {c["valores"]["nome"]: c["id"] for c in p.carga.entidades["clientes"]}
    assert clientes == {"Thopen": "1", "Athon": "2"}
    brod = sorted((u["id"], u["valores"]["cliente"]) for u in p.carga.entidades["usinas"]
                  if u["valores"]["nome"] == "Brodowski")
    assert brod == [("1", "1"), ("3", "2")]          # mesmo nome, ids e clientes diferentes
    assert "clientes" not in p.carga.listas            # deixou de ser lista: é cadastro


def test_reimportar_com_as_linhas_em_outra_ordem_casa_pelo_idusina(srv, proposta):
    srv.aplicar_carga(proposta.carga)
    antes = {u.valores["id_bd"]: u.id for u in srv.registros("usinas")}
    de_novo = montar(ler_xlsx(_xlsx(inverter=True)), srv, arquivo="x.xlsx")
    assert {u["valores"]["id_bd"]: u["id"] for u in de_novo.carga.entidades["usinas"]} == antes


# ── comparação Excel × Nexus (o valor que a planilha mostrava contra o que o Nexus calcula) ──

def _cel(v, f=None):
    return Celula(valor=v, formula=f)


def test_comparacao_explica_quando_o_excel_pegava_desligado():
    ops = Aba(["IDUsina", "OPERAÇÃO", "STATUS", "Equipe Cluster", "Técnico O&M"], [
        [_cel("UFV-1"), _cel("A"), _cel("OPERAÇÃO"), _cel("SP Leste 01"), _cel("Ana Lima", "=XLOOKUP()")],
        [_cel("UFV-2"), _cel("B"), _cel("OPERAÇÃO"), _cel("CE Sul 01"), _cel("Rui Dias", "=XLOOKUP()")],
    ])
    pes = Aba(["Nome", "Nome Padrão", "Cluster", "Cargo", "Status de Contratação"], [
        [_cel("Ana Maria Lima"), _cel("Ana Lima", "=LEFT()"), _cel("SP Leste 01"), _cel("Técnico O&M"), _cel("Ativo")],
        [_cel("Rui Dias"), _cel("Rui Dias", "=LEFT()"), _cel("CE Sul 01"), _cel("Técnico O&M"), _cel("Desligado")],
        [_cel("Eva Luz"), _cel("Eva Luz", "=LEFT()"), _cel("CE Sul 01"), _cel("Técnico O&M"), _cel("Ativo")],
    ])
    p = montar({"Operações": ops, "Relação Geral Colaboradores": pes}, None, arquivo="x.xlsx")
    tec = next(c for c in p.comparacao if c["campo"] == "tecnico_om")
    assert tec["total"] == 2 and tec["iguais"] == 1
    dif = tec["diferentes"][0]
    assert dif["id"] == "2" and dif["excel"] == "Rui Dias" and dif["nexus"] == "Eva Luz"
    assert "desligad" in dif["motivo"]


def test_comparacao_trata_zero_do_excel_como_vazio():
    # Pessoa sem cluster: o PROCX do Excel casa com uma linha vazia e mostra 0; o Nexus não acha ninguém.
    assert comparar("ref", 0, [])
    assert comparar("telefone", 0, None)
    # ...mas vazio de um lado e valor do outro é diferença (senão a comparação aprova tudo).
    assert not comparar("ref", 0, ["Ana Lima"])
    assert not comparar("telefone", None, "11912345678")
    # Espaço sobrando ou quebra de linha não é valor diferente (o Nexus apara o texto na importação).
    assert comparar("texto", "Brodowski ,SP,Brasil", "Brodowski,SP,Brasil")
    assert comparar("texto", "Ana Lima\n", "Ana Lima")
    assert comparar("texto", "#N/A", ERRO)
    assert comparar("moeda", 704700.0, 704700.004)
    assert not comparar("texto", "Brodowski/SP", "Batatais/SP")
