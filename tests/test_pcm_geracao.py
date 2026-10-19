"""Etapa 2: o Nexus gera a semana em SOMBRA, com o motor do PCM copiado sem mudança (nexus/pcm/motor), e compara
com a semana oficial do PCM. Nada é publicado.

Os testes usam um MOTOR FALSO (um script que escreve uma planilha no formato do gerador): nenhum teste vai ao Fracttal.
"""
import json
import re
import sys
import textwrap
from datetime import date

import openpyxl
import pytest

from nexus.pcm import auxiliar as A
from nexus.pcm import comparar as C
from nexus.pcm import geracao as G
from nexus.pcm import insumos as I
from test_pcm_auxiliar import cadastro_de_teste
from test_pcm_insumos import criar_pasta_pcm

CAB = ["Equipe", "Dia", "OSs ID", "Ativo (Usina)", "Código Equipamento", "Tipo", "Tarefa", "Reprogramada",
       "Nº vezes programada", "Hora Início", "Hora Fim", "Duração (h)", "RPN/Prioridade"]
LINHAS = [
    ["SP Leste 02", "Segunda-feira (05/10)", 15020, "Thopen - Altair 1 - SP", "ALT1-INV01", "Corretiva",
     "Fusível queimado", "Não", 1, "07:30", "09:00", 1.5, 12],
    ["SP Leste 02", "Terça-feira (06/10)", 15021, "Thopen - Altair 1 - SP", "ALT1-CAB01", "Preventiva",
     "MPM cabine", "Sim", 3, "07:30", "08:36", 1.1, 20],
]


def planilha(caminho, linhas=LINHAS, pendentes=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SP Leste 02"
    ws.append(CAB)
    for l in linhas:
        ws.append(l)
    wp = wb.create_sheet("_Pendentes")
    wp.append(["Equipe", "OSs ID", "Código Equipamento", "Tarefa", "Motivo"])
    for p in pendentes:
        wp.append(p)
    wb.save(caminho)


def historico(caminho, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["task_key", "first_week", "last_week", "count", "weeks"])
    for l in linhas:
        ws.append(l)
    wb.save(caminho)


@pytest.fixture
def origem(tmp_path):
    """A pasta do PCM de mentira, com tudo o que o motor lê, no formato dos arquivos de verdade."""
    o = criar_pasta_pcm(tmp_path / "pcm")
    return o


def motor_falso(tmp_path, codigo_saida=0, espera=0.0):
    """Faz o papel do programacao_v7.py: confere os insumos na PCM_PROG_DIR e o ambiente que o Nexus passa, e escreve
    em PCM_OUTPUT a mesma planilha das LINHAS (sem pendentes)."""
    m = tmp_path / f"motor_falso_{codigo_saida}_{int(espera * 10)}.py"
    m.write_text(textwrap.dedent(f"""
        import json, os, sys, time
        import openpyxl
        base = os.environ["PCM_PROG_DIR"]
        for nome in ("{A.NOME}", "Historico_Programacoes.xlsx", "Observacoes_Semana.txt"):
            assert os.path.exists(os.path.join(base, nome)), nome
        assert os.environ.get("FRACTTAL_CLIENT_ID") == "id-de-teste"
        assert os.environ.get("FRACTTAL_BASE_URL", "").endswith("/api/")
        assert not any(k.startswith("NEXUS_") for k in os.environ)
        print("semana pedida:", sys.argv[sys.argv.index("--semana") + 1])
        time.sleep({espera})
        if {codigo_saida}:
            print("Traceback: deu ruim no motor falso")
            sys.exit({codigo_saida})
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SP Leste 02"
        ws.append(json.loads({json.dumps(json.dumps(CAB))}))
        for l in json.loads({json.dumps(json.dumps(LINHAS))}):
            ws.append(l)
        wp = wb.create_sheet("_Pendentes")
        wp.append(["Equipe", "OSs ID", "Código Equipamento", "Tarefa", "Motivo"])
        wb.save(os.environ["PCM_OUTPUT"])
        print("OK")
    """), encoding="utf-8")
    return m


# o fato_programacao do banco de mentira: a 15021|ALT1-CAB01 programada na W39 e na W40, a 15020|ALT1-INV01 na W40.
# Gravado em 06/10: a W40 já tinha fechado (a W41 é a semana gerada nos testes)
FATO = {"linhas": [{"data_id_semana": "20260921", "os": "15021", "codigo_ativo": "ALT1-CAB01",
                    "lido_em": "2026-10-06T10:00:00-03:00"},
                   {"data_id_semana": "20260928", "os": "15021", "codigo_ativo": "ALT1-CAB01",
                    "lido_em": "2026-10-06T10:00:00-03:00"},
                   {"data_id_semana": "20260928", "os": "15020", "codigo_ativo": "ALT1-INV01",
                    "lido_em": "2026-10-06T10:00:00-03:00"}],
        "gerado_em": "2026-10-06T10:00:00-03:00"}


@pytest.fixture
def config(tmp_path, origem):
    # na sombra, o Nexus importa os arquivos da pasta do PCM antes de gerar (insumos.importar); o histórico vem do banco
    I.importar(tmp_path / "trabalho", origem, "2026-W41")
    return {"TESTING": True, "NEXUS_PCM_ORIGEM": str(origem), "NEXUS_PCM_TRABALHO": str(tmp_path / "trabalho"),
            "NEXUS_PCM_MOTOR": str(motor_falso(tmp_path)),
            "NEXUS_PCM_CADASTRO_TESTE": cadastro_de_teste(tmp_path / "cadastro.json"),
            "NEXUS_PCM_FRACTTAL_TESTE": {"FRACTTAL_CLIENT_ID": "id-de-teste", "FRACTTAL_CLIENT_SECRET": "segredo"},
            "NEXUS_PCM_HISTORICO_TESTE": FATO}


def guardar_reserva(config, linhas):
    """A reserva do histórico no Nexus (o que era importado do arquivo do PC até 08/10)."""
    trab = G.pasta_trabalho(config)
    d = I.carregar(trab)
    d["historico"] = {"aba": None, "preambulo": [], "colunas": ["task_key", "first_week", "last_week", "count", "weeks"],
                      "linhas": linhas, "meta": {"importado_em": "2026-10-05T09:49:19-03:00"}}
    I.salvar(trab, d)


def historico_da_rodada(config, rid):
    import pandas as pd
    p = G.pasta_trabalho(config) / "geracoes" / rid / "Historico_Programacoes.xlsx"
    return {str(r["task_key"]): (int(r["count"]), str(r["weeks"])) for _, r in pd.read_excel(p).iterrows()}


# ── regras ───────────────────────────────────────────────────────────────────────────────────────────────

def test_semana_padrao_e_a_proxima_segunda():
    assert G.semana_padrao(date(2026, 9, 30)) == "2026-W41"      # quarta -> semana de 05/10
    assert G.semana_padrao(date(2026, 10, 1)) == "2026-W41"      # quinta, dia de gerar
    assert G.semana_padrao(date(2026, 10, 5)) == "2026-W42"      # segunda -> a próxima


def test_historico_volta_para_antes_da_semana_gerada(tmp_path):
    """Gerar a semana de novo marcava tudo como reprogramado (medido na semana 40: 415 tarefas). O Nexus gera a partir do
    histórico de ANTES da semana, então tanto faz rodar antes ou depois do PCM."""
    src, dst = tmp_path / "h.xlsx", tmp_path / "h2.xlsx"
    historico(src, [["A|1", "2026-W39", "2026-W41", 3, "2026-W39,2026-W40,2026-W41"],
                    ["B|2", "2026-W41", "2026-W41", 1, "2026-W41"],
                    ["C|3", "2026-W38", "2026-W38", 1, "2026-W38"]])
    G.preparar_historico(src, dst, "2026-W41")
    ws = openpyxl.load_workbook(dst).active
    linhas = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(linhas) == {"A|1", "C|3"}                 # a B só existia por causa da W41
    assert linhas["A|1"][1:] == ("2026-W39", "2026-W40", 2, "2026-W39,2026-W40")
    assert linhas["C|3"][4] == "2026-W38"


def test_credencial_vem_do_ambiente_antes_dos_arquivos(config):
    c = G.credencial(config)
    assert c.ok and c.onde == "teste"
    assert "id-de-teste" not in repr(c) and "segredo" not in repr(c)   # o valor nunca aparece
    sem = dict(config, NEXUS_PCM_FRACTTAL_TESTE={})
    assert not G.credencial(sem).ok


def test_conferir_lista_o_que_falta(config, origem, tmp_path):
    ok = G.conferir(config)
    assert ok["pronto"], ok
    # o histórico vem do banco: o arquivo da pasta do PCM não faz falta à geração
    (origem / "Historico_Programacoes.xlsx").unlink()
    assert G.conferir(config)["pronto"]
    # a AUXILIAR sai do cadastro do Nexus: a planilha da pasta não faz falta, o cadastro faz
    sem_cadastro = {k: v for k, v in config.items() if k != "NEXUS_PCM_CADASTRO_TESTE"}
    falta = G.conferir(sem_cadastro)
    assert [i["nome"] for i in falta["itens"] if not i["ok"]] == ["Cadastro de usinas (AUXILIAR)"]
    # Nexus sem os insumos importados: cada um aparece como falta (o histórico não: é do banco, e a reserva é opcional)
    vazio = dict(config, NEXUS_PCM_TRABALHO=str(tmp_path / "outro"))
    nomes = [i["nome"] for i in G.conferir(vazio)["itens"] if not i["ok"]]
    assert nomes == [I.NOMES[k] for k in ("prioridades", "confiabilidade", "feriados")]
    hist = [i for i in G.conferir(vazio)["itens"] if i["nome"] == I.NOMES["historico"]][0]
    assert hist["ok"] and not hist["obrigatorio"] and "do banco" in hist["detalhe"]


def test_historico_da_geracao_vem_do_banco(config):
    """08/10/2026 ("PUXE O HISTÓRICO"): o motor recebe o histórico montado do fato_programacao, com as semanas de
    antes da gerada, e a rodada guarda de onde veio cada semana."""
    guardar_reserva(config, [["15021|ALT1-CAB01", "2026-W20", "2026-W41", 9, "2026-W20,2026-W41"]])
    st = G.aguardar(config, G.iniciar(config, "2026-W41")["id"], 60)
    assert st["estado"] == "ok", st
    assert historico_da_rodada(config, st["id"]) == {"15021|ALT1-CAB01": (2, "2026-W39,2026-W40"),
                                                     "15020|ALT1-INV01": (1, "2026-W40")}
    carimbo = [i for i in st["insumos"] if i["nome"] == I.NOMES["historico"]][0]
    assert "do banco" in carimbo["detalhe"] and not carimbo["aviso"]
    rel = json.loads((G.pasta_trabalho(config) / "geracoes" / st["id"] / "historico_banco.json").read_text("utf-8"))
    assert rel["semanas"]["2026-W40"]["origem"] == "banco"


def test_banco_fora_do_ar_usa_a_reserva_de_antes_da_semana(config):
    guardar_reserva(config, [["15021|ALT1-CAB01", "2026-W39", "2026-W41", 2, "2026-W39,2026-W41"]])
    cfg = dict(config, NEXUS_PCM_HISTORICO_TESTE=RuntimeError("503"))
    st = G.aguardar(cfg, G.iniciar(cfg, "2026-W41")["id"], 60)
    assert st["estado"] == "ok", st
    assert historico_da_rodada(cfg, st["id"]) == {"15021|ALT1-CAB01": (1, "2026-W39")}   # a W41 saiu
    carimbo = [i for i in st["insumos"] if i["nome"] == I.NOMES["historico"]][0]
    assert carimbo["aviso"] and "o banco não respondeu" in carimbo["detalhe"]


def test_sem_banco_e_sem_reserva_nao_gera(config):
    cfg = dict(config, NEXUS_PCM_HISTORICO_TESTE=RuntimeError("503"))
    with pytest.raises(G.NaoPronto) as erro:
        G.iniciar(cfg, "2026-W41")
    assert "histórico" in str(erro.value) and "reserva" in str(erro.value)
    assert not list((G.pasta_trabalho(cfg) / "geracoes").iterdir())     # a pasta da rodada não fica para trás


def test_servidor_sem_a_pasta_do_pcm_gera_mesmo_assim(config, tmp_path):
    """No servidor não existe o OneDrive do PCM. Desde 02/10 a AUXILIAR sai do cadastro e os insumos moram no Nexus:
    a pasta só serve para comparar com a oficial e trazer as durações (sombra). Ela era obrigatória e travava a
    geração no servidor."""
    cfg = dict(config, NEXUS_PCM_ORIGEM=str(tmp_path / "nao-existe"))
    conf = G.conferir(cfg)
    assert conf["pronto"], [i for i in conf["itens"] if not i["ok"]]
    st = G.aguardar(cfg, G.iniciar(cfg, "2026-W41")["id"], 60)
    assert st["estado"] == "ok", st
    assert "comparacao" not in st                     # sem a oficial, não há com o que comparar


def test_gerar_em_sombra_e_comparar_com_a_oficial(config, origem):
    planilha(origem / "Programação Semana 41.xlsx")               # a semana oficial (a do PCM)
    r = G.iniciar(config, "2026-W41")
    st = G.aguardar(config, r["id"], 60)
    assert st["estado"] == "ok", st
    assert st["resumo"]["linhas"] == 2
    assert st["comparacao"]["identicas"] is True
    assert "semana pedida: 2026-W41" in G.log(config, r["id"])
    # os insumos são carimbados: de onde veio cada um e de quando é o dado
    assert {i["nome"] for i in st["insumos"]} >= {"Cadastro de usinas (AUXILIAR)", I.NOMES["historico"], I.NOMES["feriados"]}


def test_diferenca_com_a_oficial_aparece_por_campo(config, origem):
    outra = [list(LINHAS[0]), list(LINHAS[1])]
    outra[1][9] = "13:12"                                          # a oficial saiu às 13:12
    planilha(origem / "Programação Semana 41.xlsx", outra)
    r = G.iniciar(config, "2026-W41")
    st = G.aguardar(config, r["id"], 60)
    cmp = st["comparacao"]
    assert cmp["identicas"] is False and cmp["diferencas"].get("Hora Início") == 1


def _linhas_da_planilha(p):
    wb = openpyxl.load_workbook(p, read_only=True)
    out = {}
    for ws in wb.worksheets:
        cab = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        out[ws.title] = [(r[cab.index("OSs ID")], r[cab.index("Tarefa")]) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return out


def test_sem_das_observacoes_tira_a_tarefa_da_planilha_gerada(tmp_path, monkeypatch):
    """Tirar uma tarefa da semana (15/10/2026) é o "sem:". O motor só o aplica nas OS de plano (modo A); na OS comum
    ele programa a tarefa mesmo assim. Depois do motor, a rodada tira da planilha as linhas da tarefa como o robô do PCM
    tira na semana em curso (atualizacao_semanal, STEP A3): em toda aba com "OSs ID" e "Tarefa", a da OS cuja tarefa
    contém a palavra, sem diferenciar maiúscula."""
    p = tmp_path / "sombra.xlsx"
    outra = list(LINHAS[1])
    outra[6] = "MPM trafo"
    vizinha = list(LINHAS[1])
    vizinha[2] = 15022
    planilha(p, [LINHAS[0], LINHAS[1], outra, vizinha],
             pendentes=[["SP Leste 02", 15021, "ALT1-CAB02", "MPM CABINE 2", "Sem capacidade"],
                        ["SP Leste 02", "15020", "", "Outra", ""]])
    obs = tmp_path / "Observacoes_Semana.txt"
    obs.write_text("@usina Altair 1 = seg\n15021; sem: MPM cabine\n15020; não\n", encoding="utf-8")
    assert G.aplicar_sem(obs, p) == {"linhas": 2, "os": ["15021"]}
    assert _linhas_da_planilha(p) == {"SP Leste 02": [(15020, "Fusível queimado"), (15021, "MPM trafo"),
                                                      (15022, "MPM cabine")],
                                      "_Pendentes": [("15020", "Outra")]}
    # sem "sem:" (ou sem linha a tirar), a planilha nem é regravada
    antes = p.read_bytes()
    obs.write_text("15020; não\n", encoding="utf-8")
    assert G.aplicar_sem(obs, p) == {"linhas": 0, "os": []} and p.read_bytes() == antes
    obs.write_text("15021; sem: tarefa que não existe\n", encoding="utf-8")
    monkeypatch.setattr(openpyxl.Workbook, "save", lambda *a, **k: pytest.fail("regravou a planilha sem tirar nada"))
    assert G.aplicar_sem(obs, p) == {"linhas": 0, "os": []} and p.read_bytes() == antes
    obs.unlink()
    assert G.aplicar_sem(obs, p) == {"linhas": 0, "os": []} and p.read_bytes() == antes


def test_sem_da_subtarefa_tira_so_o_ativo_dela(tmp_path):
    """A subtarefa (19/10/2026): "nome @ CÓDIGO" tira só a linha daquele ativo, como o robô do PCM (`_casa_tarefa`); o
    mesmo nome em outro ativo fica. Na _Pendentes do motor a coluna é "Código"."""
    p = tmp_path / "sombra.xlsx"
    outro_ativo = list(LINHAS[1])
    outro_ativo[4] = "ALT1-CAB09"
    planilha(p, [LINHAS[0], LINHAS[1], outro_ativo],
             pendentes=[["SP Leste 02", 15021, "ALT1-CAB01", "MPM cabine", "Sem capacidade"],
                        ["SP Leste 02", 15021, "ALT1-CAB02", "MPM cabine", "Sem capacidade"]])
    obs = tmp_path / "Observacoes_Semana.txt"
    obs.write_text("15021; sem: MPM cabine @ ALT1-CAB01\n", encoding="utf-8")
    assert G.aplicar_sem(obs, p) == {"linhas": 2, "os": ["15021"]}
    wb = openpyxl.load_workbook(p, read_only=True)
    ficam = {ws.title: [(r[2], r[4]) if ws.title != "_Pendentes" else (r[1], r[2])
                        for r in ws.iter_rows(min_row=2, values_only=True)] for ws in wb.worksheets}
    wb.close()
    assert ficam == {"SP Leste 02": [(15020, "ALT1-INV01"), (15021, "ALT1-CAB09")], "_Pendentes": [(15021, "ALT1-CAB02")]}


def test_rodada_aplica_o_sem_das_observacoes(config):
    I.salvar_observacoes(G.pasta_trabalho(config), "2026-W41", "15021; sem: MPM cabine\n", autor="teste")
    st = G.aguardar(config, G.iniciar(config, "2026-W41")["id"], 60)
    assert st["estado"] == "ok", st
    assert st["sem"] == {"linhas": 1, "os": ["15021"]} and st["resumo"]["linhas"] == 1


def test_motor_que_falha_vira_estado_falhou_com_o_motivo(config, tmp_path):
    cfg = dict(config, NEXUS_PCM_MOTOR=str(motor_falso(tmp_path, codigo_saida=1)))
    st = G.aguardar(cfg, G.iniciar(cfg, "2026-W41")["id"], 60)
    assert st["estado"] == "falhou" and "deu ruim" in st["erro"]


def test_uma_geracao_por_vez(config, tmp_path):
    cfg = dict(config, NEXUS_PCM_MOTOR=str(motor_falso(tmp_path, espera=2.0)))
    r = G.iniciar(cfg, "2026-W41")
    with pytest.raises(G.Ocupado):
        G.iniciar(cfg, "2026-W41")
    G.aguardar(cfg, r["id"], 60)


# ── a geração com a FOTO do Fracttal (sem ler o Fracttal) ─────────────────────────────────────────────────────

def motor_da_foto(tmp_path, regrava=False):
    """O motor falso da rodada com foto: confere que ele NÃO tem como ler o Fracttal (sem credencial, endereço numa
    porta fechada) e que a foto está na pasta com a idade de antes (o TTL do motor é pela idade do arquivo). Com
    `regrava`, faz o que o motor de verdade faz quando LÊ o Fracttal: regrava o cache."""
    m = tmp_path / f"motor_foto_{int(regrava)}.py"
    m.write_text(textwrap.dedent(f"""
        import json, os, time
        import openpyxl
        base = os.environ["PCM_PROG_DIR"]
        assert os.environ["FRACTTAL_CLIENT_ID"] == "" and os.environ["FRACTTAL_CLIENT_SECRET"] == ""
        assert os.environ["FRACTTAL_BASE_URL"].startswith("http://127.0.0.1:9/")
        assert int(os.environ["PROG_API_TTL_MIN"]) > 60 * 24 * 365
        foto = os.path.join(base, ".cache_semanal_api.pkl")
        assert time.time() - os.path.getmtime(foto) > 3600, "a foto perdeu a idade"
        assert os.path.exists(os.path.join(base, "Historico_Programacoes.xlsx"))
        if {regrava}:
            open(foto, "ab").write(b"lido de novo")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SP Leste 02"
        ws.append(json.loads({json.dumps(json.dumps(CAB))}))
        for l in json.loads({json.dumps(json.dumps(LINHAS))}):
            ws.append(l)
        wb.save(os.environ["PCM_OUTPUT"])
        print("OK")
    """), encoding="utf-8")
    return m


@pytest.fixture
def foto(origem):
    """A foto que o motor do PCM grava na pasta dele ao ler o Fracttal, com a idade de uma foto de verdade."""
    import os
    import time
    for nome in (".cache_semanal_api.pkl", "_usinas_coordenadas_cache.json"):
        (origem / nome).write_bytes(b"foto de mentira")
        velho = time.time() - 3 * 3600
        os.utime(origem / nome, (velho, velho))
    return origem


def test_gerar_com_foto_nao_le_o_fracttal(config, foto, tmp_path):
    cfg = dict(config, NEXUS_PCM_MOTOR=str(motor_da_foto(tmp_path)))
    st = G.gerar_com_foto(cfg, "2026-W41")
    assert st["estado"] == "ok", st
    assert st["foto"]["intacta"] and set(st["foto"]["arquivos"]) == {".cache_semanal_api.pkl",
                                                                     "_usinas_coordenadas_cache.json"}
    assert st["credencial"].startswith("nenhuma") and st["id"].endswith("-2026-W41-foto")
    assert st["conferencia"]["blocos"] == 2 and st["conferencia"]["fora_da_semana"] == 0
    # é prova, não geração: fica fora da lista da tela
    assert st["id"] not in [u["id"] for u in G.ultimas(cfg)]
    # e a pasta do PCM só foi lida
    assert (foto / ".cache_semanal_api.pkl").read_bytes() == b"foto de mentira"


def test_foto_regravada_pelo_motor_nao_esta_intacta(config, foto, tmp_path):
    """Se o motor tivesse lido o Fracttal, ele regravaria o cache: a rodada diz que a foto não ficou intacta."""
    cfg = dict(config, NEXUS_PCM_MOTOR=str(motor_da_foto(tmp_path, regrava=True)))
    st = G.gerar_com_foto(cfg, "2026-W41", rotulo="Teste 1")
    assert st["estado"] == "ok" and st["foto"]["intacta"] is False
    assert st["id"].endswith("-foto-teste-1")


def test_foto_com_historico_da_reserva(config, foto, tmp_path):
    guardar_reserva(config, [["15021|ALT1-CAB01", "2026-W39", "2026-W41", 2, "2026-W39,2026-W41"]])
    cfg = dict(config, NEXUS_PCM_MOTOR=str(motor_da_foto(tmp_path)))
    st = G.gerar_com_foto(cfg, "2026-W41", historico="nexus")
    assert st["estado"] == "ok" and st["historico"] == "nexus"
    assert historico_da_rodada(cfg, st["id"]) == {"15021|ALT1-CAB01": (1, "2026-W39")}


def test_sem_foto_nao_roda(config, tmp_path):
    with pytest.raises(G.NaoPronto):
        G.gerar_com_foto(config, "2026-W41", foto=tmp_path / "vazia")
    with pytest.raises(ValueError):
        G.gerar_com_foto(config, "2026-W41", historico="planilha")


def test_conferencia_da_semana_gerada(tmp_path):
    """Os casos da W42 de 09/10/2026: a corretiva de MPA na janela da noite não conta no dia; a corretiva forçada que
    passa da meia-noite começa fora da jornada; a tarefa partida ([PARCIAL]) tem dois blocos de propósito."""
    cab = CAB + ["Paralelo (terceirizada)"]
    linhas = [
        ["SP Leste 02", "Segunda-feira (05/10)", 1, "U", "A", "Corretiva", "t1", "Não", 1, "07:30", "11:30", 4, 9, "Não"],
        ["SP Leste 02", "Segunda-feira (05/10) [PARCIAL]", 2, "U", "B", "MPM", "t2", "Não", 1, "13:12", "16:12", 3, 9, "Não"],
        ["SP Leste 02", "Terça-feira (06/10)", 2, "U", "B", "MPM", "t2", "Não", 1, "07:30", "08:00", 0.5, 9, "Não"],
        ["SP Leste 02", "Segunda-feira (05/10) [NOTURNO]", 3, "U", "C", "Corretiva", "MPA x", "Sim", 2, "15:45", "23:45", 8, 9,
         "Não"],
        ["SP Leste 02", "Terça-feira (06/10) [EXCEDE HH]", 4, "U", "D", "Corretiva", "t4", "Não", 1, "03:15 (D+1)",
         "09:15 (D+1)", 6, 9, "Não"],
        ["SP Leste 02", "Terça-feira (06/10) [ZELADORIA/PARALELO]", 5, "U", "E", "Zeladoria", "t5", "Não", 1, "07:30",
         "17:30", 20, 9, "Sim"],
        ["SP Leste 02", "Segunda-feira (12/10)", 6, "U", "F", "MPM", "t6", "Não", 1, "07:30", "08:30", 1, 9, "Não"],
        # o último MPS do dia, cortado pelo fim do expediente: 32 min no relógio e 1,3 h na coluna, sem [PARCIAL]
        ["SP Leste 02", "Quarta-feira (07/10)", 7, "U", "G", "MPS", "t7", "Não", 1, "16:18", "16:50", 1.3, 9, "Não"],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SP Leste 02"
    ws.append(cab)
    for l in linhas:
        ws.append(l)
    wb.create_sheet("_Pendentes").append(["Equipe", "OSs ID", "Código Equipamento", "Tarefa", "Motivo"])
    wb["_Pendentes"].append(["SP Leste 02", 9, "Z", "t9", "não coube"])
    wb.save(tmp_path / "s.xlsx")
    c = C.conferir(tmp_path / "s.xlsx", "2026-W41")
    assert c["blocos"] == 8 and c["pendentes"] == 1 and c["fora_da_semana"] == 1     # o 12/10 é da semana seguinte
    # segunda: 4 h + 3 h (a corretiva de MPA das 15:45 ficou na janela da noite); terça: 0,5 h + 6 h (zeladoria fora).
    # Com os 80% da W42 os dois dias cabiam nas 7 h; com a carga de 70% (16/10/2026, 6,15 h) passam
    assert c["horas_dia_max"] == 7.0 and c["dias_de_equipe"] == 3 and c["capacidade_txt"] == "6,15"
    assert c["acima_da_capacidade"] == 2 and [e["horas"] for e in c["acima_da_capacidade_exemplos"]] == [7.0, 6.5]
    assert c["acima_da_capacidade_pelo_relogio"] == 2 and c["horas_dia_max_pelo_relogio"] == 7.0
    assert c["duracao_maior_que_o_horario"] == 1 and c["duracao_maior_que_o_horario_exemplos"][0]["os"] == "7"
    assert c["horas_noite_max"] == 8
    assert c["fora_da_jornada"] == 1                      # a forçada que começa às 03:15 do dia seguinte
    assert c["tarefas_em_mais_de_um_bloco"] == 1 and c["repetidas_sem_parcial"] == 0
    assert c["marcacoes"]["ZELADORIA (PARALELO)"] == 1 and c["reprogramadas"] == 1
    assert C.entre_semanas(tmp_path / "s.xlsx", tmp_path / "s.xlsx") == {     # a chave do motor: OS + código
        "anterior": 7, "atual": 7, "continuam": 7, "novas": 0, "sairam": 0}
    assert C._no_relogio("11:15", "13:27") == 60 and C._no_relogio("03:15 (D+1)", "09:15 (D+1)") == 360


def test_comparar_planilhas(tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    planilha(a)
    planilha(b)
    assert C.comparar(a, b)["identicas"] is True
    planilha(b, LINHAS[:1], pendentes=[["SP Leste 02", 15021, "ALT1-CAB01", "MPM cabine", "não coube"]])
    r = C.comparar(a, b)
    assert r["identicas"] is False and r["so_a"] == 1 and r["pendentes_b"] == 1


# ── telas ────────────────────────────────────────────────────────────────────────────────────────────────

def test_tela_gerar_mostra_o_que_falta(app, logado, config):
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    html = logado.get("/t/pcm/gerar").get_data(as_text=True)
    assert "Gerar a semana" in html and "2026-W" in html
    assert "Credencial do Fracttal" in html and "Em construção" not in html


def test_tela_gerar_roda_e_mostra_o_resultado(app, logado, config, origem):
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    planilha(origem / "Programação Semana 41.xlsx")
    resp = logado.post("/t/pcm/gerar", data={"semana": "2026-W41", "confirmo": "1"})
    assert resp.status_code == 303
    rodada = resp.headers["Location"].split("rodada=")[1]
    G.aguardar(app.config, rodada, 60)
    html = logado.get("/t/pcm/gerar?rodada=" + rodada).get_data(as_text=True)
    assert "Idênticas" in html
    estado = json.loads(logado.get("/t/pcm/gerar/estado?rodada=" + rodada).get_data(as_text=True))
    assert estado["estado"] == "ok"


def test_tela_gerar_sem_credencial_nao_roda(app, logado, config):
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    app.config["NEXUS_PCM_FRACTTAL_TESTE"] = {}
    resp = logado.post("/t/pcm/gerar", data={"semana": "2026-W41", "confirmo": "1"})
    assert resp.status_code == 400
    assert "Credencial do Fracttal" in resp.get_data(as_text=True)


def test_semana_invalida_nao_roda(app, logado, config):
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    assert logado.post("/t/pcm/gerar", data={"semana": "41; rm -rf", "confirmo": "1"}).status_code == 400


def test_a_conferencia_usa_a_carga_maxima_do_motor():
    # 16/10/2026 (Levi): "carga máxima no programador de sexta de 70% HH por dia" (era 80%, 7 h). A conferência do Nexus
    # conta os dias de equipe acima da carga com a MESMA conta do motor: carga nova no motor do PCM quebra este teste.
    txt = G.MOTOR.read_text(encoding="utf-8")
    m = re.search(r"^CARGA_MAX_DIA\s*=\s*([0-9.]+)", txt, re.M)
    assert m and float(m.group(1)) == C.CARGA_MAX_DIA
    assert re.search(r"^USEFUL_DAY_MIN\s*=.*\*\s*CARGA_MAX_DIA\)", txt, re.M)
    assert re.search(r"^MPA_DAY_USEFUL_MIN\s*=.*\*\s*CARGA_MAX_DIA\)", txt, re.M)
    assert re.search(r"^NIGHT_CAPACITY_MIN\s*=.*\*\s*0\.80\)", txt, re.M)      # a noite ficou nos 80%
    assert (C.DIA_DO_MOTOR_MIN, C.CAPACIDADE_H, C.CAPACIDADE_TXT) == (369, 6.15, "6,15")


def test_motor_le_o_fracttal_a_1_pedido_por_segundo_em_todos_os_leitores(monkeypatch, tmp_path):
    # W43 de prova, 16/10/2026 15:14: o Nexus só mandava a FRACTTAL_INTERVALO_S; a leitura da semana (fonte_bd_api,
    # 5 páginas em paralelo) lê a PROG_API_INTERVALO_S e, sem ela, andou a 0,42 s e levou HTTP 429 em 9 s. Toda variável
    # de ritmo que um leitor do motor lê tem de ir a 1 s, inclusive a que o ambiente do processo já tiver.
    monkeypatch.setenv("PROG_API_INTERVALO_S", "0.1")
    env = G._ambiente(G.Credencial(True, "teste", "id", "seg"), tmp_path, tmp_path / "s.xlsx")
    lidas = set()
    for leitor in (G.MOTOR.parent).glob("*.py"):
        lidas |= set(re.findall(r'os\.environ\.get\(\s*"([A-Z_]*INTERVALO_S)"', leitor.read_text(encoding="utf-8")))
    assert {"FRACTTAL_INTERVALO_S", "PROG_API_INTERVALO_S"} <= lidas
    assert {v: env.get(v) for v in lidas} == {v: "1.0" for v in lidas}
    assert G._ambiente_foto(tmp_path, tmp_path / "s.xlsx")["PROG_API_INTERVALO_S"] == "1.0"


def test_fuso_do_motor_so_vai_pelo_tz_fora_do_windows(monkeypatch, tmp_path):
    # S41 de teste (02/10/2026): no Windows o TZ=America/Sao_Paulo deixava o relógio do motor 4 h adiantado
    cred = G.Credencial(True, "teste", "id", "seg")
    monkeypatch.setenv("TZ", "UTC")
    monkeypatch.setattr(G.os, "name", "nt")
    assert "TZ" not in G._ambiente(cred, tmp_path, tmp_path / "s.xlsx")
    monkeypatch.setattr(G.os, "name", "posix")
    assert G._ambiente(cred, tmp_path, tmp_path / "s.xlsx")["TZ"] == "America/Sao_Paulo"

def test_baixar_a_planilha_da_rodada(app, logado, config, origem):
    """Levi, 19/10/2026: "tem que ter a opção de baixar!". A sombra.xlsx de cada rodada (a mesma que o Publicar no App
    manda como "Programação Semana NN.xlsx") baixa pela tela: o botão na rodada aberta e em cada linha das Últimas
    gerações. Rodada que não existe, id fora do formato ou rodada sem a planilha: 404."""
    from urllib.parse import unquote
    app.config.update({k: v for k, v in config.items() if k != "TESTING"})
    planilha(origem / "Programação Semana 41.xlsx")
    rodada = logado.post("/t/pcm/gerar", data={"semana": "2026-W41", "confirmo": "1"}).headers["Location"].split(
        "rodada=")[1]
    G.aguardar(app.config, rodada, 60)
    url = f"/t/pcm/gerar/planilha?rodada={rodada}"
    html = logado.get("/t/pcm/gerar?rodada=" + rodada).get_data(as_text=True)
    assert html.count(f'href="{url}"') == 2 and "Baixar a planilha" in html      # na rodada e na linha da tabela
    r = logado.get(url)
    assert r.status_code == 200 and r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    disp = unquote(r.headers["Content-Disposition"])
    assert "attachment" in disp and f"Programação Semana 41 - rodada {rodada[:15]}.xlsx" in disp
    sombra = G.pasta_trabalho(app.config) / "geracoes" / rodada / "saida" / "sombra.xlsx"
    assert r.data == sombra.read_bytes()
    r.close()                      # no Windows, a resposta segura o arquivo aberto até fechar
    for ruim in ("../../etc/passwd", "20261019-102300-2026-W41", "", rodada + "x"):
        assert logado.get("/t/pcm/gerar/planilha", query_string={"rodada": ruim}).status_code == 404, ruim
    sombra.unlink()
    assert logado.get(url).status_code == 404
    assert f'href="{url}"' not in logado.get("/t/pcm/gerar?rodada=" + rodada).get_data(as_text=True)
