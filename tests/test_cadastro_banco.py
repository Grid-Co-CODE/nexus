"""O cadastro no banco (nexus/cadastro/banco.py): ID numérico, ligação por ID, sensível cifrado e nada recifrado à toa."""
import io
import json

import openpyxl
import pytest

from nexus.cadastro import banco as B
from nexus.cadastro.armazem import ArmazemLocal
from nexus.cadastro.cifra import Cofre, gerar_chave
from nexus.cadastro.servico import Carga, Servico
from nexus.cadastro.tipos import Legado

CPF, EMAIL, FONE, CNPJ = "52998224725", "ana.lima@gridco.com.br", "11912345678", "12345678000199"


@pytest.fixture
def cofre():
    return Cofre(gerar_chave())


@pytest.fixture
def srv(tmp_path, cofre):
    s = Servico(ArmazemLocal(tmp_path / "c.json"), cofre)
    s.aplicar_carga(Carga(
        entidades={
            "clientes": [{"id": "1", "ordem": 1, "valores": {"nome": "Thopen"}},
                         {"id": "2", "ordem": 2, "valores": {"nome": "Athon"}}],
            "equipes": [{"id": "1", "ordem": 1, "valores": {"nome": "SP Norte 02"}}],
            "pessoas": [{"id": "7", "ordem": 1, "valores": {
                "nome": "Ana Maria Lima", "cargo": "Técnico O&M", "equipe": "1", "status": "Ativo", "cpf": CPF,
                "email": EMAIL, "telefone_corporativo": FONE, "vinculo": "Colaborador de campo"}}],
            "usinas": [
                {"id": "1", "ordem": 1, "valores": {
                    "nome": "Altair", "cliente": "1", "status": "OPERAÇÃO", "equipe": "1", "responsavel_om": "7",
                    "codigo": "THPN-ALT100", "id_bd": "UFV-001", "potencia_contratual": 6.898, "cnpj": CNPJ,
                    "receita_mensal": 11745.0, "latitude": -20.5, "n_inversores": "N/A", "ucs": 3012345678}},
                {"id": "2", "ordem": 2, "valores": {
                    "nome": "Altair", "cliente": "2", "status": "A MOBILIZAR", "responsavel_om": Legado("Fulano Antigo"),
                    "codigo": "ATHN-MAB100"}},
            ],
        },
        listas={"status_usina": ["OPERAÇÃO", "A MOBILIZAR"], "cargos": ["Técnico O&M"], "uf": ["SP"],
                "status_contratacao": ["Ativo"], "vinculo": ["Colaborador de campo"]},
        origem={"arquivo": "teste.xlsx"}), quem="importação")
    return s


def _tab(t, aba):
    cab, linhas = t[aba]
    return [dict(zip(cab, l)) for l in linhas]


def test_id_numerico_e_ligacao_por_id(srv, cofre):
    t = B.tabelas(srv, cofre)
    us = {u["usina_id"]: u for u in _tab(t, "usinas")}
    assert list(t["usinas"][0][:1]) == ["usina_id"] and set(us) == {1, 2}
    assert us[1]["cliente_id"] == 1 and us[1]["equipe_id"] == 1 and us[1]["responsavel_om_id"] == 7
    # o inner join fecha: toda chave apontada existe na tabela dela
    ids = {aba: {r[c] for r in _tab(t, aba)} for _, aba, c in B.TABELAS}
    for u in us.values():
        assert u["cliente_id"] in ids["clientes"]
    assert us[1]["n_inversores"] is None                       # "N/A" numa coluna de número vira vazio
    assert us[2]["responsavel_om_id"] is None and t["_sem_id"] == 1   # nome que não casou: sem ID, não inventa


def test_calculo_com_erro_vira_vazio(srv, cofre, monkeypatch):
    """Achado no ensaio real (04/10): "Cidade/Estado (base)" dava #ERRO numa pessoa e o envio inteiro quebrava."""
    from nexus.cadastro import calculos
    monkeypatch.setattr(calculos, "cidade_estado_base", lambda *a: calculos.ERRO)
    srv._cache = None
    p = _tab(B.tabelas(srv, cofre), "pessoas")[0]
    assert "cidade_estado_base" not in json.loads(cofre.decifrar(p[B.COL_CIFRA], "banco/pessoas/7"))


def test_sensivel_e_pessoa_nao_vao_em_claro(srv, cofre):
    t = B.tabelas(srv, cofre)
    tudo = json.dumps({k: v for k, v in t.items() if k != "_sem_id"}, ensure_ascii=False, default=str)
    # "ucs": no BD é o número da UC (achado no 1º envio real), não a quantidade
    for segredo in (CPF, EMAIL, FONE, CNPJ, "Ana Maria Lima", "11745", "Fulano Antigo", "-20.5", "3012345678"):
        assert segredo not in tudo, segredo
    p = _tab(t, "pessoas")[0]
    assert p["pessoa_id"] == 7 and p["equipe_id"] == 1 and p["cargo"] == "Técnico O&M"
    corpo = json.loads(cofre.decifrar(p[B.COL_CIFRA], "banco/pessoas/7"))
    assert corpo["cpf"] == CPF and corpo["email"] == EMAIL
    u2 = _tab(t, "usinas")[1]
    assert json.loads(cofre.decifrar(u2[B.COL_CIFRA], "banco/usinas/2"))["responsavel_om_texto"] == "Fulano Antigo"


def test_cifra_que_nao_mudou_e_reaproveitada(srv, cofre):
    """Nonce aleatório: recifrar daria outro texto e a API gravaria histórico em toda linha, a cada envio."""
    t1 = B.tabelas(srv, cofre)
    ant = {"usinas": {r["usina_id"]: r[B.COL_CIFRA] for r in _tab(t1, "usinas")},
           "pessoas": {r["pessoa_id"]: r[B.COL_CIFRA] for r in _tab(t1, "pessoas")}}
    t2 = B.tabelas(srv, cofre, ant)
    assert _tab(t2, "usinas")[0][B.COL_CIFRA] == ant["usinas"][1]
    assert _tab(t2, "pessoas")[0][B.COL_CIFRA] == ant["pessoas"][7]
    outro = {"usinas": {1: Cofre(gerar_chave()).cifrar("{}", "banco/usinas/1")}}
    assert _tab(B.tabelas(srv, cofre, outro), "usinas")[0][B.COL_CIFRA] != outro["usinas"][1]


def test_de_para_casa_por_codigo_e_por_nome_com_cliente(srv):
    fontes = {"BD_Performance": [{"chave": "MAB100", "codigo": "MAB100"}, {"chave": "XYZ999", "codigo": "XYZ999"}],
              "Fracttal": [{"chave": "Thopen - Altair - SP", "nome": "Thopen - Altair - SP"},
                           {"chave": "Altair", "nome": "Altair"}]}           # "Altair" sozinho serve a duas usinas
    dp = B.de_para(srv, fontes)
    assert [2, "BD_Performance", "MAB100", "código"] in dp
    assert [1, "Fracttal", "Thopen - Altair - SP", "nome"] in dp
    # sem par ou ambíguo: não liga, mas fica na tabela dizendo por quê (a conta do "% ligado" sai daqui)
    assert [None, "BD_Performance", "XYZ999", B.SEM_PAR] in dp and [None, "Fracttal", "Altair", B.SEM_PAR] in dp
    assert [1, "BD_Operações", "UFV-001", "id do BD"] in dp


def test_codigo_que_serve_a_duas_usinas_nao_liga():
    """As bases da API gravam o código sem o prefixo do cliente: "THPN-ALT100" e "ATHN-ALT100" viram os dois "ALT100"."""
    class U:
        def __init__(self, id_, cod):
            self.id, self._v = id_, {"codigo": cod, "nome": f"Usina {id_}", "cliente": None, "id_bd": None}

        def valor(self, c):
            return self._v.get(c)

    class S:
        def registros(self, ent):
            return [U("1", "THPN-ALT100"), U("2", "ATHN-ALT100"), U("3", "THPN-JCD100")]

        def titulo_de(self, ent, ref):
            return ""

    dp = B.de_para(S(), {"API": [{"chave": "ALT100", "codigo": "ALT100"}, {"chave": "JCD100", "codigo": "JCD100"}]})
    assert dp == [[3, "API", "JCD100", "código"], [None, "API", "ALT100", B.SEM_PAR]]
    # com o prefixo do cliente, o mesmo "ALT100" deixa de ser ambíguo (caso real: 2C-IPX100 x THPN-IPX100)
    dp = B.de_para(S(), {"Fracttal": [{"chave": "Athon - Usina 2", "codigo": "ATHN-ALT100"}]})
    assert dp == [[2, "Fracttal", "Athon - Usina 2", "código"]]


def test_regra_de_ignorar_tira_da_conta_e_diz_o_motivo(srv):
    """Decisões do Levi (04/10): "TESTE é teste"; "Porteiras não entra no BD"."""
    regras = [{"contem": "teste", "motivo": "teste"},
              {"contem": "Porteiras", "sistema": "Fracttal", "motivo": "não entra no BD"}]
    fontes = {"Fracttal": [{"chave": "Thopen - Teste 1 - PA", "nome": "Thopen - Teste 1 - PA"},
                           {"chave": "(sem Classificação 1) TESTE100", "codigo": "TESTE100"},
                           {"chave": "Thopen - Porteiras 1 - CE", "codigo": "THPN-PRT100"}],
              "BD_Thopen": [{"chave": "Porteiras 1", "nomes": ["Porteiras 1"], "cliente": "Thopen"}]}
    dp = B.de_para(srv, fontes, regras)
    assert [None, "Fracttal", "Thopen - Teste 1 - PA", "ignorado: teste"] in dp
    assert [None, "Fracttal", "(sem Classificação 1) TESTE100", "ignorado: teste"] in dp
    assert [None, "Fracttal", "Thopen - Porteiras 1 - CE", "ignorado: não entra no BD"] in dp
    assert [None, "BD_Thopen", "Porteiras 1", B.SEM_PAR] in dp          # a regra era só do Fracttal


def test_base_sem_codigo_casa_por_nome_so_dentro_do_cliente(srv):
    """E1 e Thopen têm usinas de mesmo nome e são usinas diferentes (Levi, 04/10): o casamento por nome não cruza
    cliente. Aqui "Altair" existe na Thopen (1) e na Athon (2)."""
    fontes = {"BD_Thopen": [{"chave": "Altair", "nomes": ["Altair"], "cliente": "Thopen"}],
              "Outra": [{"chave": "Altair", "nomes": ["Altair"], "cliente": "E1"}]}
    dp = B.de_para(srv, fontes)
    assert [1, "BD_Thopen", "Altair", "nome"] in dp
    assert [None, "Outra", "Altair", B.SEM_PAR] in dp


def test_dica_do_de_para_de_trackers_liga_ou_trava(srv):
    """A dica (o de-para de trackers, já conferido) liga quando o nome não casa e trava quando aponta outra usina."""
    fr = [{"chave": "Thopen - Altair 1 - SP", "codigo": "THPN-ALT100"}]
    so_dica = {"Fracttal": fr, "BD_Thopen": [{"chave": "Santo Inácio XII", "nomes": ["Santo Inácio XII"],
                                              "cliente": "Thopen", "dica": "Thopen - Altair 1 - SP"}]}
    assert [1, "BD_Thopen", "Santo Inácio XII", "de-para de trackers"] in B.de_para(srv, so_dica)
    contra = {"Fracttal": [{"chave": "Athon - Altair - XX", "codigo": "ATHN-MAB100"}],
              "BD_Thopen": [{"chave": "Altair", "nomes": ["Altair"], "cliente": "Thopen", "dica": "Athon - Altair - XX"}]}
    assert [None, "BD_Thopen", "Altair", "conflito com o de-para de trackers"] in B.de_para(srv, contra)


@pytest.mark.parametrize("bruto, esperado", [
    ("MAB100-INVR2.4", "MAB100"), ("THPN-SDI100-INVR11.1", "THPN-SDI100"), ("2C-IPX100-TR1", "2C-IPX100"),
    ("FRN100-CABN2", "FRN100"), ("ALLM-QGBT", None), ("", None), (None, None)])
def test_codigo_da_usina_nos_dois_formatos_do_fracttal(bruto, esperado):
    """Ler só o 1º pedaço tomava "THPN" por código: a medição de 04/10 deu 47% em vez de 93%."""
    assert B.codigo_do_equipamento(bruto) == esperado


def test_planilha_tem_cabecalho_e_nao_tem_texto_vazio(srv, cofre):
    wb = openpyxl.load_workbook(io.BytesIO(B.xlsx_bytes(B.montar(srv, cofre))))
    # a região de campo (estrutura de O&M de 10/2026) entrou como aba nova, depois das de antes
    assert wb.sheetnames == ["clientes", "equipes", "pessoas", "usinas", "regioes_campo", "de_para", "atualizacao"]
    assert next(wb["usinas"].iter_rows(values_only=True))[0] == "usina_id"
    assert all(v != "" for ws in wb for row in ws.iter_rows(values_only=True) for v in row)


class _Sessao:
    def __init__(self, existe):
        self.existe, self.chamadas = existe, []

    def _r(self, corpo):
        class R:
            def raise_for_status(self):
                pass

            def json(self):
                return corpo
        return R()

    def get(self, url, **k):
        self.chamadas.append(("GET", url, k))
        return self._r([{"key": B.WORKBOOK}] if self.existe else [])

    def post(self, url, **k):
        self.chamadas.append(("POST", url, k))
        return self._r({"sheets": 6})


def test_sincronizar_cria_uma_vez_e_manda_com_token():
    nova = _Sessao(existe=False)
    B.sincronizar(b"x", base="https://api/", token="tok", sessao=nova)
    assert [c[1] for c in nova.chamadas if c[0] == "POST"] == ["https://api/api/workbooks",
                                                                f"https://api/api/workbooks/{B.WORKBOOK}/sync-xlsx"]
    assert nova.chamadas[-1][2]["params"] == {"replace": "true"}
    assert nova.chamadas[-1][2]["headers"]["Authorization"] == "Bearer tok"
    ja = _Sessao(existe=True)
    B.sincronizar(b"x", base="https://api", token="tok", sessao=ja)
    assert [c[1] for c in ja.chamadas if c[0] == "POST"] == [f"https://api/api/workbooks/{B.WORKBOOK}/sync-xlsx"]
